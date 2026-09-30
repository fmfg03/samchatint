"""DR paper balance and provenance from the accountant's sample layout."""

from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from devnous.gastos.routes import user_routes
from devnous.gastos.services.informe_poliza_workpaper import (
    InformeWorkpaperExpense,
    generate_informe_poliza_workpaper,
)


def _expense(**changes):
    values = dict(
        source_id="expense-1",
        reference="G-1",
        date="2026-09-28",
        description="Gasolina / F-1",
        amount=700.0,
        vat=94.12,
        expense_account="5100-006-004",
        vat_account="1200-001-001",
        counterpart_account="1170-001-009",
        company_amex=False,
        cfdi_uuid="CFDI-1",
    )
    values.update(changes)
    return InformeWorkpaperExpense(**values)


def test_workpaper_has_one_dr_with_live_balance_and_traceable_expenses():
    payload = generate_informe_poliza_workpaper(
        [_expense(), _expense(
            source_id="expense-2", reference="G-2", amount=200.0,
            vat=None, expense_account="", counterpart_account="",
        )],
        reference="REF 3", title="Gastos Edgar", currency="MXN",
    )
    wb = load_workbook(BytesIO(payload))
    dr = wb["Papel DR"]
    assert (dr["A3"].value, dr["B3"].value, dr["D3"].value) == ("Dr", 3, 4)
    assert dr["F4"].value == 605.88
    assert dr["F5"].value == 94.12
    assert dr["F6"].value == 200.0
    assert dr["B6"].value is None
    assert dr["G7"].value == 900.0
    assert dr["B7"].value is None  # Mixed/missing counterparts are reviewed.
    assert dr["B8"].value == "FIN_PARTIDAS"
    assert wb["Cuadre"]["B5"].value == "=SUM('Papel DR'!F4:F7)"
    assert wb["Cuadre"]["B6"].value == "=SUM('Papel DR'!G4:G7)"
    assert wb["Cuadre"]["B7"].value == "=ROUND(B5-B6,2)"
    assert "COUNTBLANK" in wb["Cuadre"]["B8"].value
    assert "B8=0" in wb["Cuadre"]["B9"].value
    assert [wb["Origen y revisión"].cell(i, 1).value for i in (2, 3)] == [
        "expense-1", "expense-2",
    ]


def test_workpaper_keeps_amex_liability_separate_and_does_not_guess_iva():
    payload = generate_informe_poliza_workpaper(
        [_expense(), _expense(
            source_id="expense-2", amount=250, vat=None,
            company_amex=True, counterpart_account="2120-002-062",
        )],
        reference="REF 3", title="Mixto", currency="MXN",
    )
    dr = load_workbook(BytesIO(payload))["Papel DR"]
    assert dr["F6"].value == 250.0
    assert dr["G7"].value == 700.0
    assert dr["G8"].value == 250.0
    assert dr["B7"].value == "1170-001-009"
    assert dr["B8"].value == "2120-002-062"


@pytest.mark.parametrize("change", [dict(amount=0), dict(vat=800)])
def test_invalid_amounts_do_not_produce_a_plausible_dr(change):
    with pytest.raises(ValueError, match="Importe o IVA inválido"):
        generate_informe_poliza_workpaper(
            [_expense(**change)], reference="REF 3", title="X", currency="MXN"
        )


@pytest.mark.asyncio
async def test_finance_route_includes_unclassified_expenses_without_writing(monkeypatch):
    cuenta_id = uuid4()
    owner_id = uuid4()
    cuenta = SimpleNamespace(
        id=cuenta_id, empleado_id=owner_id, currency="MXN",
        nombre="Gastos Edgar", referencia_base="IG-3",
    )
    expense = SimpleNamespace(
        id=uuid4(), numero_referencia="G-1", fecha=None,
        concepto="Gasolina", gasto_cantidad=700, iva=94.12,
        cuenta_contable=None, cuenta_iva=None, contra_cuenta_contable=None,
        cfdi_report=None, currency="MXN", pagado_con_amex_empresa=False,
    )
    session = SimpleNamespace(
        get=AsyncMock(return_value=cuenta),
        execute=AsyncMock(return_value=SimpleNamespace(
            scalars=lambda: SimpleNamespace(all=lambda: [expense])
        )),
    )
    monkeypatch.setattr(
        user_routes, "_informe_documento_for_cuenta",
        AsyncMock(return_value=SimpleNamespace(referencia_operaciones="3")),
    )
    beneficiary = SimpleNamespace(id=uuid4(), nombre="Edgar Ejemplo")
    debtor_account = SimpleNamespace(
        id=uuid4(), codigo="1170-001-009", nombre="Edgar Ejemplo"
    )
    monkeypatch.setattr(
        user_routes,
        "resolve_cuenta_debtor_empleado",
        AsyncMock(return_value=beneficiary),
    )
    monkeypatch.setattr(
        user_routes,
        "resolve_cuenta_debtor_account",
        AsyncMock(return_value=debtor_account),
    )
    actor = SimpleNamespace(id=uuid4(), rol="finanzas")

    response = await user_routes.exportar_papel_poliza_informe(
        cuenta_id=cuenta_id, session=session, current_empleado=actor
    )

    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.body))
    assert wb["Papel DR"]["F4"].value == 605.88
    assert wb["Papel DR"]["B4"].value is None
    assert wb["Papel DR"]["B6"].value == "1170-001-009"
    assert wb["Origen y revisión"]["A2"].value == str(expense.id)
    assert wb["Origen y revisión"]["K2"].value == "1170-001-009"
    session.execute.assert_awaited_once()


@pytest.mark.asyncio
async def test_finance_route_rejects_missing_informe(monkeypatch):
    cuenta_id = uuid4()
    cuenta = SimpleNamespace(id=cuenta_id, empleado_id=uuid4())
    session = SimpleNamespace(get=AsyncMock(return_value=cuenta))
    monkeypatch.setattr(
        user_routes, "_informe_documento_for_cuenta", AsyncMock(return_value=None)
    )
    with pytest.raises(HTTPException) as error:
        await user_routes.exportar_papel_poliza_informe(
            cuenta_id=cuenta_id,
            session=session,
            current_empleado=SimpleNamespace(id=uuid4(), rol="finanzas"),
        )
    assert error.value.status_code == 409


@pytest.mark.asyncio
async def test_finance_route_blocks_generic_counterpart_without_beneficiary_detail(
    monkeypatch,
):
    cuenta_id = uuid4()
    cuenta = SimpleNamespace(
        id=cuenta_id,
        empleado_id=uuid4(),
        currency="MXN",
        nombre="Gastos",
        referencia_base="IG-4",
    )
    expense = SimpleNamespace(
        id=uuid4(),
        numero_referencia="G-2",
        fecha=None,
        concepto="Gasolina",
        gasto_cantidad=100,
        iva=None,
        cuenta_contable=SimpleNamespace(codigo="5100-006-004"),
        cuenta_iva=None,
        contra_cuenta_contable=SimpleNamespace(codigo="2120-000-000"),
        cfdi_report=None,
        currency="MXN",
        pagado_con_amex_empresa=False,
    )
    session = SimpleNamespace(
        get=AsyncMock(return_value=cuenta),
        execute=AsyncMock(
            return_value=SimpleNamespace(
                scalars=lambda: SimpleNamespace(all=lambda: [expense])
            )
        ),
    )
    monkeypatch.setattr(
        user_routes,
        "_informe_documento_for_cuenta",
        AsyncMock(return_value=SimpleNamespace(referencia_operaciones="4")),
    )
    beneficiary = SimpleNamespace(id=uuid4(), nombre="Sin subcuenta")
    monkeypatch.setattr(
        user_routes,
        "resolve_cuenta_debtor_empleado",
        AsyncMock(return_value=beneficiary),
    )
    monkeypatch.setattr(
        user_routes,
        "resolve_cuenta_debtor_account",
        AsyncMock(return_value=None),
    )

    with pytest.raises(HTTPException) as error:
        await user_routes.exportar_papel_poliza_informe(
            cuenta_id=cuenta_id,
            session=session,
            current_empleado=SimpleNamespace(id=uuid4(), rol="finanzas"),
        )

    assert error.value.status_code == 409
    assert "subcuenta contable de detalle" in error.value.detail
    assert "1170-001" in error.value.detail
