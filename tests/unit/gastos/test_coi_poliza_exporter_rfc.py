"""Real COI XLSX RFC cells must not contain layout padding."""

import csv
import io
import zipfile
from datetime import datetime

import pytest
from openpyxl import load_workbook

from devnous.gastos.services.coi_poliza_exporter import (
    ExpenseCFDI,
    generate_coi_poliza_csv,
    generate_coi_poliza_xlsx,
    generate_coi_poliza_zip,
)


def expense(document_type, emitter, receiver):
    return ExpenseCFDI(
        fecha=datetime(2026, 10, 14),
        total=116,
        iva_amount=16,
        subtotal_amount=100,
        concepto="Prueba de RFC",
        cuenta_contable="5300-001-001",
        cuenta_contrapartida="2100-001-001",
        cuenta_iva="1180-001-001",
        cfdi_uuid="10000000-0000-4000-8000-000000000001",
        rfc_emisor=emitter,
        rfc_receptor=receiver,
        poliza_group_key=document_type,
        poliza_reference=document_type,
        poliza_description="Prueba " + document_type,
    )


def rfcs(content, uuid="10000000-0000-4000-8000-000000000001"):
    sheet = load_workbook(io.BytesIO(content))["Poliza COI"]
    row = next(row for row in sheet if row[8].value == uuid)
    return tuple(row[column].value for column in (5, 6))


@pytest.mark.parametrize("document_type", ["SOLICITUD", "INFORME", "AMEX"])
@pytest.mark.parametrize("padding", ["", " ", "\t", "\u00a0"])
def test_xlsx_rfc_cells_are_clean_without_mutating_source(document_type, padding):
    emitter, receiver = "AAA010101AAA", "BBB010101BBB"
    item = expense(
        document_type, padding + emitter + padding, padding + receiver + padding
    )
    original = (item.rfc_emisor, item.rfc_receptor)
    assert rfcs(generate_coi_poliza_xlsx([item])) == (emitter, receiver)
    assert (item.rfc_emisor, item.rfc_receptor) == original


@pytest.mark.parametrize("value", [None, "", "  ", "\t\u00a0"])
def test_xlsx_missing_rfc_stays_empty(value):
    assert rfcs(generate_coi_poliza_xlsx([expense("INFORME", value, value)])) == (
        None,
        None,
    )


def test_zip_workbooks_also_remove_rfc_padding():
    items = [
        expense(kind, " AAA010101AAA ", " BBB010101BBB ")
        for kind in ("SOLICITUD", "INFORME")
    ]
    with zipfile.ZipFile(io.BytesIO(generate_coi_poliza_zip(items))) as archive:
        assert len(archive.namelist()) == 2
        for name in archive.namelist():
            assert rfcs(archive.read(name)) == ("AAA010101AAA", "BBB010101BBB")


def test_xlsx_normalization_does_not_change_csv_layout():
    item = expense("INFORME", " AAA010101AAA ", " BBB010101BBB ")
    before = generate_coi_poliza_csv([item])
    generate_coi_poliza_xlsx([item])
    assert generate_coi_poliza_csv([item]) == before
    rows = list(csv.reader(io.StringIO(before.decode("latin-1"))))
    marker = next(index for index, row in enumerate(rows) if row[2] == "INICIO_CFDI")
    assert rows[marker + 1][5:7] == ["  AAA010101AAA ", "  BBB010101BBB "]


def test_marker_text_in_policy_description_does_not_change_movements():
    item = expense("INFORME", "AAA010101AAA", "BBB010101BBB")
    item.poliza_description = "INICIO_CFDI"
    sheet = load_workbook(io.BytesIO(generate_coi_poliza_xlsx([item])))["Poliza COI"]
    assert sheet["C3"].value == "INICIO_CFDI"
    movement = next(row for row in sheet if row[2].value == "0")
    assert isinstance(movement[5].value, (int, float))
    assert movement[5].value == 100
    assert rfcs(generate_coi_poliza_xlsx([item])) == (
        "AAA010101AAA",
        "BBB010101BBB",
    )


def test_rfc_cleanup_keeps_formula_like_source_as_text():
    item = expense("SOLICITUD", " =1+1 ", "BBB010101BBB")
    sheet = load_workbook(io.BytesIO(generate_coi_poliza_xlsx([item])))["Poliza COI"]
    record = next(row for row in sheet if row[8].value == item.cfdi_uuid)
    assert record[5].data_type == "s"
    assert record[5].value == "'=1+1"
    assert item.rfc_emisor == " =1+1 "
