from io import BytesIO

from openpyxl import load_workbook

from devnous.gastos.services.payment_run_exporter import (
    generate_payment_run_order_xlsx,
)


def test_payment_order_export_contains_cutoff_instructions_and_payable_amount() -> None:
    payload = generate_payment_run_order_xlsx(
        closure={
            "id": "0f702abc-3341-4ad6-bcc5-ef697baf235a",
            "run_date": "2026-09-07",
            "total_amount": "56020.00",
            "items": [
                {
                    "numero_referencia": "S-260005",
                    "referencia_operaciones": "101",
                    "solicitante": "Miguel Soria",
                    "beneficiario": "Miguel Soria",
                    "banco": "Santander",
                    "cuenta_bancaria": "1234567890",
                    "cuenta_clabe": "012345678901234567",
                    "concepto_pago": "Reembolso de saldo a favor",
                    "fecha_pago": "2026-09-08",
                    "currency": "MXN",
                    "monto": "56020.00",
                    "payment_data_status": "Listo",
                }
            ],
        }
    )

    workbook = load_workbook(BytesIO(payload), data_only=True)
    sheet = workbook["Orden de pago"]

    assert sheet["A7"].value == "S-260005"
    assert sheet["B7"].value == "101"
    assert sheet["G7"].value == "012345678901234567"
    assert sheet["K7"].value == 56020.00
    assert sheet["L7"].value == "Listo"


def test_payment_order_export_marks_missing_payment_instructions() -> None:
    payload = generate_payment_run_order_xlsx(
        closure={
            "id": "0f702abc-3341-4ad6-bcc5-ef697baf235a",
            "items": [
                {
                    "numero_referencia": "S-260006",
                    "currency": "MXN",
                    "monto": "900.00",
                    "payment_data_status": "Falta banco, cuenta o CLABE.",
                }
            ],
        }
    )

    workbook = load_workbook(BytesIO(payload), data_only=True)
    assert workbook["Orden de pago"]["L7"].value == "Falta banco, cuenta o CLABE."


def test_payment_order_export_escapes_formula_like_text() -> None:
    payload = generate_payment_run_order_xlsx(
        closure={
            "id": "0f702abc-3341-4ad6-bcc5-ef697baf235a",
            "items": [
                {
                    "numero_referencia": "=HYPERLINK(\"https://example.test\")",
                    "currency": "MXN",
                    "monto": "1.00",
                    "payment_data_status": "Listo",
                }
            ],
        }
    )

    workbook = load_workbook(BytesIO(payload), data_only=False)
    assert workbook["Orden de pago"]["A7"].value.startswith("'=")
