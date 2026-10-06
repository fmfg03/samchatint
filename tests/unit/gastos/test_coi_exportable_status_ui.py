import csv
import io
import zipfile
from datetime import datetime
from html import unescape
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.dialects import postgresql

from devnous.gastos.models import Documento, ExpenseReport
from devnous.gastos.routes import admin_routes, user_routes
from devnous.gastos.services.coi_poliza_exporter import (
    ExpenseCFDI,
    generate_coi_poliza_csv,
    generate_coi_poliza_xlsx,
    generate_coi_poliza_zip,
)
from devnous.gastos.services.expense_coi_export_service import (
    expense_coi_batch_period_condition,
    informe_coi_period_condition,
)


class _ScalarRows:
    def __init__(self, rows):
        self.rows = rows

    def scalar_one_or_none(self):
        return self.rows[0] if self.rows else None

    def scalars(self):
        return self

    def all(self):
        return self.rows


@pytest.mark.parametrize("relation", ["documento", "informe", "cuenta"])
@pytest.mark.parametrize("amex", [False, True])
def test_cross_month_policy_is_discovered_in_only_one_batch_period(relation, amex):
    """Execute both discovery predicates, including legacy ownership, in SQL."""
    report_id, account_id = uuid4(), uuid4()
    expense_ids = [uuid4(), uuid4()]
    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        connection.exec_driver_sql(
            "CREATE TABLE documentos (id TEXT, tipo TEXT, aprobado_en DATETIME, "
            "cuenta_gastos_id TEXT)"
        )
        connection.exec_driver_sql(
            "CREATE TABLE expense_reports (id TEXT, estado_gasto TEXT, fecha DATETIME, "
            "documento_id TEXT, informe_documento_id TEXT, cuenta_gastos_id TEXT, "
            "pagado_con_amex_empresa BOOLEAN, origen TEXT)"
        )
        connection.exec_driver_sql(
            "CREATE TABLE amex_accounting_cuts (id TEXT, informe_id TEXT, kind TEXT, "
            "accounting_date DATE)"
        )
        connection.exec_driver_sql(
            "INSERT INTO documentos VALUES (?, 'INFORME', ?, ?)",
            (report_id.hex, "2026-09-20" if amex else "2026-10-20", account_id.hex),
        )
        for expense_id, expense_date in zip(expense_ids, ("2026-08-31", "2026-09-01")):
            connection.exec_driver_sql(
                "INSERT INTO expense_reports VALUES (?, 'activo', ?, ?, ?, ?, ?, ?)",
                (
                    expense_id.hex,
                    expense_date,
                    report_id.hex if relation == "documento" else None,
                    report_id.hex if relation == "informe" else None,
                    account_id.hex if relation == "cuenta" else None,
                    amex,
                    "amex_batch" if amex else None,
                ),
            )
        if amex:
            connection.exec_driver_sql(
                "INSERT INTO amex_accounting_cuts VALUES (?, ?, 'initial', '2026-10-15')",
                (uuid4().hex, report_id.hex),
            )
        for month in (8, 9, 10, 11):
            start, end = datetime(2026, month, 1), datetime(2026, month + 1, 1)
            reports = (
                connection.execute(
                    select(Documento.id).where(informe_coi_period_condition(start, end))
                )
                .scalars()
                .all()
            )
            expenses = (
                connection.execute(
                    select(ExpenseReport.id).where(
                        expense_coi_batch_period_condition(start, end)
                    )
                )
                .scalars()
                .all()
            )
            assert reports == ([report_id] if month == 10 else [])
            assert set(expenses) == (set(expense_ids) if month == 10 else set())
    engine.dispose()


def test_coi_exportable_rows_include_selection_and_status_controls():
    expense_id = uuid4()
    documento_id = uuid4()
    html = user_routes._render_coi_exportable_lote_rows_html(
        [
            {
                "tipo_lote": "INFORME",
                "documento": SimpleNamespace(
                    id=documento_id,
                    numero_referencia="I-123456",
                    estado="aprobado",
                ),
                "expenses": [
                    SimpleNamespace(
                        id=expense_id,
                        numero_referencia="O-26000001",
                        concepto="Gasolina",
                        gasto_cantidad=123.45,
                        fecha=datetime(2026, 8, 20),
                        coi_estado="reversar",
                    )
                ],
                "period_label": "2026-08-20",
                "can_export": True,
                "block_reason": "",
            }
        ]
    )

    assert 'name="selected_documento_id"' in html
    assert f'value="{documento_id}"' in html
    assert (
        f'class="coi-selection-checkbox" type="checkbox" '
        f'form="coi-export-form" name="selected_documento_id" '
        f'value="{documento_id}"'
    ) in html
    assert f'value="{documento_id}" checked' not in html
    assert 'name="coi_estado"' in html
    assert 'value="reversar" selected' in html
    assert f"/admin/contabilidad/coi/gastos/{expense_id}/estado" in html
    assert f"/documentos/{documento_id}/exportar-coi.xlsx" in html
    assert "Clasificación: No registrado" in html
    assert "Exportación: No registrado" in html


def test_coi_blocked_report_has_no_selectable_partial_policy():
    documento_id = uuid4()
    html = user_routes._render_coi_exportable_lote_rows_html(
        [
            {
                "tipo_lote": "INFORME",
                "documento": SimpleNamespace(
                    id=documento_id,
                    numero_referencia="I-123456",
                    estado="aprobado",
                ),
                "expenses": [],
                "period_label": "2026-08-20",
                "can_export": False,
                "block_reason": "O-1: falta cuenta contable",
            }
        ]
    )

    assert 'name="selected_documento_id"' not in html
    assert "Bloqueada: O-1: falta cuenta contable" in html


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("role", "owns_account", "can_download"),
    [
        ("finanzas", False, True),
        ("operaciones", True, True),
        ("operaciones", False, False),
    ],
)
async def test_blocked_amex_report_links_only_readable_review_workpaper(
    role, owns_account, can_download
):
    cuenta_id, documento_id, expense_id = uuid4(), uuid4(), uuid4()
    rows = [
        {
            "tipo_lote": "INFORME",
            "documento": SimpleNamespace(
                id=documento_id,
                cuenta_gastos_id=cuenta_id,
                numero_referencia="I-AMEX",
                estado="aprobado",
            ),
            "expenses": [
                SimpleNamespace(
                    id=expense_id,
                    numero_referencia="O-AMEX",
                    concepto="Cargo AMEX",
                    gasto_cantidad=100,
                    fecha=datetime(2026, 8, 20),
                    pagado_con_amex_empresa=True,
                )
            ],
            "period_label": "2026-08",
            "can_export": False,
            "block_reason": "Falta corte AMEX",
        }
    ]
    session = SimpleNamespace(
        execute=AsyncMock(
            return_value=SimpleNamespace(
                scalars=lambda: SimpleNamespace(
                    all=lambda: [cuenta_id] if owns_account else []
                )
            )
        )
    )
    allowed_ids = await user_routes._coi_readable_workpaper_account_ids(
        session,
        rows,
        SimpleNamespace(
            id=uuid4(),
            rol=role,
            correo="",
            permissions={"finanzas.manage"} if role == "operaciones" else set(),
        ),
    )
    html = user_routes._render_coi_exportable_lote_rows_html(
        rows, readable_workpaper_account_ids=allowed_ids
    )

    assert "Bloqueada: Falta corte AMEX" in html
    assert (f'/informes-de-gastos/{cuenta_id}/papel-poliza.xlsx' in html) is can_download
    assert ("Descargar papel de revisión" in html) is can_download
    assert 'name="selected_documento_id"' not in html


def test_coi_status_form_keeps_month_filters_and_expense_anchor():
    expense_id = uuid4()
    return_to = (
        "/admin/contabilidad/coi?year=2026&month=8&tipo=Eg"
        "&q=tarjeta+%26+taxis"
    )
    html = unescape(
        user_routes._render_coi_expense_status_form(
            SimpleNamespace(
                id=expense_id,
                numero_referencia="O-AMEX",
                coi_estado="pendiente",
            ),
            {},
            return_to=return_to,
        )
    )

    assert f'id="coi-expense-{expense_id}"' in html
    assert f'name="next" value="{return_to}"' in html


@pytest.mark.asyncio
async def test_coi_status_save_returns_to_filtered_expense():
    expense_id = uuid4()
    expense = SimpleNamespace(id=expense_id, coi_exported_at=None)
    session = SimpleNamespace(
        execute=AsyncMock(return_value=_ScalarRows([expense])),
        commit=AsyncMock(),
    )

    response = await user_routes.actualizar_estado_coi_gasto(
        expense_id,
        request=SimpleNamespace(),
        session=session,
        current_empleado=SimpleNamespace(id=uuid4()),
        coi_estado="reversar",
        next=(
            "/admin/contabilidad/coi?year=2026&month=8&tipo=Eg"
            "&q=tarjeta+%26+taxis"
        ),
    )

    location = response.headers["location"]
    assert response.status_code == 303
    assert location.startswith(
        "/admin/contabilidad/coi?year=2026&month=8&tipo=Eg"
        "&q=tarjeta+%26+taxis&success_msg="
    )
    assert location.endswith(f"#coi-expense-{expense_id}")
    assert expense.coi_estado == "reversar"
    session.commit.assert_awaited_once()


def test_coi_status_return_path_rejects_external_or_other_admin_pages():
    for raw in ("//example.org/path", "/admin/contabilidad/manual"):
        assert user_routes._coi_status_return_path(raw) == "/admin/contabilidad/coi"


def test_coi_export_page_requires_explicit_visible_selection():
    from pathlib import Path

    source = Path("src/devnous/gastos/routes/user_routes.py").read_text()

    assert '@router.post("/admin/contabilidad/coi/exportar-gastos-lote.xlsx"' in source
    assert 'id="coi-select-all"' in source
    assert 'id="coi-selected-count"' in source
    assert 'id="coi-confirmed-selection-count"' in source
    assert "selected_documento_id" in source
    assert "Exportar todo" not in source
    assert "Confirma nuevamente el número exacto de pólizas" in source
    assert 'coi_estado = "contabilizado"' in source


def test_coi_view_separates_preparation_export_and_history_tasks():
    from pathlib import Path

    source = Path("src/devnous/gastos/routes/user_routes.py").read_text()
    start = source.index("async def contabilidad_coi_view")
    end = source.index(
        '@router.post("/admin/contabilidad/coi/gastos/{expense_id}/estado")',
        start,
    )
    view_source = source[start:end]

    assert '<nav class="task-journey" aria-label="Flujo de trabajo COI">' in view_source
    assert (
        'href="/admin/gastos/sin-cuenta-contable?period={selected_year}-{selected_month:02d}"'
        in view_source
    )
    assert "1. Preparar COI" in view_source
    assert "2. Revisar y exportar" in view_source
    assert "3. Consultar historial" in view_source
    assert 'id="coi-exportacion"' in view_source
    assert (
        '<section id="coi-historial" aria-labelledby="coi-historial-heading">'
        in view_source
    )
    assert "2. Revisar y exportar gastos listos" in view_source
    assert "3. Historial e imports" in view_source


def test_coi_lote_preserves_standalone_expense_date_filters():
    from pathlib import Path

    source = Path("src/devnous/gastos/routes/user_routes.py").read_text()
    start = source.index("async def _load_coi_lote_informe_documentos")
    end = source.index("def _coi_lote_documento_period_label", start)
    filter_source = source[start:end]

    assert "ExpenseReport.fecha >= start_dt" in filter_source
    assert "ExpenseReport.fecha < end_dt" in filter_source
    assert "informe_coi_period_condition(start_dt, end_dt)" in filter_source
    assert "Documento.pagado_en >= start_dt" not in filter_source


@pytest.mark.asyncio
async def test_coi_batch_export_rejects_empty_or_hidden_selection(monkeypatch):
    visible_document = SimpleNamespace(id=uuid4())

    async def build_visible_rows(*_args, **_kwargs):
        return [{"documento": visible_document, "can_export": True}]

    monkeypatch.setattr(
        user_routes, "_build_coi_exportable_lote_rows", build_visible_rows
    )
    common = {
        "session": SimpleNamespace(),
        "current_empleado": SimpleNamespace(id=uuid4()),
        "year": 2026,
        "month": 9,
        "q": "",
    }

    empty = await user_routes.exportar_coi_gastos_lote_xlsx(
        **common, selected_documento_id=None, confirmed_selection_count=0
    )
    hidden = await user_routes.exportar_coi_gastos_lote_xlsx(
        **common, selected_documento_id=[uuid4()], confirmed_selection_count=1
    )

    assert empty.status_code == 303
    assert "Selecciona%20al%20menos%20una" in empty.headers["location"]
    assert hidden.status_code == 303
    assert "selecci%C3%B3n%20ya%20no%20coincide" in hidden.headers["location"]


@pytest.mark.asyncio
async def test_coi_batch_export_commits_only_the_confirmed_visible_selection(
    monkeypatch,
):
    visible_expense = SimpleNamespace(id=uuid4())
    visible_document = SimpleNamespace(id=uuid4())
    hidden_document = SimpleNamespace(id=uuid4())
    commits = []

    class Session:
        async def commit(self):
            commits.append(True)

    async def build_visible_rows(*_args, **_kwargs):
        return [
            {
                "documento": visible_document,
                "expenses": [visible_expense],
                "can_export": True,
            }
        ]

    async def collect_selected_rows(_session, rows):
        assert [row["documento"].id for row in rows] == [visible_document.id]
        return [], [], 1, {visible_expense.id}

    monkeypatch.setattr(
        user_routes, "_build_coi_exportable_lote_rows", build_visible_rows
    )
    monkeypatch.setattr(
        user_routes, "_collect_coi_lote_expense_cfdis", collect_selected_rows
    )
    monkeypatch.setattr(
        user_routes,
        "generate_coi_poliza_xlsx",
        lambda *_args, **_kwargs: b"xlsx",
    )

    response = await user_routes.exportar_coi_gastos_lote_xlsx(
        session=Session(),
        current_empleado=SimpleNamespace(id=uuid4()),
        year=2026,
        month=9,
        q="",
        selected_documento_id=[visible_document.id],
        confirmed_selection_count=1,
    )

    assert hidden_document.id != visible_document.id
    assert response.status_code == 200
    assert commits == [True]


@pytest.mark.asyncio
async def test_informe_lote_query_uses_approval_or_frozen_cut_period():
    class EmptyResult:
        def scalars(self):
            return self

        def all(self):
            return []

    class CaptureSession:
        statement = None

        async def execute(self, statement):
            self.statement = statement
            return EmptyResult()

    session = CaptureSession()
    await user_routes._load_coi_lote_informe_documentos(
        session, datetime(2026, 9, 1), datetime(2026, 10, 1)
    )
    compiled = str(
        session.statement.compile(
            dialect=postgresql.dialect(), compile_kwargs={"literal_binds": True}
        )
    )

    assert "EXISTS" in compiled
    assert "documentos.aprobado_en >= '2026-09-01 00:00:00'" in compiled
    assert "documentos.aprobado_en < '2026-10-01 00:00:00'" in compiled
    assert "amex_accounting_cuts.accounting_date >= '2026-09-01'" in compiled
    assert "amex_accounting_cuts.accounting_date < '2026-10-01'" in compiled


@pytest.mark.asyncio
async def test_document_bundle_blocks_entire_informe_when_one_expense_is_not_ready(
    monkeypatch,
):
    employee_id = uuid4()
    document = SimpleNamespace(
        id=uuid4(),
        empleado_id=employee_id,
        tipo="INFORME",
        estado="aprobado",
        cuenta_gastos_id=None,
        numero_referencia="I-26000001",
        aprobado_en=datetime(2026, 9, 20),
    )
    expense = SimpleNamespace(
        id=uuid4(),
        numero_referencia="G-26000001",
        estado_gasto="activo",
        fecha=datetime(2026, 9, 10),
    )

    class Session:
        calls = 0

        async def execute(self, _statement):
            self.calls += 1
            return _ScalarRows([document] if self.calls == 1 else [expense])

    async def not_ready(_session, _expense):
        return False, ["Falta cuenta contable"]

    monkeypatch.setattr(user_routes, "assess_expense_coi_cleanup_ready", not_ready)

    with pytest.raises(HTTPException) as exc_info:
        await user_routes._build_documento_coi_bundle(
            document.id,
            Session(),
            SimpleNamespace(id=employee_id, rol="usuario"),
            require_complete_informe=True,
        )

    assert exc_info.value.status_code == 400
    assert "No se generó una póliza parcial" in exc_info.value.detail
    assert "G-26000001: Falta cuenta contable" in exc_info.value.detail


@pytest.mark.asyncio
@pytest.mark.parametrize("second_date", [datetime(2026, 9, 1), datetime(2027, 1, 1)])
@pytest.mark.parametrize("has_approval_date", [False, True])
async def test_document_bundle_exports_informe_spanning_accounting_months(
    monkeypatch, second_date, has_approval_date
):
    employee_id = uuid4()
    document = SimpleNamespace(
        id=uuid4(),
        empleado_id=employee_id,
        tipo="INFORME",
        estado="aprobado",
        cuenta_gastos_id=None,
        numero_referencia="I-26000001",
        aprobado_en=datetime(2026, 9, 20) if has_approval_date else None,
    )
    expenses = [
        SimpleNamespace(
            id=uuid4(),
            numero_referencia=reference,
            fecha=expense_date,
        )
        for reference, expense_date in (
            ("G-AUG", datetime(2026, 8, 31)),
            ("G-LATER", second_date),
        )
    ]

    class Session:
        calls = 0

        async def execute(self, _statement):
            self.calls += 1
            return _ScalarRows([document] if self.calls == 1 else expenses)

    async def ready(*_args, **_kwargs):
        return True, []

    async def build(_session, expense, **_kwargs):
        return ExpenseCFDI(
            fecha=expense.fecha,
            total=100,
            iva_amount=0,
            subtotal_amount=100,
            concepto=expense.numero_referencia,
            cuenta_contable="5000",
            cuenta_contrapartida="1170",
            cfdi_uuid=str(expense.id),
            cfdi_date=expense.fecha,
        )

    monkeypatch.setattr(user_routes, "assess_expense_coi_cleanup_ready", ready)
    monkeypatch.setattr(user_routes, "build_expense_cfdi_for_export", build)
    if not has_approval_date:
        with pytest.raises(HTTPException, match="Falta fecha de aprobación"):
            await user_routes._build_documento_coi_bundle(
                document.id,
                Session(),
                SimpleNamespace(id=employee_id, rol="usuario"),
                require_complete_informe=True,
            )
        return
    _, loaded, payloads = await user_routes._build_documento_coi_bundle(
        document.id,
        Session(),
        SimpleNamespace(id=employee_id, rol="usuario"),
        require_complete_informe=True,
    )

    assert loaded == expenses
    assert [p.fecha for p in payloads] == [e.fecha for e in expenses]
    assert [p.cfdi_date for p in payloads] == [e.fecha for e in expenses]
    assert [p.cfdi_uuid for p in payloads] == [str(e.id) for e in expenses]
    assert {p.poliza_group_key for p in payloads} == {f"informe:{document.id}"}

    csv_rows = list(
        csv.reader(io.StringIO(generate_coi_poliza_csv(payloads).decode("utf-8-sig")))
    )
    workbook = load_workbook(io.BytesIO(generate_coi_poliza_xlsx(payloads)))
    with zipfile.ZipFile(io.BytesIO(generate_coi_poliza_zip(payloads))) as archive:
        assert archive.namelist() == [f"Poliza_COI_{document.numero_referencia}.xlsx"]
        zipped_workbook = load_workbook(io.BytesIO(archive.read(archive.namelist()[0])))
    for rows in (
        csv_rows,
        list(workbook["Poliza COI"].values),
        list(zipped_workbook["Poliza COI"].values),
    ):
        assert sum(row[0] == "Eg" for row in rows) == 1
        assert sum(row[1] == "FIN_PARTIDAS" for row in rows) == 1
        assert {row[8] for row in rows if row[8]} == {str(e.id) for e in expenses}
        assert {row[2] for row in rows if row[8]} == {
            e.fecha.strftime("%d/%m/%y") for e in expenses
        }
        movements = [row for row in rows if str(row[4]) == "1"]
        assert sum(float(row[5] or 0) for row in movements) == 200
        assert sum(float(row[6] or 0) for row in movements) == 200


@pytest.mark.asyncio
async def test_single_expense_export_redirects_to_owning_informe(monkeypatch):
    employee_id = uuid4()
    informe_id = uuid4()
    expense = SimpleNamespace(
        id=uuid4(),
        empleado_id=employee_id,
        informe_documento_id=informe_id,
        documento_id=None,
        cuenta_gastos_id=None,
    )
    informe = SimpleNamespace(id=informe_id, tipo="INFORME")

    async def load_expense(_session, _expense_id):
        return expense

    class Session:
        async def get(self, _model, entity_id):
            assert entity_id == informe_id
            return informe

    monkeypatch.setattr(user_routes, "load_expense_for_coi_export", load_expense)

    response = await user_routes.exportar_coi_poliza_gasto_excel(
        expense.id,
        Session(),
        SimpleNamespace(id=employee_id, rol="usuario"),
    )

    assert response.status_code == 303
    assert response.headers["location"] == f"/documentos/{informe_id}/exportar-coi.xlsx"


@pytest.mark.asyncio
@pytest.mark.parametrize("outside_ready", [True, False])
@pytest.mark.parametrize("has_approval_date", [True, False])
async def test_monthly_batch_checks_all_informe_expenses_across_periods(
    monkeypatch,
    outside_ready,
    has_approval_date,
):
    document = SimpleNamespace(
        id=uuid4(),
        tipo="INFORME",
        estado="aprobado",
        numero_referencia="I-26000001",
        aprobado_en=datetime(2026, 9, 20) if has_approval_date else None,
        creado_en=datetime(2026, 9, 1),
    )
    september = SimpleNamespace(
        id=uuid4(),
        numero_referencia="G-SEP",
        concepto="Gasolina",
        proyecto="Proyecto",
        fecha=datetime(2026, 9, 10),
    )
    august = SimpleNamespace(
        id=uuid4(),
        numero_referencia="G-AUG",
        concepto="Casetas",
        proyecto="Proyecto",
        fecha=datetime(2026, 8, 31),
    )

    async def load_informes(*_args, **_kwargs):
        return [document]

    async def load_terceros(*_args, **_kwargs):
        return []

    async def load_expenses(_session, _document, start_dt=None, end_dt=None):
        if start_dt is not None and end_dt is not None:
            return [september]
        return [august, september]

    async def ready(_session, expense):
        if expense.id == august.id and not outside_ready:
            return False, ["Falta cuenta contable"]
        return True, []

    monkeypatch.setattr(user_routes, "_load_coi_lote_informe_documentos", load_informes)
    monkeypatch.setattr(
        user_routes, "_load_coi_lote_terceros_documentos", load_terceros
    )
    monkeypatch.setattr(
        user_routes, "_load_documento_active_coi_expenses", load_expenses
    )
    monkeypatch.setattr(user_routes, "assess_expense_coi_cleanup_ready", ready)

    rows = await user_routes._build_coi_exportable_lote_rows(
        SimpleNamespace(),
        start_dt=datetime(2026, 9, 1),
        end_dt=datetime(2026, 10, 1),
        start_date=datetime(2026, 9, 1).date(),
        end_date=datetime(2026, 10, 1).date(),
    )

    assert len(rows) == 1
    assert rows[0]["expenses"] == [august, september]
    assert rows[0]["can_export"] is (outside_ready and has_approval_date)
    if not has_approval_date:
        assert "Falta fecha de aprobación" in rows[0]["block_reason"]
        assert rows[0]["period_label"] == "-"
    elif outside_ready:
        assert rows[0]["block_reason"] == ""
    else:
        assert "G-AUG: Falta cuenta contable" in rows[0]["block_reason"]


@pytest.mark.asyncio
@pytest.mark.parametrize("has_approval_date", [True, False])
async def test_finance_batch_groups_all_report_expenses_into_one_policy(
    monkeypatch, has_approval_date
):
    informe = SimpleNamespace(
        id=uuid4(),
        tipo="INFORME",
        estado="aprobado",
        numero_referencia="I-26000001",
        aprobado_en=datetime(2026, 9, 20) if has_approval_date else None,
        cuenta_gastos_id=None,
    )
    expenses = [
        SimpleNamespace(
            id=uuid4(),
            fecha=datetime(2026, 9, day),
            numero_referencia=f"G-{day}",
            cuenta_gastos_id=None,
            informe_documento=informe,
            documento=None,
        )
        for day in (10, 11)
    ]

    class Session:
        calls = 0

        async def execute(self, _statement):
            self.calls += 1
            return _ScalarRows(expenses)

    async def ready(*_args, **_kwargs):
        return True, []

    async def build(_session, expense, **_kwargs):
        return ExpenseCFDI(
            fecha=expense.fecha,
            total=100,
            iva_amount=0,
            subtotal_amount=100,
            concepto=expense.numero_referencia,
            cuenta_contable="5000-001",
            cuenta_contrapartida="1170-001",
        )

    monkeypatch.setattr(admin_routes, "assess_expense_coi_cleanup_ready", ready)
    monkeypatch.setattr(admin_routes, "build_expense_cfdi_for_export", build)

    if not has_approval_date:
        with pytest.raises(ValueError, match="Falta fecha de aprobación"):
            await admin_routes._build_finance_coi_batch_expenses(
                Session(), year=2026, month=9
            )
        return

    year, month, payloads = await admin_routes._build_finance_coi_batch_expenses(
        Session(), year=2026, month=9
    )

    assert (year, month) == (2026, 9)
    assert len(payloads) == 2
    assert {item.poliza_group_key for item in payloads} == {f"informe:{informe.id}"}
    assert {item.poliza_reference for item in payloads} == {"I-26000001"}


@pytest.mark.asyncio
async def test_finance_batch_blocks_unapproved_informe():
    informe = SimpleNamespace(
        id=uuid4(),
        tipo="INFORME",
        estado="enviado",
        numero_referencia="I-26000002",
        cuenta_gastos_id=None,
    )
    expense = SimpleNamespace(
        id=uuid4(),
        fecha=datetime(2026, 9, 10),
        numero_referencia="G-10",
        cuenta_gastos_id=None,
        informe_documento=informe,
        documento=None,
    )

    class Session:
        async def execute(self, _statement):
            return _ScalarRows([expense])

    with pytest.raises(ValueError, match="debe estar aprobado"):
        await admin_routes._build_finance_coi_batch_expenses(
            Session(), year=2026, month=9
        )


@pytest.mark.asyncio
@pytest.mark.parametrize("cut_state", ["missing", "invalid", "frozen"])
async def test_finance_amex_report_requires_frozen_cut(monkeypatch, cut_state):
    informe = SimpleNamespace(
        id=uuid4(),
        tipo="INFORME",
        estado="aprobado",
        numero_referencia="I-AMEX",
        cuenta_gastos_id=None,
    )
    expenses = [
        SimpleNamespace(
            id=uuid4(),
            fecha=expense_date,
            cuenta_gastos_id=None,
            informe_documento=informe,
            documento=None,
        )
        for expense_date in (datetime(2026, 8, 31), datetime(2026, 9, 1))
    ]
    cut = SimpleNamespace(id=uuid4())
    responses = [expenses, expenses, [] if cut_state == "missing" else [cut]]

    class Session:
        async def execute(self, _statement):
            return _ScalarRows(responses.pop(0))

    def frozen(found_cut):
        assert found_cut is cut
        if cut_state == "invalid":
            raise ValueError("El corte contiene evidencia incompleta o inválida.")
        return [payload]

    def forbidden(*_args, **_kwargs):
        pytest.fail("AMEX export must not rebuild mutable classifications")

    payload = ExpenseCFDI(
        fecha=datetime(2026, 9, 30),
        total=200,
        iva_amount=0,
        subtotal_amount=200,
        concepto="Corte congelado",
        cuenta_contable="5000",
        cuenta_contrapartida="1170",
        poliza_group_key=f"amex-cut:{cut.id}",
    )
    monkeypatch.setattr(admin_routes, "is_company_amex_expense", lambda _: True)
    monkeypatch.setattr(admin_routes, "cut_expense_cfdis", frozen)
    monkeypatch.setattr(admin_routes, "assess_expense_coi_cleanup_ready", forbidden)
    monkeypatch.setattr(admin_routes, "build_expense_cfdi_for_export", forbidden)
    if cut_state != "frozen":
        with pytest.raises(ValueError, match="corte"):
            await admin_routes._build_finance_coi_batch_expenses(
                Session(), year=2026, month=9
            )
    else:
        _, _, payloads = await admin_routes._build_finance_coi_batch_expenses(
            Session(), year=2026, month=9
        )
        assert payloads == [payload]


@pytest.mark.asyncio
@pytest.mark.parametrize("relation", ["documento", "cuenta"])
@pytest.mark.parametrize(
    "blocker", [None, "not_ready", "outside_period", "outside_not_ready"]
)
async def test_finance_report_ownership_and_atomic_blockers(
    monkeypatch, relation, blocker
):
    account_id = uuid4() if relation == "cuenta" else None
    informe = SimpleNamespace(
        id=uuid4(),
        tipo="INFORME",
        estado="aprobado",
        numero_referencia="I-ATOMIC",
        aprobado_en=datetime(2026, 9, 20),
        cuenta_gastos_id=account_id,
    )
    expense = SimpleNamespace(
        id=uuid4(),
        fecha=datetime(2026, 9, 10),
        numero_referencia="G-READY",
        cuenta_gastos_id=account_id,
        informe_documento=None,
        documento=informe if relation == "documento" else None,
    )
    second = SimpleNamespace(
        **{**vars(expense), "id": uuid4(), "numero_referencia": "G-BLOCKED"}
    )
    if blocker in {"outside_period", "outside_not_ready"}:
        second.fecha = datetime(2026, 8, 31)
    responses = [[expense]]
    if account_id:
        responses.append([informe])
    responses.append([expense, second])

    class Session:
        async def execute(self, _statement):
            return _ScalarRows(responses.pop(0))

    async def ready(_session, item):
        if blocker in {"not_ready", "outside_not_ready"} and item.id == second.id:
            return False, ["Falta contrapartida"]
        return True, []

    async def build(_session, item, **_kwargs):
        return ExpenseCFDI(
            fecha=item.fecha,
            total=100,
            iva_amount=0,
            subtotal_amount=100,
            concepto=item.numero_referencia,
            cuenta_contable="5000",
            cuenta_contrapartida="1170",
        )

    monkeypatch.setattr(admin_routes, "assess_expense_coi_cleanup_ready", ready)
    monkeypatch.setattr(admin_routes, "build_expense_cfdi_for_export", build)
    if blocker in {"not_ready", "outside_not_ready"}:
        with pytest.raises(ValueError, match="G-BLOCKED"):
            await admin_routes._build_finance_coi_batch_expenses(
                Session(), year=2026, month=9
            )
    else:
        _, _, payloads = await admin_routes._build_finance_coi_batch_expenses(
            Session(), year=2026, month=9
        )
        assert len(payloads) == 2
        assert {p.poliza_group_key for p in payloads} == {f"informe:{informe.id}"}
        assert [p.fecha for p in payloads] == [expense.fecha, second.fecha]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "endpoint",
    [
        admin_routes.admin_finance_coi_batch_consolidated_xlsx,
        admin_routes.admin_finance_coi_batch_xlsx,
    ],
)
async def test_finance_download_returns_blocker_without_generating_file(
    monkeypatch, endpoint
):
    async def blocked(*_args, **_kwargs):
        raise ValueError("Informe incompleto: G-2")

    def forbidden(*_args, **_kwargs):
        pytest.fail("A blocked report must not generate a partial file")

    monkeypatch.setattr(admin_routes, "_build_finance_coi_batch_expenses", blocked)
    monkeypatch.setattr(admin_routes, "generate_coi_poliza_xlsx", forbidden)
    monkeypatch.setattr(admin_routes, "generate_coi_poliza_zip", forbidden)
    response = await endpoint(
        request=SimpleNamespace(),
        session=SimpleNamespace(),
        current_empleado=SimpleNamespace(),
        year=2026,
        month=9,
    )
    assert response.status_code == 303
    assert "Informe%20incompleto%3A%20G-2" in response.headers["location"]


@pytest.mark.asyncio
@pytest.mark.parametrize("fail_second", [False, True])
async def test_selected_report_exports_all_items_or_no_status_updates(
    monkeypatch, fail_second
):
    document = SimpleNamespace(id=uuid4(), tipo="INFORME", numero_referencia="I-ATOMIC")
    expenses = [
        SimpleNamespace(
            id=uuid4(), numero_referencia=f"G-{index}", coi_estado="pendiente"
        )
        for index in (1, 2)
    ]
    commits = []

    class Session:
        async def commit(self):
            commits.append(True)

    async def visible(*_args, **_kwargs):
        return [
            {
                "documento": document,
                "tipo_lote": "INFORME",
                "expenses": expenses,
                "can_export": True,
            }
        ]

    async def build(_session, item, **_kwargs):
        if fail_second and item.id == expenses[1].id:
            raise ValueError("CFDI incompleto")
        return ExpenseCFDI(
            fecha=datetime(2026, 9, 10),
            total=100,
            iva_amount=0,
            subtotal_amount=100,
            concepto=item.numero_referencia,
            cuenta_contable="5000",
            cuenta_contrapartida="1170",
        )

    def generate(payloads, **_kwargs):
        assert len(payloads) == 2
        assert {p.poliza_group_key for p in payloads} == {f"informe:{document.id}"}
        return b"complete-report"

    monkeypatch.setattr(user_routes, "_build_coi_exportable_lote_rows", visible)
    monkeypatch.setattr(user_routes, "build_expense_cfdi_for_export", build)
    monkeypatch.setattr(user_routes, "generate_coi_poliza_xlsx", generate)
    actor = SimpleNamespace(id=uuid4())
    response = await user_routes.exportar_coi_gastos_lote_xlsx(
        session=Session(),
        current_empleado=actor,
        year=2026,
        month=9,
        q="",
        selected_documento_id=[document.id],
        confirmed_selection_count=1,
    )
    if fail_second:
        assert response.status_code == 303
        assert "CFDI%20incompleto" in response.headers["location"]
        assert commits == []
        assert {e.coi_estado for e in expenses} == {"pendiente"}
    else:
        assert response.body == b"complete-report"
        assert commits == [True]
        assert {e.coi_estado for e in expenses} == {"contabilizado"}
        assert {e.coi_exported_by_id for e in expenses} == {actor.id}


@pytest.mark.asyncio
@pytest.mark.parametrize("relation", ["documento", "cuenta"])
async def test_single_expense_resolves_legacy_report_links(monkeypatch, relation):
    actor_id = uuid4()
    informe = SimpleNamespace(id=uuid4(), tipo="INFORME")
    expense = SimpleNamespace(
        id=uuid4(),
        empleado_id=actor_id,
        informe_documento_id=None,
        documento_id=informe.id if relation == "documento" else None,
        cuenta_gastos_id=uuid4() if relation == "cuenta" else None,
    )

    async def load(*_args):
        return expense

    async def load_report(*_args):
        return informe

    class Session:
        async def get(self, _model, _entity_id):
            return informe

    monkeypatch.setattr(user_routes, "load_expense_for_coi_export", load)
    monkeypatch.setattr(user_routes, "_informe_documento_for_cuenta", load_report)
    response = await user_routes.exportar_coi_poliza_gasto_excel(
        expense.id, Session(), SimpleNamespace(id=actor_id, rol="usuario")
    )
    assert response.status_code == 303
    assert response.headers["location"] == f"/documentos/{informe.id}/exportar-coi.xlsx"


@pytest.mark.asyncio
async def test_document_bundle_rejects_late_failure_after_ready_item(monkeypatch):
    actor_id = uuid4()
    document = SimpleNamespace(
        id=uuid4(),
        empleado_id=actor_id,
        tipo="INFORME",
        estado="aprobado",
        cuenta_gastos_id=None,
        numero_referencia="I-ATOMIC",
        aprobado_en=datetime(2026, 9, 20),
    )
    expenses = [
        SimpleNamespace(
            id=uuid4(),
            numero_referencia=f"G-{i}",
            fecha=datetime(2026, 9, 10),
            coi_estado="pendiente",
        )
        for i in (1, 2)
    ]
    responses = [[document], expenses]

    class Session:
        async def execute(self, _statement):
            return _ScalarRows(responses.pop(0))

        async def commit(self):
            pytest.fail("An incomplete report must not commit status changes")

    async def ready(*_args):
        return True, []

    async def build(_session, item, **_kwargs):
        if item.id == expenses[1].id:
            raise ValueError("Falta contrapartida")
        return ExpenseCFDI(
            fecha=item.fecha,
            total=100,
            iva_amount=0,
            subtotal_amount=100,
            concepto=item.numero_referencia,
            cuenta_contable="5000",
            cuenta_contrapartida="1170",
        )

    monkeypatch.setattr(user_routes, "assess_expense_coi_cleanup_ready", ready)
    monkeypatch.setattr(user_routes, "build_expense_cfdi_for_export", build)
    with pytest.raises(HTTPException, match="G-2: Falta contrapartida"):
        await user_routes._build_documento_coi_bundle(
            document.id,
            Session(),
            SimpleNamespace(id=actor_id, rol="usuario"),
            require_complete_informe=True,
        )
    assert {e.coi_estado for e in expenses} == {"pendiente"}


@pytest.mark.asyncio
@pytest.mark.parametrize("other_ready", [False, True])
async def test_standalone_policies_remain_independently_selectable(
    monkeypatch, other_ready
):
    document = SimpleNamespace(
        id=uuid4(),
        tipo="SOLICITUD",
        estado="aprobado",
        numero_referencia="S-1",
        pagado_en=None,
        fecha_pago=None,
        aprobado_en=datetime(2026, 9, 10),
        creado_en=datetime(2026, 9, 1),
    )
    expenses = [
        SimpleNamespace(
            id=uuid4(),
            numero_referencia=f"G-{i}",
            concepto="Servicios",
            proyecto="Proyecto",
            fecha=datetime(2026, 9, 10),
            gasto_cantidad=100,
        )
        for i in range(2)
    ]

    async def informes(*args):
        return []

    async def terceros(*args):
        return [document]

    async def load(*args):
        return expenses

    async def ready(session, expense):
        ok = expense is expenses[0] or other_ready
        return ok, [] if ok else ["Falta cuenta"]

    monkeypatch.setattr(user_routes, "_load_coi_lote_informe_documentos", informes)
    monkeypatch.setattr(user_routes, "_load_coi_lote_terceros_documentos", terceros)
    monkeypatch.setattr(user_routes, "_load_documento_active_coi_expenses", load)
    monkeypatch.setattr(user_routes, "assess_expense_coi_cleanup_ready", ready)
    rows = await user_routes._build_coi_exportable_lote_rows(
        SimpleNamespace(),
        start_dt=datetime(2026, 9, 1),
        end_dt=datetime(2026, 10, 1),
        start_date=datetime(2026, 9, 1).date(),
        end_date=datetime(2026, 10, 1).date(),
    )
    assert len(rows) == (2 if other_ready else 1)
    assert rows[0]["expenses"] == [expenses[0]]
    assert all(row["can_export"] for row in rows)
    html = user_routes._render_coi_exportable_lote_rows_html(rows)
    assert f'name="selected_gasto_id" value="{expenses[0].id}"' in html
    assert f"/gastos/{expenses[0].id}/exportar-coi.xlsx" in html
    assert 'name="selected_documento_id"' not in html
    collected = []
    commits = []

    async def visible(*args, **kwargs):
        return rows

    async def collect(session, selected_rows):
        collected.extend(selected_rows)
        return [], [], len(selected_rows), {expenses[0].id}

    class Session:
        async def commit(self):
            commits.append(True)

    monkeypatch.setattr(user_routes, "_build_coi_exportable_lote_rows", visible)
    monkeypatch.setattr(user_routes, "_collect_coi_lote_expense_cfdis", collect)
    monkeypatch.setattr(
        user_routes, "generate_coi_poliza_xlsx", lambda *args, **kwargs: b"xlsx"
    )
    response = await user_routes.exportar_coi_gastos_lote_xlsx(
        session=Session(),
        current_empleado=SimpleNamespace(id=uuid4()),
        year=2026,
        month=9,
        q="",
        selected_documento_id=None,
        selected_gasto_id=[expenses[0].id],
        confirmed_selection_count=1,
    )
    assert response.status_code == 200
    assert collected == [rows[0]]
    assert commits == [True]
    assert expenses[0].coi_estado == "contabilizado"
    assert not hasattr(expenses[1], "coi_estado")


@pytest.mark.parametrize(
    "party_path",
    [
        "beneficiario_empleado",
        "beneficiario_proveedor_cliente",
        "account_employee",
        "account_provider",
        "account_legacy",
        "proveedor_cliente",
        "empleado",
    ],
)
def test_coi_search_and_display_use_effective_beneficiary(party_path):
    party = SimpleNamespace(nombre="José <Pérez>")
    document = SimpleNamespace(
        id=uuid4(),
        numero_referencia="I-42",
        estado="aprobado",
        empleado=SimpleNamespace(nombre="Solicitante diferente"),
    )
    account_paths = {
        "account_employee": "beneficiario_empleado",
        "account_provider": "beneficiario_proveedor_cliente",
        "account_legacy": "empleado",
    }
    if party_path in account_paths:
        document.cuenta_gastos = SimpleNamespace(**{account_paths[party_path]: party})
    else:
        setattr(document, party_path, party)
    assert user_routes._coi_exportable_matches_search(
        documento=document, expenses=[], search_q="JOSE <PE"
    )
    assert not user_routes._coi_exportable_matches_search(
        documento=document, expenses=[], search_q="Solicitante diferente"
    )
    html = user_routes._render_coi_exportable_lote_rows_html(
        [
            {
                "tipo_lote": "INFORME",
                "documento": document,
                "expenses": [],
                "period_label": "2026-09-01",
                "can_export": False,
            }
        ]
    )
    assert "Titular/beneficiario: José &lt;Pérez&gt;" in html
    assert "Solicitante diferente" not in html


def test_coi_search_missing_name_preserves_existing_fields():
    document = SimpleNamespace(id=uuid4(), numero_referencia="I-42", estado="aprobado")
    expense = SimpleNamespace(
        id=uuid4(),
        numero_referencia="G-12",
        concepto="Hospedaje",
        proyecto="Torneo",
    )
    for query in ("", " I-42 ", "aprobado", "hosped", "torneo", "G-12"):
        assert user_routes._coi_exportable_matches_search(
            documento=document, expenses=[expense], search_q=query
        )
    assert not user_routes._coi_exportable_matches_search(
        documento=document, expenses=[expense], search_q="persona ausente"
    )
    html = user_routes._render_coi_exportable_lote_rows_html(
        [
            {
                "tipo_lote": "INFORME",
                "documento": document,
                "expenses": [],
                "period_label": "2026-09-01",
                "can_export": False,
            }
        ]
    )
    assert "Titular/beneficiario: —" in html


def test_coi_explicit_beneficiary_precedes_account_and_requester():
    document = SimpleNamespace(
        id=uuid4(),
        numero_referencia="I-42",
        estado="aprobado",
        beneficiario_empleado=SimpleNamespace(nombre="Titular explícito"),
        cuenta_gastos=SimpleNamespace(
            beneficiario_empleado=SimpleNamespace(nombre="Titular cuenta")
        ),
        empleado=SimpleNamespace(nombre="Solicitante"),
    )
    for query, expected in (
        ("titular explicito", True),
        ("titular cuenta", False),
        ("solicitante", False),
    ):
        assert (
            user_routes._coi_exportable_matches_search(
                documento=document, expenses=[], search_q=query
            )
            is expected
        )


@pytest.mark.asyncio
@pytest.mark.parametrize("kind", ["INFORME", "SOLICITUD_TERCEROS"])
async def test_coi_loaders_eagerly_load_all_beneficiary_fallbacks(kind):
    class Session:
        statement = None

        async def execute(self, statement):
            self.statement = statement
            return _ScalarRows([])

    session = Session()
    start, end = datetime(2026, 9, 1), datetime(2026, 10, 1)
    if kind == "INFORME":
        await user_routes._load_coi_lote_informe_documentos(session, start, end)
    else:
        await user_routes._load_coi_lote_terceros_documentos(
            session, start, end, start.date(), end.date()
        )
    paths = {
        tuple(element.key for element in option.path if hasattr(element, "key"))
        for option in session.statement._with_options
    }
    assert paths == {
        ("beneficiario_empleado",),
        ("beneficiario_proveedor_cliente",),
        ("proveedor_cliente",),
        ("empleado",),
        ("cuenta_gastos", "beneficiario_empleado"),
        ("cuenta_gastos", "beneficiario_proveedor_cliente"),
        ("cuenta_gastos", "empleado"),
    }
    sql = str(
        session.statement.compile(
            dialect=postgresql.dialect(), compile_kwargs={"literal_binds": True}
        )
    )
    if kind == "INFORME":
        assert "documentos.estado = 'aprobado'" in sql
        assert "documentos.aprobado_en >= '2026-09-01 00:00:00'" in sql
    else:
        assert "expense_reports.fecha >= '2026-09-01 00:00:00'" in sql
        assert "expense_reports.fecha < '2026-10-01 00:00:00'" in sql


@pytest.mark.asyncio
async def test_coi_view_searches_beneficiary_and_preserves_period(monkeypatch):
    from starlette.requests import Request

    document = SimpleNamespace(
        id=uuid4(),
        numero_referencia="I-42",
        estado="aprobado",
        tipo="INFORME",
        cuenta_gastos_id=None,
        aprobado_en=datetime(2026, 9, 5),
        beneficiario_empleado=SimpleNamespace(nombre="José Pérez"),
        empleado=SimpleNamespace(nombre="Solicitante diferente"),
    )
    expense = SimpleNamespace(
        id=uuid4(),
        numero_referencia="G-1",
        concepto="Hospedaje",
        proyecto="Torneo",
        gasto_cantidad=100,
        fecha=datetime(2026, 9, 5),
    )
    bounds = []

    async def reports(_session, start, end):
        bounds.append((start, end))
        return [document]

    async def third_parties(*_args):
        return []

    async def expenses(*_args):
        return [expense]

    async def ready(*_args):
        return True, []

    class Session:
        async def execute(self, _statement):
            return _ScalarRows([])

    monkeypatch.setattr(user_routes, "_load_coi_lote_informe_documentos", reports)
    monkeypatch.setattr(
        user_routes, "_load_coi_lote_terceros_documentos", third_parties
    )
    monkeypatch.setattr(user_routes, "_load_documento_active_coi_expenses", expenses)
    monkeypatch.setattr(user_routes, "assess_expense_coi_cleanup_ready", ready)
    monkeypatch.setattr(user_routes, "render_top_navigation", lambda *_args: "")
    request = Request({"type": "http", "query_string": b""})
    for query, visible in (("jose pe", True), ("Solicitante diferente", False)):
        html = await user_routes.contabilidad_coi_view(
            request,
            Session(),
            SimpleNamespace(),
            year=2026,
            month=9,
            tipo="Eg",
            q=query,
        )
        assert ("Titular/beneficiario: José Pérez" in html) is visible
        assert ('name="selected_documento_id"' in html) is visible
        assert 'name="year" value="2026"' in html
        assert 'name="month" value="9"' in html
        assert f'name="q" value="{query}"' in html
    assert bounds == [(datetime(2026, 9, 1), datetime(2026, 10, 1))] * 2
