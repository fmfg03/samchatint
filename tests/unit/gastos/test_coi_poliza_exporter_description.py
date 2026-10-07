import io
import zipfile
from datetime import datetime
from types import SimpleNamespace

from openpyxl import load_workbook

from devnous.gastos.services.coi_poliza_exporter import (
    ExpenseCFDI,
    _expense_description,
    build_coi_poliza_rows,
    generate_coi_poliza_xlsx,
    generate_coi_poliza_zip,
)
from devnous.gastos.services.expense_coi_export_service import (
    group_expense_cfdis_for_document,
)


def _expense(**overrides) -> ExpenseCFDI:
    values = {
        "fecha": datetime(2026, 9, 6),
        "total": 40334.64,
        "iva_amount": 0.0,
        "subtotal_amount": 40334.64,
        "concepto": "Hospedaje Fase Nacional LTTB",
        "cuenta_contable": "5300-012-031",
        "cuenta_contrapartida": "1200-001-001",
    }
    values.update(overrides)
    return ExpenseCFDI(**values)


def test_coi_description_uses_accounting_order_and_removes_provider_duplication():
    expense = _expense(
        export_reference="03/269",
        nombre_emisor="ANTONIO GARCES HINOJOZA",
        concepto=(
            "Pago a proveedor: ANTONIO GARCES HINOJOZA - "
            "HOSPEDAJE FASE NACIONAL LTTB"
        ),
        proyecto="SEDE ECATEPEC",
    )

    assert _expense_description(expense) == (
        "03/269 / ANTONIO GARCES HINOJOZA / "
        "HOSPEDAJE FASE NACIONAL LTTB / SEDE ECATEPEC"
    )


def test_coi_description_does_not_emit_empty_fields_or_provider_wrapper_without_cfdi():
    expense = _expense(
        export_reference="03/270",
        concepto="Pago a proveedor: Servicio médico",
    )

    assert _expense_description(expense) == "03/270 / Servicio médico"


def test_coi_header_and_movement_rows_share_compact_description():
    expense = _expense(
        export_reference="03/269",
        nombre_emisor="ANTONIO GARCES HINOJOZA",
        concepto="Pago a proveedor: ANTONIO GARCES HINOJOZA - Hospedaje",
        proyecto="LTTB",
    )

    rows = build_coi_poliza_rows([expense])
    description = "03/269 / ANTONIO GARCES HINOJOZA / Hospedaje / LTTB"

    assert rows[2][2] == description
    assert rows[3][3] == description
    assert rows[4][3] == description


def test_expense_report_partidas_share_one_coi_policy_header_and_closure():
    first = _expense(
        export_reference="G-1",
        concepto="Hospedaje",
        poliza_group_key="informe:1",
        poliza_reference="I-26000001",
        poliza_description="Informe de Gastos I-26000001",
    )
    second = _expense(
        export_reference="G-2",
        concepto="Alimentos",
        poliza_group_key="informe:1",
        poliza_reference="I-26000001",
        poliza_description="Informe de Gastos I-26000001",
    )

    rows = build_coi_poliza_rows([first, second])

    headers = [row for row in rows if row[0] == "Eg"]
    closures = [row for row in rows if row[1] == "FIN_PARTIDAS"]
    assert headers == [
        [
            "Eg",
            "1",
            "Informe de Gastos I-26000001",
            "4",
            "",
            "",
            "",
            "",
            "",
        ]
    ]
    assert len(closures) == 1


def test_ungrouped_standalone_expenses_keep_one_policy_each():
    rows = build_coi_poliza_rows(
        [
            _expense(export_reference="G-1"),
            _expense(export_reference="G-2"),
        ]
    )

    assert [row[1] for row in rows if row[0] == "Eg"] == ["1", "2"]
    assert len([row for row in rows if row[1] == "FIN_PARTIDAS"]) == 2


def test_grouped_report_produces_one_workbook_in_zip():
    expenses = [
        _expense(
            export_reference=reference,
            poliza_group_key="informe:1",
            poliza_reference="I-26000001",
            poliza_description="Informe de Gastos I-26000001",
        )
        for reference in ("G-1", "G-2")
    ]

    payload = generate_coi_poliza_zip(expenses)

    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        assert archive.namelist() == ["Poliza_COI_I-26000001.xlsx"]


def test_grouped_report_preserves_each_cfdi_block_inside_single_policy():
    expenses = [
        _expense(
            export_reference=reference,
            cfdi_uuid=cfdi_uuid,
            rfc_emisor="AAA010101AAA",
            rfc_receptor="BBB010101BBB",
            poliza_group_key="informe:1",
            poliza_reference="I-26000001",
            poliza_description="Informe de Gastos I-26000001",
        )
        for reference, cfdi_uuid in (("G-1", "UUID-1"), ("G-2", "UUID-2"))
    ]

    rows = build_coi_poliza_rows(expenses)

    assert len([row for row in rows if row[0] == "Eg"]) == 1
    assert len([row for row in rows if row[2] == "INICIO_CFDI"]) == 2
    assert len([row for row in rows if row[2] == "FIN_CFDI"]) == 2
    assert {row[8] for row in rows if row[8]} == {"UUID-1", "UUID-2"}


def test_report_policy_preserves_tax_movements_and_balance():
    expenses = [
        _expense(
            export_reference=f"G-{index}",
            total=116,
            subtotal_amount=100,
            iva_amount=16,
            cuenta_iva="1180",
            poliza_group_key="informe:tax",
            poliza_reference="I-TAX",
            poliza_description="Informe I-TAX",
        )
        for index in (1, 2)
    ]
    rows = build_coi_poliza_rows(expenses)
    movements = [row for row in rows if row[4] == "1"]

    assert len([row for row in rows if row[0] == "Eg"]) == 1
    assert len([row for row in movements if row[1] == "1180"]) == 2
    assert sum(float(row[5] or 0) for row in movements) == 232
    assert sum(float(row[6] or 0) for row in movements) == 232


def test_standalone_zip_preserves_references_and_duplicate_filename_suffix():
    payload = generate_coi_poliza_zip(
        [
            _expense(export_reference="G-1"),
            _expense(export_reference="G-1"),
        ]
    )
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        assert archive.namelist() == ["Poliza_COI_G-1.xlsx", "Poliza_COI_002_G-1.xlsx"]


def _coi_sheet(expenses):
    workbook = load_workbook(io.BytesIO(generate_coi_poliza_xlsx(expenses)))
    return workbook["Poliza COI"]


def test_solicitud_xlsx_writes_operation_provider_and_tournament_in_c2_to_e2():
    expense = _expense(export_reference="O-26000001")
    document = SimpleNamespace(
        id="solicitud-1",
        tipo="SOLICITUD",
        referencia_operaciones="104",
        proveedor_cliente=SimpleNamespace(nombre="Servicios Deportivos SA de CV"),
        beneficiario_empleado=None,
        beneficiario_proveedor_cliente=None,
        empleado=None,
        torneo=SimpleNamespace(name="Copa Telmex 2026"),
        proyecto_otro=None,
        cuenta_gastos=None,
    )

    sheet = _coi_sheet(group_expense_cfdis_for_document([expense], document))

    assert sheet["A2"].value == "|||"
    assert sheet["C2"].value == "104"
    assert sheet["D2"].value == "Servicios Deportivos SA de CV"
    assert sheet["E2"].value == "Copa Telmex 2026"
    assert sheet["A3"].value == "Eg"


def test_solicitud_xlsx_uses_beneficiary_when_project_is_missing():
    expense = _expense(export_reference="O-26000002")
    document = SimpleNamespace(
        id="solicitud-2",
        tipo="SOLICITUD",
        referencia_operaciones="107",
        proveedor_cliente=SimpleNamespace(nombre="Proveedor Dos"),
        beneficiario_empleado=SimpleNamespace(nombre="Beneficiaria Dos"),
        beneficiario_proveedor_cliente=None,
        empleado=None,
        torneo=None,
        proyecto_otro=None,
        cuenta_gastos=None,
    )

    sheet = _coi_sheet(group_expense_cfdis_for_document([expense], document))

    assert sheet["D2"].value == "Proveedor Dos"
    assert sheet["E2"].value == "Beneficiaria Dos"


def test_informe_xlsx_uses_beneficiary_and_expense_reason_fallback():
    expenses = [_expense(export_reference="G-1"), _expense(export_reference="G-2")]
    document = SimpleNamespace(
        id="informe-1",
        tipo="INFORME",
        numero_referencia="I-26000001",
        referencia_operaciones="105",
        proveedor_cliente=None,
        beneficiario_empleado=SimpleNamespace(nombre="Ana Pérez"),
        beneficiario_proveedor_cliente=None,
        empleado=SimpleNamespace(nombre="Solicitante"),
        torneo=None,
        proyecto_otro=None,
        cuenta_gastos=SimpleNamespace(
            nombre="Viáticos para eliminatoria nacional",
            torneo=None,
            beneficiario_empleado=None,
            beneficiario_proveedor_cliente=None,
            empleado=None,
        ),
    )

    sheet = _coi_sheet(group_expense_cfdis_for_document(expenses, document))

    assert sheet["C2"].value == "105"
    assert sheet["D2"].value == "Ana Pérez"
    assert sheet["E2"].value == "Viáticos para eliminatoria nacional"
    assert len([row for row in sheet.iter_rows() if row[0].value == "Eg"]) == 1


def test_mixed_document_xlsx_marks_metadata_as_multiple_documents():
    first = _expense(
        poliza_document_id="documento-1",
        poliza_operation_reference="104",
        poliza_party_name="Proveedor Uno",
        poliza_context_description="Torneo Uno",
    )
    second = _expense(
        poliza_document_id="documento-2",
        poliza_operation_reference="105",
        poliza_party_name="Beneficiario Dos",
        poliza_context_description="Torneo Dos",
    )

    sheet = _coi_sheet([first, second])

    assert [sheet[cell].value for cell in ("C2", "D2", "E2")] == [
        "Múltiples documentos",
        "Múltiples documentos",
        "Múltiples documentos",
    ]
