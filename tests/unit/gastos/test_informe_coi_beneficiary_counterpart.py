from datetime import datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from devnous.gastos.models import CuentaDeGastos
from devnous.gastos.services import expense_coi_export_service as coi
from devnous.gastos.services.coi_poliza_exporter import generate_coi_poliza_xlsx


def _account(code: str, name: str = ""):
    return SimpleNamespace(
        id=uuid4(),
        codigo=code,
        nombre=name or code,
        activo=True,
    )


def _expense(*, company_amex: bool = False):
    cfdi = SimpleNamespace(
        cfdi_uuid="UUID-TEST",
        fecha=datetime(2026, 9, 30),
        emisor_rfc="AAA010101AAA",
        receptor_rfc="BBB010101BBB",
        folio="F-1",
        emisor_nombre="Proveedor",
        receptor_uso_cfdi="G03",
    )
    return SimpleNamespace(
        id=uuid4(),
        fecha=datetime(2026, 9, 30),
        gasto_cantidad=116.0,
        concepto="Hospedaje",
        proyecto="Copa Telmex",
        numero_referencia="G-TEST",
        cuenta_gastos_id=uuid4(),
        pagado_con_amex_empresa=company_amex,
        origen="informe_quick_entry",
        cuenta_contable=_account("5300-010-001"),
        contra_cuenta_contable=_account("2120-000-000", "ACREEDORES DIVERSOS"),
        cfdi_report=cfdi,
        cfdi_report_id=uuid4(),
    )


def _preview(counterpart: str):
    return {
        "contra_account": {
            "codigo": counterpart,
            "cuenta_contable_id": str(uuid4()),
            "nombre": "Contrapartida",
        },
        "taxes": {
            "base_gasto": 100.0,
            "iva_trasladado": 16.0,
            "iva_account": {
                "codigo": "1180-001-001",
                "cuenta_contable_id": str(uuid4()),
                "nombre": "IVA",
            },
            "retenciones": [],
            "retenciones_total": 0.0,
            "impuestos_locales": [],
            "gastos_no_deducibles": [],
            "neto_contrapartida": 116.0,
        },
    }


class _Session:
    def __init__(self, cuenta):
        self.cuenta = cuenta

    async def get(self, model, identifier):
        if model is CuentaDeGastos and identifier == self.cuenta.id:
            return self.cuenta
        return None


@pytest.mark.asyncio
@pytest.mark.parametrize("expense_code", ["5300-010-001", "5300-010-031"])
async def test_informe_coi_uses_beneficiary_detail_account_over_generic_payable(
    monkeypatch,
    expense_code,
):
    expense = _expense()
    expense.cuenta_contable = _account(expense_code)
    cuenta = SimpleNamespace(id=expense.cuenta_gastos_id)
    employee = SimpleNamespace(id=uuid4(), nombre="Persona beneficiaria")
    debtor = _account("1170-001-007", "Persona beneficiaria")

    monkeypatch.setattr(
        coi,
        "build_cleanup_preview",
        AsyncMock(return_value={"status": "Listo COI", "issues": []}),
    )
    monkeypatch.setattr(
        coi,
        "build_expense_accounting_preview",
        AsyncMock(return_value=_preview("2120-000-000")),
    )
    employee_resolver = AsyncMock(return_value=employee)
    debtor_resolver = AsyncMock(return_value=debtor)
    monkeypatch.setattr(coi, "resolve_cuenta_debtor_empleado", employee_resolver)
    monkeypatch.setattr(coi, "resolve_cuenta_debtor_account", debtor_resolver)

    payload = await coi.build_expense_cfdi_for_export(_Session(cuenta), expense)

    assert payload.cuenta_contrapartida == "1170-001-007"
    assert payload.cuenta_contable == expense_code
    workbook = load_workbook(BytesIO(generate_coi_poliza_xlsx([payload])))
    movement_rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    assert any(row[1] == expense_code and row[5] == 100 for row in movement_rows)
    assert any(row[1] == "1170-001-007" and row[6] == 116 for row in movement_rows)
    assert not any(row[1] == "2120-000-000" for row in movement_rows)
    employee_resolver.assert_awaited_once()
    debtor_resolver.assert_awaited_once()


@pytest.mark.asyncio
async def test_informe_coi_blocks_when_beneficiary_has_no_detail_account(monkeypatch):
    expense = _expense()
    cuenta = SimpleNamespace(id=expense.cuenta_gastos_id)
    employee = SimpleNamespace(id=uuid4(), nombre="Persona beneficiaria")

    monkeypatch.setattr(
        coi,
        "build_cleanup_preview",
        AsyncMock(return_value={"status": "Listo COI", "issues": []}),
    )
    monkeypatch.setattr(
        coi, "resolve_cuenta_debtor_empleado", AsyncMock(return_value=employee)
    )
    monkeypatch.setattr(
        coi, "resolve_cuenta_debtor_account", AsyncMock(return_value=None)
    )

    ready, issues = await coi.assess_expense_coi_cleanup_ready(
        _Session(cuenta), expense
    )

    assert ready is False
    assert any("subcuenta contable de detalle" in issue for issue in issues)


@pytest.mark.asyncio
async def test_company_amex_in_informe_keeps_its_configured_counterpart(monkeypatch):
    expense = _expense(company_amex=True)
    cuenta = SimpleNamespace(id=expense.cuenta_gastos_id)
    debtor_resolver = AsyncMock(
        side_effect=AssertionError("AMEX empresa must not use employee debtor account")
    )

    monkeypatch.setattr(
        coi,
        "build_cleanup_preview",
        AsyncMock(return_value={"status": "Listo COI", "issues": []}),
    )
    monkeypatch.setattr(
        coi,
        "build_expense_accounting_preview",
        AsyncMock(return_value=_preview("2130-010-001")),
    )
    monkeypatch.setattr(coi, "resolve_cuenta_debtor_account", debtor_resolver)

    payload = await coi.build_expense_cfdi_for_export(_Session(cuenta), expense)

    assert payload.cuenta_contrapartida == "2130-010-001"
    debtor_resolver.assert_not_awaited()


@pytest.mark.asyncio
async def test_preassessed_batch_reuses_cached_beneficiary_counterpart(monkeypatch):
    expense = _expense()
    cuenta = SimpleNamespace(id=expense.cuenta_gastos_id)
    employee = SimpleNamespace(id=uuid4(), nombre="Persona beneficiaria")
    debtor = _account("1170-001-007", "Persona beneficiaria")
    session = _Session(cuenta)

    monkeypatch.setattr(
        coi,
        "build_cleanup_preview",
        AsyncMock(return_value={"status": "Listo COI", "issues": []}),
    )
    monkeypatch.setattr(
        coi,
        "build_expense_accounting_preview",
        AsyncMock(return_value=_preview("2120-000-000")),
    )
    employee_resolver = AsyncMock(return_value=employee)
    debtor_resolver = AsyncMock(return_value=debtor)
    monkeypatch.setattr(coi, "resolve_cuenta_debtor_empleado", employee_resolver)
    monkeypatch.setattr(coi, "resolve_cuenta_debtor_account", debtor_resolver)

    ready, issues = await coi.assess_expense_coi_cleanup_ready(session, expense)
    assert ready is True
    assert issues == []

    payload = await coi.build_expense_cfdi_for_export(
        session, expense, require_cleanup_ready=False
    )

    assert payload.cuenta_contrapartida == "1170-001-007"
    employee_resolver.assert_awaited_once()
    debtor_resolver.assert_awaited_once()
