"""Synthetic context and workbook invariants against release 612bd50bf.

The three reported references identify scenarios only. No attached workbook or
real supplier text is reproduced; missing historical text is never reconstructed.
"""

import io
import zipfile
from copy import deepcopy
from datetime import datetime
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook

from devnous.gastos.services import coi_poliza_exporter as exporter
from devnous.gastos.services.import_coi_service import parse_coi_workbook


def _expense(**overrides):
    fields = dict(
        fecha=datetime(2026, 10, 2),
        total=116.0,
        iva_amount=16.0,
        subtotal_amount=100.0,
        concepto="Traslado médico y alimentación",
        cuenta_contable="5300-001-001",
        cuenta_contrapartida="2100-001-001",
        export_reference="O-FIXTURE",
        nombre_emisor="Proveedor ficticio",
        proyecto="Proyecto de prueba",
        poliza_document_id="fixture-document",
        poliza_operation_reference="OP-PRUEBA",
        poliza_party_name="Proveedor ficticio",
        poliza_context_description="Proyecto de prueba",
    )
    fields.update(overrides)
    return exporter.ExpenseCFDI(**fields)


def _workbook(expenses, **kwargs):
    return load_workbook(
        io.BytesIO(exporter.generate_coi_poliza_xlsx(expenses, **kwargs))
    )


def _header(expense):
    return _workbook([expense])["Poliza COI"]["C3"].value


@pytest.mark.parametrize(
    "reference,length,ending",
    [
        ("O-26001241", 115, "motivo de prueba"),
        ("F-26000090", 120, "Dirección G"),
        ("F-26000091", 120, "Dirección Gen"),
    ],
)
def test_reported_lengths_preserve_supplied_text_without_reconstructing_it(
    reference, length, ending
):
    prefix = f"{reference} / Proveedor ficticio / "
    supplied = prefix + "á" * (length - len(prefix) - len(ending)) + ending
    expense = _expense(
        export_reference=reference,
        poliza_description=supplied,
        poliza_context_description="",
        poliza_operation_reference="",
        posting_movements=[
            dict(cuenta="5300", concepto=supplied, debe=116, haber=0),
            dict(cuenta="2100", concepto="Concepto manual distinto", debe=0, haber=116),
        ],
    )
    sheet = _workbook([expense])["Poliza COI"]
    assert len(supplied) == length
    assert sheet["D4"].value == supplied
    assert sheet["D5"].value == "Concepto manual distinto"
    assert sheet["C3"].value == "Beneficiario: Proveedor ficticio / " + supplied
    assert sheet["C3"].value.endswith(ending)
    assert supplied in sheet["C3"].value


def test_header_deduplicates_fields_while_reason_precedes_additional_operation():
    expense = _expense(export_reference="O-26001241")
    header = _header(expense)
    assert header == (
        "Beneficiario: Proveedor ficticio / O-26001241 / "
        "Traslado médico y alimentación / Proyecto de prueba / OP-PRUEBA"
    )
    for value in ("Proveedor ficticio", "Proyecto de prueba", "OP-PRUEBA"):
        assert header.count(value) == 1
    assert "Operaciones:" not in header
    assert "Contexto:" not in header


def test_long_accented_reason_is_not_truncated_at_120():
    reason = "Atención médica, alimentación y transporte. " * 100
    expense = _expense(concepto=reason)
    header = _header(expense)
    assert reason.strip() in header
    assert len(header) > 4000
    assert _workbook([expense])["Poliza COI"]["D4"].value == (
        f"O-FIXTURE / Proveedor ficticio / {reason.strip()} / Proyecto de prueba"
    )


@pytest.mark.parametrize(
    "operation,party,context,expected,beneficiary,summary",
    [
        ("", "", "", "Motivo", None, "Motivo"),
        (" OP ", "", "", "Operaciones: OP / Motivo", None, "Motivo"),
        ("", "", " Proyecto ", "Contexto: Proyecto / Motivo", None, "Motivo"),
        (
            "",
            " Ana  Pérez ",
            "",
            "Beneficiario: Ana Pérez / Motivo",
            "Ana Pérez",
            "Motivo",
        ),
        (
            "OP",
            "",
            "Proyecto",
            "Operaciones: OP / Motivo / Proyecto",
            None,
            "Motivo / Proyecto",
        ),
        (
            "",
            "Ana",
            "Proyecto",
            "Beneficiario: Ana / Motivo / Proyecto",
            "Ana",
            "Motivo / Proyecto",
        ),
    ],
)
def test_empty_fields_and_whitespace_keep_importer_unambiguous(
    operation, party, context, expected, beneficiary, summary
):
    expense = _expense(
        poliza_description="Motivo",
        poliza_operation_reference=operation,
        poliza_party_name=party,
        poliza_context_description=context,
    )
    assert _header(expense) == expected
    parsed = parse_coi_workbook(
        "synthetic.xlsx", exporter.generate_coi_poliza_xlsx([expense])
    )
    assert parsed[0].beneficiario_nombre == beneficiary
    assert parsed[0].concepto_resumen == summary


def test_no_metadata_keeps_manual_policy_description_exactly():
    expense = _expense(
        poliza_description="Manual  sin / sustitución",
        poliza_document_id=None,
        poliza_operation_reference=None,
        poliza_party_name=None,
        poliza_context_description=None,
    )
    assert _header(expense) == expense.poliza_description


def test_explicit_policy_description_is_not_deduplicated_as_metadata():
    expense = _expense(
        poliza_description=(
            "  REF / Ana  Pérez / Ana Pérez Norte / MÉDICO / "
            "Proyecto de prueba ampliado"
        ),
        poliza_party_name="Ana Pérez",
        poliza_context_description="Proyecto de prueba",
    )
    assert _header(expense) == (
        "Beneficiario: Ana Pérez / REF / Ana Pérez / Ana Pérez Norte / MÉDICO / "
        "Proyecto de prueba ampliado / Proyecto de prueba / OP-PRUEBA"
    )


@pytest.mark.parametrize(
    "reference,counterpart,reason,project,expected_body",
    [
        ("O-1", "Proveedor X", "Proveedor X", "", "O-1 / Proveedor X"),
        ("", "", "Proveedor X", "", "Proveedor X"),
        (
            "Proveedor X",
            "Proveedor X",
            "Proveedor X",
            "Proveedor X",
            "Proveedor X / Proveedor X / Proveedor X",
        ),
        (
            "O-1",
            "Proveedor X Norte",
            "Motivo",
            "Proyecto",
            "O-1 / Proveedor X Norte / Motivo / Proyecto",
        ),
        (
            "REF / Proveedor X / SUFIJO",
            "Proveedor X",
            "Motivo",
            "",
            "REF / Proveedor X / SUFIJO / Proveedor X / Motivo",
        ),
    ],
)
def test_only_generated_counterpart_is_removed_when_values_match(
    reference, counterpart, reason, project, expected_body
):
    expense = _expense(
        export_reference=reference,
        nombre_emisor=counterpart,
        concepto=reason,
        proyecto=project,
        poliza_party_name="Proveedor X",
        poliza_operation_reference="",
        poliza_context_description="",
    )
    before = exporter._expense_description(expense)
    payload = exporter.generate_coi_poliza_xlsx([expense])
    sheet = load_workbook(io.BytesIO(payload))["Poliza COI"]
    assert sheet["C3"].value == "Beneficiario: Proveedor X / " + expected_body
    assert sheet["D4"].value == before
    parsed = parse_coi_workbook("synthetic.xlsx", payload)
    assert parsed[0].beneficiario_nombre == "Proveedor X"
    assert parsed[0].concepto_resumen == expected_body


def test_same_operation_and_context_are_emitted_once():
    expense = _expense(
        poliza_operation_reference="O-FIXTURE",
        poliza_context_description="Proveedor ficticio",
    )
    header = _header(expense)
    assert header.count("O-FIXTURE") == 1
    assert header.count("Proveedor ficticio") == 1


def _release_header(group):
    """Frozen header algorithm from 612bd50bf, used only as a diff oracle."""
    metadata = {}
    has_metadata_free_policy = False
    for index, expense in enumerate(group.expenses, start=1):
        values = tuple(
            " ".join((value or "").split())
            for value in (
                expense.poliza_operation_reference,
                expense.poliza_party_name,
                expense.poliza_context_description,
            )
        )
        document_id = " ".join((expense.poliza_document_id or "").split())
        if not document_id and not any(values):
            has_metadata_free_policy = True
            continue
        metadata.setdefault(document_id or f"metadata:{index}:{values!r}", values)
    if not metadata:
        return group.description
    if has_metadata_free_policy or len(metadata) > 1:
        parts = ["Múltiples documentos"]
    else:
        operation, party, context = next(iter(metadata.values()))
        parts = [
            f"Operaciones: {operation}" if operation else "",
            f"Beneficiario: {party}" if party else "",
            f"Contexto: {context}" if context else "",
        ]
    return " / ".join(
        " ".join(value.split())
        for value in [*parts, group.description]
        if value and value.strip()
    )


@pytest.mark.parametrize(
    "scenario", ["single", "report", "mixed", "frozen", "metadata_free_member"]
)
def test_release_equivalence_outside_header_text(monkeypatch, scenario):
    expenses = [
        _expense(
            base_amount=95.0,
            cfdi_uuid="00000000-0000-4000-8000-000000000001",
            cfdi_date=datetime(2026, 9, 30),
            rfc_emisor="AAA010101AAA",
            rfc_receptor="BBB010101BBB",
            retenciones=[dict(label="ISR", importe=10, cuenta_contable="2160")],
            impuestos_locales=[dict(label="ISH", importe=2, cuenta_contable="5301")],
            gastos_no_deducibles=[
                dict(label="Propina", importe=3, cuenta_contable="5302")
            ],
        )
    ]
    if scenario in {"report", "mixed", "metadata_free_member"}:
        expenses[0].poliza_group_key = "informe:fixture"
        expenses[0].poliza_description = "Informe de Gastos I-FIXTURE"
        expenses.append(
            _expense(
                export_reference="F-FIXTURE",
                concepto="Concepto individual de segunda operación",
                poliza_group_key="informe:fixture",
                poliza_description="Informe de Gastos I-FIXTURE",
                cfdi_uuid="00000000-0000-4000-8000-000000000002",
            )
        )
    if scenario == "mixed":
        expenses[1].poliza_document_id = "other-document"
        expenses[1].poliza_operation_reference = "OTHER-OP"
    if scenario == "metadata_free_member":
        for field in (
            "poliza_document_id",
            "poliza_operation_reference",
            "poliza_party_name",
            "poliza_context_description",
        ):
            setattr(expenses[1], field, None)
    if scenario == "frozen":
        expenses[0].poliza_group_key = "amex-cut:fixture"
        expenses[0].posting_movements = [
            dict(
                cuenta="5300", concepto="Manual / motivo particular", debe=116, haber=0
            ),
            dict(cuenta="2100", concepto="Otra partida específica", debe=0, haber=116),
        ]

    preimage = deepcopy(expenses)
    assert not [
        issue
        for issue in exporter.validate_coi_poliza(expenses)
        if issue["severity"] == "error"
    ]
    manifest = [["referencia", "proyecto"], ["O-FIXTURE", "Proyecto de prueba"]]
    csv_before = exporter.generate_coi_poliza_csv(expenses)
    rows_before = exporter.build_coi_poliza_rows(expenses)
    with monkeypatch.context() as baseline:
        baseline.setattr(exporter, "_coi_policy_xlsx_description", _release_header)
        before = exporter.generate_coi_poliza_xlsx(
            expenses, lote_manifest_rows=manifest
        )
    after = exporter.generate_coi_poliza_xlsx(expenses, lote_manifest_rows=manifest)
    assert expenses == preimage
    assert exporter.build_coi_poliza_rows(expenses) == rows_before
    assert exporter.generate_coi_poliza_csv(expenses) == csv_before

    wb_before = load_workbook(io.BytesIO(before))
    wb_after = load_workbook(io.BytesIO(after))
    assert wb_before.sheetnames == wb_after.sheetnames
    header_cells = {
        row[2].coordinate
        for row in wb_before["Poliza COI"].iter_rows()
        if row[0].value == "Eg"
    }
    for name in wb_before.sheetnames:
        old, new = wb_before[name], wb_after[name]
        assert old.calculate_dimension() == new.calculate_dimension()
        for old_row, new_row in zip(old.iter_rows(), new.iter_rows(), strict=True):
            for a, b in zip(old_row, new_row, strict=True):
                assert a._style == b._style
                assert a.number_format == b.number_format
                if name != "Poliza COI" or a.coordinate not in header_cells:
                    assert a.value == b.value
                    assert a.data_type == b.data_type

    # Full package comparison verifies sheets, identifiers, grouping, row order,
    # styles and dimensions; only Eg column C text and generation times may vary.
    with zipfile.ZipFile(io.BytesIO(before)) as old_zip, zipfile.ZipFile(
        io.BytesIO(after)
    ) as new_zip:
        assert old_zip.namelist() == new_zip.namelist()
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        for name in old_zip.namelist():
            old, new = old_zip.read(name), new_zip.read(name)
            if name == "xl/worksheets/sheet1.xml":
                roots = [ET.fromstring(old), ET.fromstring(new)]
                for root in roots:
                    for cell in root.findall(".//s:c", ns):
                        if cell.attrib["r"] in header_cells:
                            cell.find("s:is/s:t", ns).text = "HEADER_TEXT"
                assert ET.tostring(roots[0]) == ET.tostring(roots[1])
            elif name == "docProps/core.xml":
                roots = [ET.fromstring(old), ET.fromstring(new)]
                for root in roots:
                    for element in root:
                        if element.tag.endswith(("}created", "}modified")):
                            element.text = "GENERATION_TIME"
                assert ET.tostring(roots[0]) == ET.tostring(roots[1])
            else:
                assert old == new, name
