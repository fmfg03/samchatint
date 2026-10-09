from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

from openpyxl import load_workbook

from devnous.gastos.services.diot_exporter import (
    build_diot_export,
    create_diot_excel,
)


def test_monthly_export_uses_effective_date_and_prorates_shared_cfdi():
    cfdi = SimpleNamespace(
        cfdi_uuid="SHARED-UUID",
        emisor_rfc="AAA010101AAA",
        emisor_nombre="Proveedor",
        receptor_rfc="BBB010101BBB",
        descripcion_concepto_principal="Servicio",
        subtotal=Decimal("1000"),
        total=Decimal("1040"),
        total_impuestos_trasladados=Decimal("160"),
        impuestos_detalle={
            "traslados": [
                {
                    "base": 1000,
                    "importe": 160,
                    "impuesto": "002",
                    "tipo_factor": "Tasa",
                    "tasa_o_cuota": 0.16,
                }
            ],
            "retenciones": [
                {"importe": 120, "impuesto": "002"},
            ],
        },
    )
    expense = SimpleNamespace(
        id=uuid4(),
        numero_referencia="O-TEST",
        fecha=date(2026, 8, 1),
        gasto_cantidad=Decimal("520"),
        propina_no_deducible=Decimal("0"),
        iva=Decimal("80"),
        concepto="Servicio",
        archivo_nombre="",
        link_xml="",
        cuenta_contable_base="",
        cfdi_report=cfdi,
        cfdi_compartido_confirmado=True,
    )

    export = build_diot_export(
        [expense],
        effective_payment_dates={str(expense.id): date(2026, 9, 15)},
    )

    row = export.detail_rows[0]
    assert row.fecha_gasto == date(2026, 9, 15)
    assert row.subtotal_cfdi == Decimal("500.00")
    assert row.base_16 == Decimal("500.00")
    assert row.iva_trasladado == Decimal("80.00")
    assert row.iva_retenido == Decimal("60.00")
    assert row.total_cfdi == Decimal("520.00")


def test_excel_adds_monthly_blocker_and_undated_sheets():
    export = build_diot_export([])
    content = create_diot_excel(
        export,
        audit_issues={
            "Bloqueos": [
                {
                    "referencia": "O-1",
                    "fecha_pago_efectiva": "2026-09-01",
                    "codigo": "cfdi_faltante",
                    "detalle": "Falta CFDI",
                    "uuid_cfdi": "",
                }
            ],
            "Sin fecha efectiva": [],
        },
    )

    workbook = load_workbook(BytesIO(content), data_only=False)
    assert "Bloqueos" in workbook.sheetnames
    assert "Sin fecha efectiva" in workbook.sheetnames
    assert workbook["Bloqueos"]["A2"].value == "O-1"


def test_monthly_export_does_not_infer_retention_without_xml_retenciones():
    cfdi = SimpleNamespace(
        cfdi_uuid="NO-RETENTION-UUID",
        emisor_rfc="AAA010101AAA",
        emisor_nombre="Proveedor",
        receptor_rfc="BBB010101BBB",
        descripcion_concepto_principal="Servicio",
        subtotal=Decimal("10000"),
        total=Decimal("11600"),
        total_impuestos_trasladados=Decimal("1600"),
        impuestos_detalle={
            "traslados": [
                {
                    "base": 10000,
                    "importe": 1600,
                    "impuesto": "002",
                    "tipo_factor": "Tasa",
                    "tasa_o_cuota": 0.16,
                }
            ],
            "retenciones": [],
        },
    )
    expense = SimpleNamespace(
        id=uuid4(),
        numero_referencia="S-26000416",
        fecha=date(2026, 9, 1),
        gasto_cantidad=Decimal("11600"),
        propina_no_deducible=Decimal("0"),
        iva=Decimal("1600"),
        concepto="Servicio",
        archivo_nombre="",
        link_xml="",
        cuenta_contable_base="",
        cfdi_report=cfdi,
        cfdi_compartido_confirmado=False,
    )

    export = build_diot_export([expense])

    assert export.detail_rows[0].iva_retenido == Decimal("0")
    assert export.summary_rows[0].amounts["iva_retenido"] == Decimal("0")


def test_invalid_shared_application_is_not_prorated_and_warns():
    cfdi = SimpleNamespace(
        cfdi_uuid="SHARED-INVALID",
        emisor_rfc="AAA010101AAA",
        emisor_nombre="Proveedor",
        receptor_rfc="BBB010101BBB",
        descripcion_concepto_principal="Servicio",
        subtotal=Decimal("100"),
        total=Decimal("116"),
        total_impuestos_trasladados=Decimal("16"),
        impuestos_detalle={"traslados": [], "retenciones": []},
    )
    expense = SimpleNamespace(
        id=uuid4(),
        numero_referencia="O-SHARED-INVALID",
        fecha=date(2026, 9, 1),
        gasto_cantidad=Decimal("200"),
        propina_no_deducible=Decimal("0"),
        iva=Decimal("0"),
        concepto="Servicio",
        archivo_nombre="",
        link_xml="",
        cuenta_contable_base="",
        cfdi_report=cfdi,
        cfdi_compartido_confirmado=True,
    )

    export = build_diot_export([expense])

    assert export.detail_rows[0].total_cfdi == Decimal("116.00")
    assert "no se prorrateó" in export.detail_rows[0].warnings[0]
