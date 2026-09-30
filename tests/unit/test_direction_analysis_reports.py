"""Financial arithmetic, signed continuity and human export parity for #432/#433."""

import copy
import io
from datetime import date
from decimal import Decimal
from unittest.mock import AsyncMock

import fitz
import pytest
from openpyxl import load_workbook
from test_direction_home import ACTOR, T1, row, scope, snapshot
from test_direction_home_routes import ask

from devnous.gastos.routes import client_executive_routes as routes
from samchat.client_executive import home
from samchat.client_executive.analysis import (
    AnalysisError,
    calculate_scenario,
    deviations,
    scenario_from_question,
)
from samchat.client_executive.conversation import (
    ContextError,
    answer_snapshot,
    load_analysis,
    sign_analysis,
)
from samchat.client_executive.reports import (
    build_report,
    render_published_report,
    render_report,
)
from samchat.executive.exporter import (
    generate_direction_report_pdf,
    generate_direction_report_xlsx,
)

pytest_plugins = ["test_direction_home_routes"]


def rich_snapshot():
    data = snapshot()
    for r in data["tournaments"]:
        r["concepts"] = {
            "status": "available",
            "rows": [
                {
                    "id": "abc",
                    "label": "Operación",
                    "budget": "10",
                    "actual": r["values"]["actual"],
                    "excess": str(Decimal(r["values"]["actual"]) - 10),
                }
            ],
        }
        r["values"]["receivables"] = "80"
    return data


def test_expense_decimal_rounding_and_continuity_without_mutation():
    data = rich_snapshot()
    before = copy.deepcopy(data)
    answer = answer_snapshot(data, "actual", "¿Qué pasa si reduzco el gasto 10%?")
    assert answer["scenario"]["base"] == "100.00"
    assert answer["scenario"]["result"] == "90.00"
    revised = answer_snapshot(data, "actual", "Mejor 5%", previous=answer)
    assert revised["scenario"]["result"] == "95.00"
    assert data == before
    assert scenario_from_question("mejor 5%", {"scenario": None}) is None
    single = snapshot([row(actual="0.05")])
    assert (
        calculate_scenario(single, {"kind": "expense_reduction", "percent": "10"})[
            "effect"
        ]
        == "0.01"
    )


@pytest.mark.parametrize(
    "spec",
    [
        {"kind": "bad"},
        {"kind": "expense_reduction"},
        {"kind": "expense_reduction", "percent": "-1"},
        {"kind": "expense_reduction", "percent": "101"},
        {"kind": "expense_reduction", "percent": "10", "tournament_id": "foreign"},
        {"kind": "expense_reduction", "percent": "10", "concept_id": "foreign"},
        {"kind": "payment_delay", "days": True},
        {"kind": "payment_delay", "days": 366},
        {"kind": "collection_acceleration", "percent": "10", "concept_id": "abc"},
    ],
)
def test_invalid_scenarios_have_no_result(spec):
    with pytest.raises(AnalysisError):
        calculate_scenario(rich_snapshot(), spec)


def test_future_concept_and_collection_bases_are_distinct():
    data = rich_snapshot()
    future = calculate_scenario(
        data, {"kind": "expense_reduction", "basis": "future_expense", "percent": "10"}
    )
    assert future["base"] == "140.00" and future["result"] == "126.00"
    concept = calculate_scenario(
        data,
        {
            "kind": "expense_reduction",
            "percent": "10",
            "tournament_id": T1,
            "concept_id": "abc",
        },
    )
    assert concept["result"] == "36.00"
    collected = calculate_scenario(
        data, {"kind": "collection_acceleration", "percent": "25"}
    )
    assert collected["effect"] == "40.00" and collected["result"] == "120.00"
    data["tournaments"][0]["values"]["actual"] = None
    with pytest.raises(AnalysisError):
        calculate_scenario(data, {"kind": "expense_reduction", "percent": "10"})
    data["tournaments"][0]["values"]["receivables"] = None
    with pytest.raises(AnalysisError):
        calculate_scenario(data, {"kind": "collection_acceleration", "percent": "10"})


def test_payment_calendar_reconciles_and_never_cancels_debt():
    data = snapshot([row()])
    r = data["tournaments"][0]
    r["values"]["obligations"] = "50"
    r["payment_evidence"] = [
        {"date": "2026-10-20", "value": "50", "reference": "REQ-1"}
    ]
    result = calculate_scenario(data, {"kind": "payment_delay", "days": 15})
    assert result["base"] == "50.00" and result["result"] == "0"
    assert any("no reduce la deuda" in v for v in result["limits"])
    r["payment_evidence"][0]["value"] = "49"
    with pytest.raises(AnalysisError):
        calculate_scenario(data, {"kind": "payment_delay", "days": 15})
    r["payment_evidence"][0].update(value="50", date="bad")
    with pytest.raises(AnalysisError):
        calculate_scenario(data, {"kind": "payment_delay", "days": 15})


def test_deviations_sorted_and_response_has_facts_hypotheses_advice():
    data = rich_snapshot()
    items, gaps = deviations(data)
    assert [Decimal(i["amount"]) for i in items] == sorted(
        [Decimal(i["amount"]) for i in items], reverse=True
    )
    assert {"forecast", "concept"} == {i["basis"] for i in items}
    for q in (
        "¿Dónde me estoy pasando?",
        "¿Por qué?",
        "¿Cómo vamos a cerrar?",
        "¿Qué harías?",
        "Compara periodos anteriores",
        "Haz un reporte para socios",
    ):
        answer = answer_snapshot(data, "actual", q)
        assert answer["conclusion"] and answer["recommendation"] and answer["next_step"]
        assert "Hechos" in answer["assistant_message"]
    answer = answer_snapshot(data, "actual", "Por qué")
    assert "por validar" in answer["hypotheses"][0]
    assert answer_snapshot(data, "actual", "Haz un reporte para socios")[
        "report_requested"
    ]
    data["tournaments"][0]["previous_period"] = {
        "period": {"start": "2026-01-01", "end": "2026-01-31"},
        "values": {"actual": "20"},
        "gaps": [],
    }
    assert any(
        "diferencia" in v
        for v in answer_snapshot(data, "actual", "Compara periodos anteriores")["facts"]
    )
    assert answer_snapshot(data, "budget", "Compara periodos anteriores")[
        "missing_evidence"
    ]


def test_concept_reconciliation_requires_all_rows_and_complete_money():
    values = {"actual": Decimal("40"), "budget": Decimal("100")}
    source = {
        "executive_concepts": [
            {
                "concept_id": "canonical-1",
                "label": "Partida",
                "actual_total": "40",
                "budget_total": "100",
            }
        ]
    }
    assert home.concept_evidence(source, values)["status"] == "available"
    source["executive_concepts"][0]["actual_total"] = "39"
    assert home.concept_evidence(source, values)["rows"] == []
    source["executive_concepts"][0]["actual_total"] = None
    assert home.concept_evidence(source, values)["status"] == "unavailable"


@pytest.mark.asyncio
async def test_previous_period_equal_duration_and_edition_boundary(monkeypatch):
    previous = AsyncMock(return_value={})
    monkeypatch.setattr(home, "_optional_read", previous)
    monkeypatch.setattr(
        home,
        "budget_values",
        lambda *a, **k: (
            {"actual": Decimal("20"), "committed": None, "paid": Decimal("5")},
            [],
        ),
    )
    result = await home.previous_period_values(
        None, {}, date(2026, 3, 1), date(2026, 3, 10), 2026
    )
    assert result["period"] == {"start": "2026-02-19", "end": "2026-02-28"}
    assert result["values"]["actual"] == "20"
    previous.reset_mock()
    result = await home.previous_period_values(
        None, {}, date(2026, 1, 1), date(2026, 1, 10), 2026
    )
    assert result["gaps"] and not result["values"]
    previous.assert_not_awaited()


def test_signed_analysis_binds_actor_cut_and_tamper(monkeypatch):
    monkeypatch.setenv("SESSION_SECRET_KEY", "test-only-analysis")
    data = snapshot()
    answer = answer_snapshot(data, "actual", "Reduce 10%")
    receipt = sign_analysis(answer, data["snapshot_id"], ACTOR)
    assert (
        load_analysis(receipt, data["snapshot_id"], ACTOR)["scenario"]["result"]
        == "90.00"
    )
    for token, cut, actor in (
        (receipt + "x", data["snapshot_id"], ACTOR),
        (receipt, "other", ACTOR),
        (receipt, data["snapshot_id"], "other"),
    ):
        with pytest.raises(ContextError):
            load_analysis(token, cut, actor)


def test_pdf_workbook_share_numbers_cut_gaps_and_formula_cache(tmp_path):
    data = rich_snapshot()
    answer = answer_snapshot(data, "actual", "Reduce 10%")
    report = build_report(data, answer)
    pdf = generate_direction_report_pdf(report)
    xlsx = generate_direction_report_xlsx(report)
    document = fitz.open(stream=pdf, filetype="pdf")
    text = "".join(page.get_text() for page in document)
    assert (
        data["snapshot_id"] in text and "$100.00 MXN" in text and "$90.00 MXN" in text
    )
    assert (
        "Riesgos" in text and "Definiciones" in text and "Cobranza desconocida" in text
    )
    book = load_workbook(io.BytesIO(xlsx), data_only=True)
    assert book["Consejo"]["B11"].value == 100
    assert book["Escenario"]["B6"].value == 10 and book["Escenario"]["B7"].value == 90
    assert (
        load_workbook(io.BytesIO(xlsx), data_only=False)["Escenario"]["B7"].value
        == "=B4-B6"
    )
    assert report["snapshot_id"] == answer["snapshot_id"]
    assert "None" not in text and "samchat." not in text
    # Rendered QA fixture; synthetic evidence, never business UAT.
    (tmp_path / "consejo.pdf").write_bytes(pdf)
    (tmp_path / "consejo.xlsx").write_bytes(xlsx)


def test_missing_stays_empty_and_untrusted_labels_cannot_be_formulas():
    data = snapshot([row(actual=None)])
    data["tournaments"][0]["name"] = '=HYPERLINK("https://evil.test","click")'
    report = build_report(data)
    book = load_workbook(
        io.BytesIO(generate_direction_report_xlsx(report)), data_only=False
    )
    assert book["Consejo"]["B11"].value is None
    assert book["Comparación"]["A2"].data_type == "s"
    assert book["Respaldo"]["A2"].data_type == "s"
    assert "<script>" not in render_report(
        {**report, "conclusion": "<script>alert(1)</script>"}
    )
    assert "<pre>" not in render_published_report("Legacy", {}, {}, "corte")


@pytest.mark.parametrize("format", ["pdf", "xlsx"])
def test_export_authorization_and_signed_analysis(context_client, monkeypatch, format):
    client, page = context_client
    answer = ask(client, page, question="Reduce 10%").json()
    body = {"context_token": page["token"], "analysis_token": answer["analysis_token"]}
    url = "/direccion/reportes/exportar/" + format
    response = client.post(url, json=body, headers={"X-Direction-CSRF": page["csrf"]})
    assert response.status_code == 200
    assert response.headers["x-direction-snapshot"] == page["snapshot"]["snapshot_id"]
    assert response.headers["cache-control"] == "no-store"
    assert response.content.startswith(b"%PDF" if format == "pdf" else b"PK")
    assert client.post(url, json=body).status_code == 403
    body["analysis_token"] += "x"
    assert (
        client.post(
            url, json=body, headers={"X-Direction-CSRF": page["csrf"]}
        ).status_code
        == 409
    )
    current = scope()
    current["selected"] = current["selected"][:1]
    monkeypatch.setattr(routes, "resolve_scope", AsyncMock(return_value=current))
    assert (
        client.post(
            url,
            json={"context_token": page["token"]},
            headers={"X-Direction-CSRF": page["csrf"]},
        ).status_code
        == 409
    )


def test_export_source_permission_revoked_and_scenario_validation(
    context_client, monkeypatch
):
    client, page = context_client
    assert (
        ask(client, page, scenario={"kind": "payment_delay", "days": True}).status_code
        == 422
    )
    answer = ask(
        client,
        page,
        scenario={
            "kind": "expense_reduction",
            "percent": "10",
            "tournament_id": "foreign",
        },
    ).json()
    assert answer["scenario"] is None and answer["missing_evidence"]
    monkeypatch.setattr(
        routes,
        "_direction_source_access",
        AsyncMock(return_value={"budget": True, "finance": False}),
    )
    response = client.post(
        "/direccion/reportes/exportar/pdf",
        json={"context_token": page["token"]},
        headers={"X-Direction-CSRF": page["csrf"]},
    )
    assert response.status_code == 409


def test_complete_concepts_opt_in_preserves_existing_six_row_contract():
    from samchat.budgets.service import _build_budget_line_breakdowns

    lines = [{"concept_name": f"Concepto {i}", "budget_amount": 10} for i in range(9)]
    assert len(_build_budget_line_breakdowns(lines)["by_concept"]) == 6
    assert (
        len(_build_budget_line_breakdowns(lines, concept_limit=None)["by_concept"]) == 9
    )


def test_legacy_report_renders_persisted_business_facts_only():
    html = render_published_report(
        "Informe",
        {"message": "Revisión"},
        {
            "cards": [
                {
                    "tournament_name": "Torneo A",
                    "budget": 100,
                    "actual": 40,
                    "private": "secret",
                }
            ]
        },
        "2026-09-29",
    )
    assert "Torneo A" in html and "$100.00 MXN" in html and "$40.00 MXN" in html
    assert "secret" not in html and "<pre>" not in html


@pytest.mark.asyncio
async def test_renamed_and_same_named_concepts_reconcile_by_canonical_id():
    from test_direction_home import Session

    from samchat.budgets.service import (
        _build_budget_finance_breakdowns,
        _build_budget_line_breakdowns,
        _finalize_breakdown_store,
        _merge_breakdown_row,
    )

    lines = [
        {
            "budget_concept_id": "canonical-1",
            "concept_name": "Nombre viejo",
            "budget_amount": 100,
        },
        {
            "budget_concept_id": "canonical-2",
            "concept_name": "Nombre nuevo",
            "budget_amount": 50,
        },
    ]
    session = Session(
        [
            [],
            [],
            [],
            [
                {
                    "concept_id": "canonical-1",
                    "label": "Nombre nuevo",
                    "actual_total": 40,
                },
                {
                    "concept_id": "canonical-2",
                    "label": "Nombre nuevo",
                    "actual_total": 20,
                },
            ],
        ]
    )
    budget = _build_budget_line_breakdowns(
        lines, concept_limit=None, concept_identity=True
    )
    finance = await _build_budget_finance_breakdowns(
        session,
        edition_year=2026,
        tournament_id=T1,
        tournament_name="Torneo",
        tournament_code=None,
        concept_limit=None,
        concept_identity=True,
    )
    store = {}
    for item in budget["by_concept"] + finance["by_concept"]:
        _merge_breakdown_row(store, **item)
    source = {"executive_concepts": _finalize_breakdown_store(store, limit=None)}
    evidence = home.concept_evidence(
        source, {"budget": Decimal("150"), "actual": Decimal("60")}
    )
    assert evidence["status"] == "available" and len(evidence["rows"]) == 2
    assert {r["id"] for r in evidence["rows"]} == {"canonical-1", "canonical-2"}
    assert all(Decimal(r["excess"]) < 0 for r in evidence["rows"])
    assert all("bc.id" in str(call[0]) for call in session.calls[2:])
    source["executive_concepts"].append(
        {"label": "Sin partida", "budget_total": "0", "actual_total": "0"}
    )
    assert (
        home.concept_evidence(
            source, {"budget": Decimal("150"), "actual": Decimal("60")}
        )["status"]
        == "available"
    )
    source["executive_concepts"][0].pop("concept_id")
    assert (
        home.concept_evidence(
            source, {"budget": Decimal("150"), "actual": Decimal("60")}
        )["status"]
        == "unavailable"
    )


def test_scenario_edits_are_allowed_but_operational_writes_remain_denied():
    data = snapshot()
    previous = answer_snapshot(data, "actual", "Reduce 10%")
    revised = answer_snapshot(data, "actual", "Modifica a 5%", previous=previous)
    assert revised["supported"] and revised["scenario"]["result"] == "95.00"
    for question in ("Modifica la factura a 5%", "Paga 5%", "Aprueba 5%"):
        result = answer_snapshot(data, "actual", question, previous=previous)
        assert not result["supported"] and result["scenario"] is None


def test_payment_delay_export_has_days_and_no_invented_percentage():
    data = snapshot([row()])
    r = data["tournaments"][0]
    r["values"]["obligations"] = "50"
    r["payment_evidence"] = [
        {"date": "2026-10-20", "value": "50", "reference": "REQ-1"}
    ]
    answer = answer_snapshot(data, "obligations", "Difiere 15 días")
    report = build_report(data, answer)
    book = load_workbook(
        io.BytesIO(generate_direction_report_xlsx(report)), data_only=True
    )
    assert book["Escenario"]["B5"].value is None
    assert book["Escenario"]["B10"].value == 15
    assert book["Escenario"]["B6"].value == 50


def test_legacy_forecast_alerts_and_safe_operational_fields_remain_visible():
    html = render_published_report(
        "Informe",
        {},
        {
            "cards": [
                {
                    "tournament_name": "A",
                    "projected": 110,
                    "available": 12,
                    "requested": 80,
                    "pending_to_pay": 15,
                    "alerts": [
                        {"title": "Alerta material <script>", "severity": "high"}
                    ],
                    "secret": "private",
                }
            ]
        },
        "corte",
    )
    assert (
        "$110.00 MXN" in html
        and "$12.00 MXN" in html
        and "$80.00 MXN" in html
        and "$15.00 MXN" in html
    )
    assert (
        "Alerta material &lt;script&gt;" in html
        and "high" in html
        and "private" not in html
    )
