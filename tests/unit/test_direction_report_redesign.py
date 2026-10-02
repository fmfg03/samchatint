"""Synthetic-only regression evidence for documentary facts and report contracts."""

import io
from datetime import date
from decimal import Decimal
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook
from test_direction_home import ACTOR, P1, T1, T2, Session, scope, snapshot

from samchat.budgets import executive_facts as facts
from samchat.client_executive import home
from samchat.client_executive.analysis import calculate_scenario
from samchat.client_executive.report_layouts import (
    GROUPS,
    budget_layout,
    cashflow_layout,
)
from samchat.client_executive.reports import build_report
from samchat.executive.exporter import (
    generate_direction_report_pdf,
    generate_direction_report_xlsx,
)

START, END = date(2026, 6, 1), date(2026, 6, 30)
PERIOD = {"start": str(START), "end": str(END)}


def expense(tid=T1, identity="e1", **overrides):
    return {
        "id": identity,
        "tournament_id": tid,
        "observed_date": START,
        "currency": "MXN",
        "gasto_cantidad": "116",
        "base_amount": "100",
        "shared_cfdi": False,
        **overrides,
    }


def document(tid=T1, identity="d1", **overrides):
    return {
        "id": identity,
        "tournament_id": tid,
        "observed_date": START,
        "currency": "MXN",
        "estado": "pagado",
        "monto_total": "75",
        "monto_solicitado": "100",
        **overrides,
    }


def project(expenses=(), documents=(), **kwargs):
    return facts.project_facts(
        [T1, T2], list(expenses), list(documents), start=START, end=END, **kwargs
    )["by_tournament"]


def test_documentary_stages_deduplicate_and_do_not_require_budget():
    data = project(
        [expense(), expense(), expense(T2, "e2", base_amount="25")],
        [document(), document()],
    )
    assert data[T1]["values"] == {
        "actual": "100.00",
        "committed": "75.00",
        "paid": "75.00",
    }
    assert data[T1]["counts"] == {"actual": 1, "committed": 1, "paid": 1}
    assert data[T2]["values"]["actual"] == "25.00"
    assert data[T2]["values"]["paid"] == "0"
    assert data[T1]["temporal_fields"]["paid"] == "documentos.creado_en"
    assert data[T1]["executive_monthly_actuals"] == [
        {"month": 6, "actual_total": "100.00"}
    ]


@pytest.mark.parametrize(
    "change,gap",
    [
        ({"shared_cfdi": True}, "shared_cfdi_allocation_requires_reconciliation"),
        ({"currency": "USD"}, "mixed_or_unknown_currency"),
        ({"gasto_cantidad": None}, "expense_amount_missing"),
        ({"base_amount": "NaN"}, "expense_amount_missing"),
        ({"observed_date": None}, "Fecha de origen ausente"),
    ],
)
def test_fact_quality_is_independent_by_source(change, gap):
    data = project([expense(**change)], [document()])[T1]
    assert data["values"]["actual"] is None
    assert any(gap in g for g in data["gaps"]["actual"])
    assert data["values"]["paid"] == "75.00"


@pytest.mark.parametrize(
    "change",
    [
        {"monto_total": None, "monto_solicitado": None},
        {"monto_total": "NaN"},
        {"currency": None},
        {"concepto_pago": "Reembolso de saldo a favor", "monto_total": None},
    ],
)
def test_missing_documentary_evidence_does_not_suppress_expenses(change):
    data = project([expense()], [document(**change)])[T1]
    assert data["values"]["actual"] == "100.00"
    assert data["values"]["paid"] is None


def test_outside_period_scope_conflicts_and_truncation():
    assert (
        project([expense(observed_date=date(2025, 1, 1))])[T1]["values"]["actual"]
        == "0"
    )
    for rows in (
        [expense("foreign")],
        [expense(identity="")],
        [expense(), expense(base_amount="5")],
    ):
        with pytest.raises(ValueError):
            project(rows)
    assert all(
        v is None for v in project([expense()], truncated=True)[T1]["values"].values()
    )
    paid = project(documents=[document(estado="aprobado")])[T1]["values"]
    assert paid["committed"] == "75.00" and paid["paid"] == "0"


@pytest.mark.asyncio
async def test_set_scoped_reader_uses_canonical_base_no_budget_gate_or_global_fallback():
    session = Session([[expense()], [document()]])
    data = await facts.build_executive_facts(
        session, tournament_ids=[T2, T1, T1], start=START, end=END
    )
    assert data["by_tournament"][T1]["values"]["actual"] == "100.00"
    assert len(session.calls) == 2
    expense_sql = str(session.calls[0][0])
    assert "report.id = e.informe_documento_id" in expense_sql
    assert "report.id = e.documento_id" in expense_sql
    assert "account_report.cuenta_gastos_id = e.cuenta_gastos_id" in expense_sql
    assert "SELECT COUNT(*)" in expense_sql
    assert session.calls[1][1]["committed_states"] == sorted(facts.COMMITTED)
    for statement, params in session.calls:
        assert "= ANY(CAST(:ids AS uuid[]))" in str(statement)
        assert params["ids"] == sorted([T1, T2])
        assert "budget_versions" not in str(statement)
        assert "budget_concept_id IS NOT NULL" not in str(statement)
        assert "LIMIT :limit" in str(statement)
    from samchat.budgets.service import _budget_expense_base_amount_sql

    assert _budget_expense_base_amount_sql("e", "cfdi") in str(session.calls[0][0])
    empty = Session()
    assert (
        await facts.build_executive_facts(
            empty, tournament_ids=[], start=START, end=END
        )
    )["by_tournament"] == {}
    assert not empty.calls
    with pytest.raises(ValueError):
        await facts.build_executive_facts(
            empty, tournament_ids=["bad"], start=START, end=END
        )
    with pytest.raises(ValueError):
        await facts.build_executive_facts(
            empty, tournament_ids=[T1], start=END, end=START
        )


@pytest.mark.asyncio
async def test_multi_scope_authorizes_every_uuid_without_widening(monkeypatch):
    monkeypatch.setattr(
        home.service, "authorized_direction_portfolio_ids", AsyncMock(return_value=[P1])
    )
    monkeypatch.setattr(
        home.service,
        "_authorized_tournaments",
        AsyncMock(return_value=scope()["tournaments"]),
    )
    args = dict(actor=ACTOR, superadmin=False, portfolio_id=None, tournament_id=None)
    selected = await home.resolve_scope(Session(), **args, tournament_ids=[T2, T1, T2])
    assert [t["id"] for t in selected["selected"]] == [T1, T2]
    assert selected["tournament_ids"] == [T1, T2]
    with pytest.raises(home.service.ClientExecutiveAccessError):
        await home.resolve_scope(Session(), **args, tournament_ids=[T1, "foreign"])
    with pytest.raises(ValueError):
        await home.resolve_scope(
            Session(), **{**args, "tournament_id": T1}, tournament_ids=[T2]
        )
    # Explicit portfolio selection is still a narrowing operation for SUPERADMIN.
    with pytest.raises(home.service.ClientExecutiveAccessError):
        await home.resolve_scope(
            Session([[], [{"id": T1}]]),
            **{**args, "superadmin": True, "portfolio_id": P1},
            tournament_ids=[T2],
        )


@pytest.mark.asyncio
async def test_draft_budget_preserves_documentary_facts_and_scenario(monkeypatch):
    selected = scope()
    selected["tournament_ids"] = [T1, T2]
    monkeypatch.setattr(home, "resolve_scope", AsyncMock(return_value=selected))

    async def read(session, loader, **kwargs):
        if loader is home.build_executive_facts:
            return facts.project_facts(
                [T1, T2],
                [expense(), expense(T2, "e2", base_amount="50")],
                [document()],
                start=kwargs["start"],
                end=kwargs["end"],
            )
        return {
            "source": "budget_scope_unavailable",
            "version": {"status": "draft"},
            "summary": {"budget_total": 999},
        }

    monkeypatch.setattr(home, "_optional_read", read)
    monkeypatch.setattr(
        home, "payment_values", AsyncMock(return_value={"value": None, "gaps": []})
    )
    monkeypatch.setattr(
        home.service, "_build_operational_dossier", AsyncMock(return_value={})
    )
    data, _ = await home.build_home(
        Session(), actor=ACTOR, superadmin=True, year=2026, start=START, end=END
    )
    metrics = {m["id"]: m for m in data["indicators"]}
    assert metrics["actual"]["value"] == "150.00"
    assert metrics["paid"]["value"] == "75.00"
    assert all(
        metrics[key]["value"] is None
        for key in ("budget", "forecast", "deviation", "liquidity")
    )
    assert metrics["actual"]["source"] == facts.SOURCE
    assert not any("no concilia" in gap for gap in metrics["forecast"]["gaps"])
    scenario = calculate_scenario(
        data,
        {"kind": "expense_reduction", "basis": "observed_expense", "percent": "10"},
    )
    assert scenario["result"] == "135.00"
    assert data["scope"]["tournament_ids"] == [T1, T2]


def test_budget_contract_all_rows_six_columns_no_double_count_and_favorability():
    cells = {
        "budget_month": "10",
        "budget_ytd": "20",
        "actual_month": "8",
        "actual_ytd": "15",
    }
    source = {f"group-{row}/part-{i}": cells for row, _ in GROUPS for i in (1, 2)}
    source.update(
        income={**cells, "budget_month": "300", "actual_month": "350"},
        broadcast=cells,
        sponsorship=cells,
    )
    layout = budget_layout(period=PERIOD, source_rows=source)
    assert [r["source_row"] for r in layout["rows"]] == list(range(6, 32))
    assert len(layout["columns"]) == 6
    rows = {r["id"]: r for r in layout["rows"]}
    assert rows["direct"]["values"]["budget_month"] == "140"
    assert rows["remainder"]["values"]["budget_month"] == "140"
    assert rows["income"]["values"]["variance_month"] == "-50"
    assert rows["income"]["favorability"]["month"] == "favorable"
    assert rows["broadcast"]["favorability"]["month"] == "favorable"
    blank = budget_layout(period=PERIOD)
    assert all(v is None for r in blank["rows"] for v in r["values"].values())
    assert all(r["percentages"]["month"] is None for r in blank["rows"])
    zero = budget_layout(
        period=PERIOD,
        source_rows={"income": {"budget_month": "0", "actual_month": "10"}},
    )
    assert zero["rows"][0]["percentages"]["month"] is None


def test_cashflow_inputs_periods_and_units_never_invent_zeros_or_previous_years():
    explicit = {
        "start": str(START),
        "end": str(END),
        "opening": "10000",
        "origins": ["1000"] * 6,
        "applications": ["500"] * 6,
    }
    layout = cashflow_layout(period=PERIOD, periods=[explicit])
    row = layout["periods"][0]
    assert row["opening"] == "10" and row["origins_total"] == "6"
    assert row["applications_total"] == "3" and row["closing"] == "13"
    blank = {**explicit, "origins": [None] * 6, "applications": [None] * 6}
    assert (
        cashflow_layout(period=PERIOD, periods=[blank])["periods"][0]["closing"] is None
    )
    assert cashflow_layout(period=PERIOD)["periods"] == []
    for change in ({"start": None}, {"end": "2025-01-01"}, {"origins": [0]}):
        with pytest.raises(ValueError):
            cashflow_layout(period=PERIOD, periods=[{**explicit, **change}])
    next_month = {
        **explicit,
        "start": "2026-07-01",
        "end": "2026-07-31",
        "opening": None,
    }
    assert (
        cashflow_layout(period=PERIOD, periods=[explicit, next_month])["periods"][1][
            "opening"
        ]
        == "13"
    )
    with pytest.raises(ValueError):
        cashflow_layout(
            period=PERIOD, periods=[explicit, {**next_month, "opening": "10"}]
        )
    with pytest.raises(ValueError):
        cashflow_layout(period=PERIOD, periods=[explicit] * 7)


def test_truncation_only_invalidates_its_own_source():
    data = project([expense()], [document()], truncated={"expense"})[T1]
    assert data["values"]["actual"] is None and data["values"]["paid"] == "75.00"
    data = project([expense()], [document()], truncated={"document"})[T1]
    assert data["values"]["actual"] == "100.00" and data["values"]["paid"] is None


@pytest.mark.asyncio
async def test_superadmin_read_does_not_need_persistent_portfolio_assignment(
    monkeypatch,
):
    monkeypatch.setattr(
        home.service, "authorized_direction_portfolio_ids", AsyncMock(return_value=[])
    )
    monkeypatch.setattr(
        home.service,
        "_authorized_tournaments",
        AsyncMock(return_value=scope()["tournaments"]),
    )
    selected = await home.resolve_scope(
        Session(), actor=ACTOR, superadmin=True, portfolio_id=None, tournament_id=None
    )
    assert len(selected["selected"]) == 2 and selected["portfolio_ids"] == []
    with pytest.raises(home.service.ClientExecutiveAccessError):
        await home.resolve_scope(
            Session(),
            actor=ACTOR,
            superadmin=False,
            portfolio_id=None,
            tournament_id=None,
        )


def test_report_exports_keep_layout_blanks_scope_and_cut():
    data = snapshot()
    report = build_report(data)
    wb = load_workbook(
        io.BytesIO(generate_direction_report_xlsx(report)), data_only=True
    )
    sheet = wb["Presupuesto vs Real"]
    assert sheet["A6"].value == "Ingresos" and sheet["A31"].value == "Remanente"
    assert all(
        sheet[f"{col}{row}"].value is None
        for col in ("B", "C", "E", "F", "H", "I")
        for row in range(6, 32)
    )
    assert wb["EFE"]["B22"].value is None
    assert report["cut"] in sheet["A2"].value
    assert report["scope"] == sheet["A1"].value
    pdf = generate_direction_report_pdf(report)
    assert pdf.startswith(b"%PDF")
    from pypdf import PdfReader

    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)
    assert "Presupuesto vs Real" in text and "Flujo de efectivo" in text
    assert data["snapshot_id"] in text


def test_sam_report_cell_reads_only_the_signed_cell_and_rejects_foreign_keys():
    from samchat.client_executive.conversation import ContextError, answer_snapshot

    data = snapshot()
    data["reports"]["budget"] = budget_layout(
        period=PERIOD,
        source_rows={"income": {"actual_month": "123", "budget_month": "100"}},
    )
    selection = {"report": "budget", "row": "income", "column": "actual_month"}
    answer = answer_snapshot(
        data, "actual", "Explica esta cifra", report_cell=selection
    )
    assert answer["metric"]["value"] == "123"
    assert "123.00 MXN" in answer["assistant_message"]
    assert answer["snapshot_id"] == data["snapshot_id"]
    assert answer["scenario"] is None
    for bad in (
        {**selection, "row": "foreign"},
        {**selection, "column": "foreign"},
        {**selection, "report": "foreign"},
        {"report": "cashflow", "column": "opening", "period": "-1"},
    ):
        with pytest.raises(ContextError):
            answer_snapshot(data, "actual", "Fuente", report_cell=bad)
    with pytest.raises(ContextError):
        answer_snapshot(
            data,
            "actual",
            "Escenario",
            report_cell=selection,
            scenario={"kind": "expense_reduction", "percent": "10"},
        )
    entry = {
        "start": str(START),
        "end": str(END),
        "opening": "1000",
        "origins": ["100"] * 6,
        "applications": ["50"] * 6,
    }
    data["reports"]["cashflow"] = cashflow_layout(period=PERIOD, periods=[entry])
    reply = answer_snapshot(
        data,
        "liquidity",
        "Fuente",
        report_cell={"report": "cashflow", "column": "closing", "period": "0"},
    )
    assert reply["metric"]["value"] == "1.3"
    assert "Miles de MXN" in reply["assistant_message"]
    from samchat.client_executive.home_ui import render_home

    html = render_home(data, scope(), token="synthetic", csrf="synthetic")
    assert 'data-report="budget"' in html and 'data-report="cashflow"' in html
    assert 'data-display="1.30 Miles de MXN"' in html
    report = build_report(data)
    wb = load_workbook(
        io.BytesIO(generate_direction_report_xlsx(report)), data_only=True
    )
    assert wb["Presupuesto vs Real"]["E6"].value == 123
    assert wb["EFE"]["B22"].value == 1.3
    assert generate_direction_report_pdf(report).startswith(b"%PDF")


@pytest.mark.asyncio
async def test_scope_counts_distinguish_portfolio_filter_from_accessible_universe(
    monkeypatch,
):
    from samchat.client_executive.home_ui import render_home

    monkeypatch.setattr(
        home.service, "authorized_direction_portfolio_ids", AsyncMock(return_value=[P1])
    )
    monkeypatch.setattr(
        home.service,
        "_authorized_tournaments",
        AsyncMock(return_value=scope()["tournaments"]),
    )
    selected = await home.resolve_scope(
        Session([[{"id": P1, "label": "Cartera sintética"}], [{"id": T1}]]),
        actor=ACTOR,
        superadmin=True,
        portfolio_id=P1,
        tournament_id=None,
    )
    assert selected["accessible_tournament_count"] == 2
    assert len(selected["tournaments"]) == len(selected["selected"]) == 1
    data = snapshot()
    data["tournaments"] = data["tournaments"][:1]
    html = render_home(data, selected, token="synthetic", csrf="synthetic")
    assert "1 seleccionados · 1 en el filtro · 2 accesibles" in html
    assert "Todos los torneos activos de la instalación" in html
    assert "Todos los 1 de este filtro" in html


@pytest.mark.parametrize(
    "state", ["pagado", "cerrado", "reembolsado", "aplicado", "liquidado"]
)
def test_terminal_document_without_paid_date_counts_both_stages(state):
    data = project(documents=[document(estado=state, pagado_en=None)])[T1]
    assert data["values"]["committed"] == "75.00"
    assert data["values"]["paid"] == "75.00"


@pytest.mark.parametrize("state", ["rechazado", "cancelado"])
def test_paid_evidence_survives_stale_state(state):
    data = project(documents=[document(estado=state, pagado_en="2026-06-15")])[T1]
    assert data["values"]["paid"] == "75.00"
    assert data["values"]["committed"] == "75.00"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "expenses,documents", [(False, True), (True, False), (False, False)]
)
async def test_documentary_source_denials_skip_owned_queries(expenses, documents):
    rows = ([[expense()]] if expenses else []) + ([[document()]] if documents else [])
    session = Session(rows)
    data = await facts.build_executive_facts(
        session,
        tournament_ids=[T1],
        start=START,
        end=END,
        include_expenses=expenses,
        include_documents=documents,
    )
    assert len(session.calls) == int(expenses) + int(documents)
    values = data["by_tournament"][T1]["values"]
    assert values["actual"] == ("100.00" if expenses else None)
    assert values["paid"] == ("75.00" if documents else None)
    if not expenses:
        assert all("FROM expense_reports e" not in str(q) for q, _ in session.calls)
    if not documents:
        assert all("d.tipo = 'SOLICITUD'" not in str(q) for q, _ in session.calls)


@pytest.mark.parametrize(
    "question", ["aprueba este presupuesto", "paga esto", "elimina el documento"]
)
def test_report_cell_preserves_operational_write_guard(question):
    from samchat.client_executive.conversation import answer_snapshot

    data = snapshot()
    result = answer_snapshot(
        data,
        "actual",
        question,
        report_cell={"report": "budget", "row": "income", "column": "actual_month"},
    )
    assert result["supported"] is False
    assert "No se ejecutó ninguna acción" in result["assistant_message"]
    assert result["read_only"] is True


@pytest.mark.asyncio
async def test_home_keeps_finance_documents_when_budget_is_denied(monkeypatch):
    monkeypatch.setattr(home, "resolve_scope", AsyncMock(return_value=scope()))

    async def read(session, loader, **kwargs):
        assert loader is home.build_executive_facts
        assert kwargs["include_expenses"] is False
        assert kwargs["include_documents"] is True
        return await loader(session, **kwargs)

    monkeypatch.setattr(home, "_optional_read", read)
    monkeypatch.setattr(home, "payment_values", AsyncMock(return_value={}))
    monkeypatch.setattr(
        home.service, "_build_operational_dossier", AsyncMock(return_value={})
    )
    session = Session([[document()]])
    data, _ = await home.build_home(
        session,
        actor=ACTOR,
        superadmin=True,
        year=2026,
        start=date(2026, 1, 1),
        end=END,
        source_access={"budget": False, "finance": True},
    )
    values = {m["id"]: m["value"] for m in data["indicators"]}
    assert values["committed"] == "75.00" and values["paid"] == "75.00"
    assert values["actual"] is None and values["budget"] is None
    assert len(session.calls) == 1


def test_submitted_request_is_committed_but_not_paid():
    data = project(documents=[document(estado="enviado", pagado_en=None)])[T1]
    assert data["values"]["committed"] == "75.00"
    assert data["values"]["paid"] == "0"
    assert "enviado" in facts.COMMITTED
