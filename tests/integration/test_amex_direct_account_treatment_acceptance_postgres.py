"""Independent financial acceptance of the three explicit direct-AMEX routes.

Only the external SAT transport is synthetic. Decisions, CP, fiscal owners,
cut, journals, exact owner DDL and frozen exports run against real PostgreSQL.
"""

import asyncio
import copy
import csv
import importlib.util
import io
import json
import os
from collections import defaultdict
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from devnous.gastos import models
from devnous.gastos.services import amex_direct_account_treatment_service as treatment
from devnous.gastos.services import amex_direct_accounting_service as accounting
from devnous.gastos.services import amex_direct_budget_service as budget
from devnous.gastos.services import amex_direct_no_cfdi_service as fiscal
from devnous.gastos.services.amex_direct_common import DirectAccessError, DirectConflict

_spec = importlib.util.spec_from_file_location(
    "independent_treatment_fixture",
    Path(__file__).resolve().parents[1]
    / "acceptance/test_amex_sat_budget_direct_postgres.py",
)
fixture = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(fixture)
direct_pg = fixture.direct_pg


@pytest.mark.asyncio
async def test_acceptance_delegated_legacy_classifier_cannot_review_financial_treatment(
    direct_pg,
):
    from devnous.gastos.services import amex_direct_common as common

    pg = direct_pg
    await setup_policy(pg)
    view, _ = await import_rows(
        pg,
        "2026-10-02,FINANCIAL REVIEW,-116.00\n"
        "2026-10-03,DELEGATED LEGACY PROJECT,-232.00\n",
    )
    item_id = next(
        i.id for i in view.items if i.source_json["description"] == "FINANCIAL REVIEW"
    )
    legacy = next(
        i
        for i in view.items
        if i.source_json["description"] == "DELEGATED LEGACY PROJECT"
    )
    request = decision_request(pg, view, [item_id], "project_expense")
    preview = await plan(pg, request)
    replay_key = "financial-before-role-change-" + uuid4().hex
    applied = await apply(pg, request, preview, replay_key)
    assert applied.status == "ready", applied.reason
    view = await current(pg, view.statement.id)
    current_request = decision_request(pg, view, [item_id], "project_expense")
    async with pg.factory.begin() as session:
        actor = await session.get(models.Empleado, pg.actor.id)
        actor.rol = "usuario"

    delegated_actions = []

    async def explicit_delegation(session, actor, tool, action):
        assert actor.id == pg.actor.id and tool == "amex_direct"
        delegated_actions.append(action)
        return True

    pg.monkeypatch.setattr(common, "explicit_tool_decision", explicit_delegation)
    for operation in ("options", "preview", "confirm", "replay"):
        async with pg.factory.begin() as session:
            with pytest.raises(DirectAccessError, match="finance_actor_required"):
                if operation == "options":
                    await treatment.list_direct_account_treatment_options(
                        session, actor_id=pg.actor.id, import_id=view.statement.id
                    )
                elif operation == "preview":
                    await treatment.preview_direct_account_treatment(
                        session, **current_request
                    )
                else:
                    await treatment.confirm_direct_account_treatment(
                        session,
                        **(request if operation == "replay" else current_request),
                        expected_preview_hash=preview.progress["preview_hash"],
                        idempotency_key=(
                            replay_key if operation == "replay" else "denied-new-review"
                        ),
                        confirmed=True,
                    )
    assert (
        await pg.conn.fetchval(
            "SELECT count(*) FROM amex_direct_account_treatment_decisions"
        )
        == 1
    )
    assert await pg.conn.fetchval("SELECT count(*) FROM accounting_polizas") == 0
    async with pg.factory.begin() as session:
        classified = await budget.classify_direct_items(
            session,
            actor_id=pg.actor.id,
            import_id=view.statement.id,
            item_ids=[legacy.id],
            tournament_id=pg.tournament.id,
            project_code=None,
            budget_concept_id=pg.concept.id,
            edition_year=2026,
            expected_versions={str(legacy.id): legacy.source_version},
            reason="Explicitly delegated legacy project classification",
        )
    assert classified.status == "classified", classified.reason
    assert "clasificar" in delegated_actions
    row = await pg.conn.fetchrow(
        "SELECT tournament_id,budget_concept_id FROM amex_direct_items WHERE id=$1",
        legacy.id,
    )
    assert tuple(row) == (pg.tournament.id, pg.concept.id)


async def setup_policy(pg):
    account = models.CuentaContable(
        id=uuid4(),
        codigo="1170-002-789",
        nombre="Acceptance partner identity",
        tipo="deudor",
        activo=True,
    )
    async with pg.factory.begin() as session:
        session.add(account)
    cfg = json.loads(os.environ["AMEX_DIRECT_AUTHORIZED_CONTEXTS"])
    for row in cfg:
        row["benefit_budget_line_ids"] = [str(pg.line_id)]
        row["partner_debtor_account_ids"] = [str(account.id)]
        for actor in row["actor_scopes"]:
            actor["benefit_budget_line_ids"] = [str(pg.line_id)]
            actor["partner_debtor_account_ids"] = [str(account.id)]
    pg.monkeypatch.setenv("AMEX_DIRECT_AUTHORIZED_CONTEXTS", json.dumps(cfg))
    await fixture.install_account_treatment_policy(pg)
    return account, cfg


async def current(pg, import_id):
    return await fixture._direct_current(pg, import_id)


async def import_rows(pg, rows):
    content = ("Date,Description,Amount\n" + rows).encode()
    result = await fixture._direct_import(pg, content)
    assert result.status in {"created", "imported"}, result.reason
    return await current(pg, result.statement.id), content


def decision_request(pg, view, ids, kind, account=None, action="set"):
    chosen = [i for i in view.items if i.id in ids]
    return dict(
        actor_id=pg.actor.id,
        import_id=view.statement.id,
        item_ids=ids,
        treatment=kind,
        debtor_account_id=account.id if account else None,
        expected_versions={str(i.id): i.source_version for i in chosen},
        reason="Explicit independently reviewed acceptance destination",
        action=action,
    )


async def plan(pg, request):
    async with pg.factory() as session:
        result = await treatment.preview_direct_account_treatment(session, **request)
    assert result.status == "preview", result.reason
    return result


async def apply(pg, request, preview, key):
    async with pg.factory.begin() as session:
        return await treatment.confirm_direct_account_treatment(
            session,
            **request,
            expected_preview_hash=preview.progress["preview_hash"],
            idempotency_key=key,
            confirmed=True,
        )


async def choose(pg, view, ids, kind, account=None):
    request = decision_request(pg, view, ids, kind, account)
    result = await apply(
        pg, request, await plan(pg, request), "destination-" + uuid4().hex
    )
    assert result.status == "ready", result.reason
    return await current(pg, view.statement.id)


async def no_invoice(pg, view, ids):
    async with pg.factory.begin() as session:
        result = await fiscal.declare_no_cfdi_items(
            session,
            actor_id=pg.actor.id,
            import_id=view.statement.id,
            item_ids=ids,
            expected_versions={
                str(i.id): i.source_version for i in view.items if i.id in ids
            },
            expected_source_version=view.source_version,
            reason="The selected acceptance charge will have no CFDI",
            confirmed=True,
            request_key="no-invoice-" + uuid4().hex,
        )
    assert result.status == "ready", result.reason
    return await current(pg, view.statement.id)


async def cp_assign(pg, view, ids):
    async with pg.factory.begin() as session:
        sent = await budget.send_direct_items_to_budget(
            session,
            actor_id=pg.actor.id,
            import_id=view.statement.id,
            item_ids=ids,
            expected_versions={
                str(i.id): i.source_version for i in view.items if i.id in ids
            },
            confirmed=True,
        )
    assert sent.status == "sent", sent.reason
    view = await current(pg, view.statement.id)
    for row in view.items:
        if row.id not in ids:
            continue
        async with pg.factory.begin() as session:
            result = await budget.assign_direct_budget_line(
                session,
                actor_id=pg.cp_actor.id,
                item_id=row.id,
                budget_line_id=pg.line_id,
                expected_source_version=row.source_version,
                expected_source_hash=row.source_hash,
                idempotency_key="cp-" + str(row.id),
            )
        assert result.status == "assigned", result.reason
    return await current(pg, view.statement.id)


async def preview_cut(pg, view):
    async with pg.factory() as session:
        return await accounting.preview_direct_accounting(
            session,
            actor_id=pg.actor.id,
            import_id=view.statement.id,
        )


async def stable_evidence(pg, import_id):
    return {
        "header": await pg.conn.fetchval(
            "SELECT to_jsonb(t)::text FROM amex_statement_imports t WHERE id=$1",
            import_id,
        ),
        "receipts": [
            r["data"]
            for r in await pg.conn.fetch(
                "SELECT to_jsonb(t)::text data FROM amex_statement_import_items t WHERE import_id=$1 ORDER BY id",
                import_id,
            )
        ],
        "sources": [
            dict(r)
            for r in await pg.conn.fetch(
                "SELECT id,import_item_id,raw_signed_amount,raw_amount,profile,profile_version,source_json FROM amex_direct_items WHERE import_id=$1 ORDER BY id",
                import_id,
            )
        ],
        "original": [
            r["data"]
            for r in await pg.conn.fetch(
                "SELECT to_jsonb(t)::text data FROM amex_statement_source_files t WHERE import_id=$1 ORDER BY id",
                import_id,
            )
        ],
    }


@pytest.mark.asyncio
async def test_acceptance_mixed_CFDi_benefit_partner_and_project_one_exact_policy(
    direct_pg,
):
    from samchat.budgets.service import build_budget_actuals_snapshot

    pg = direct_pg
    debtor, _ = await setup_policy(pg)
    content = b"Date,Description,Amount\n2026-10-02,COMERCIO ACME CDMX,-116.00\n2026-10-03,PERSONAL SUBSCRIPTION,-232.00\n2026-10-04,PROJECT SUPPLIES,-348.00\n"
    matched, _, cfdi = await fixture._matched_statement(pg, content=content)
    view = await current(pg, matched.statement.id)
    before = await stable_evidence(pg, view.statement.id)
    ids = {i.source_json["description"]: i.id for i in view.items}
    benefit_id, partner_id, project_id = [
        ids[k]
        for k in ("COMERCIO ACME CDMX", "PERSONAL SUBSCRIPTION", "PROJECT SUPPLIES")
    ]
    fiscal_link = await pg.conn.fetchrow(
        "SELECT cfdi_id,status_receipt_id,match_proof FROM amex_direct_items WHERE id=$1",
        benefit_id,
    )

    view = await choose(pg, view, [benefit_id], "benefit_expense")
    # A genuine CFDI stays linked; the bank charge does not need a project.
    benefit = next(i for i in view.items if i.id == benefit_id)
    assert (
        benefit.cfdi_id == cfdi.id
        and benefit.tournament_id is None
        and benefit.project_code is None
    )
    assert benefit.budget_concept_id is None and benefit.budget_assignment_id is None
    blocked = await preview_cut(pg, view)
    assert (
        blocked.status == "pending"
        and "budget_assignment_required" in blocked.progress["exceptions"]
    )

    view = await choose(pg, view, [partner_id], "partner_receivable", debtor)
    view = await no_invoice(pg, view, [partner_id])
    view = await choose(pg, view, [project_id], "project_expense")
    view = await no_invoice(pg, view, [project_id])
    project = next(i for i in view.items if i.id == project_id)
    async with pg.factory.begin() as session:
        result = await budget.classify_direct_items(
            session,
            actor_id=pg.actor.id,
            import_id=view.statement.id,
            item_ids=[project_id],
            tournament_id=pg.tournament.id,
            project_code=None,
            budget_concept_id=pg.concept.id,
            edition_year=2026,
            expected_versions={str(project.id): project.source_version},
            reason="Actual project context",
        )
    assert result.status == "classified", result.reason
    view = await cp_assign(
        pg, await current(pg, view.statement.id), [benefit_id, project_id]
    )
    partner = next(i for i in view.items if i.id == partner_id)
    assert (
        partner.budget_status == "not_applicable"
        and partner.budget_assignment_id is None
    )
    row = await pg.conn.fetchrow(
        "SELECT i.tournament_id,i.project_code,e.proyecto,e.cuenta_contable_id,e.contra_cuenta_contable_id FROM amex_direct_items i JOIN amex_direct_effective_sources s ON s.id=i.id JOIN expense_reports e ON e.id=s.imported_expense_id WHERE i.id=$1",
        benefit_id,
    )
    assert tuple(row[:2]) == (None, None)
    assert not str(row["proyecto"] or "").strip()
    assert (
        row["cuenta_contable_id"] == pg.gasto.id
        and row["contra_cuenta_contable_id"] == pg.liability.id
    )

    args = await fixture._cut_args(pg, view)
    cut = await fixture._confirm(pg, args)
    assert cut.status == "recognized", cut.reason
    again = await fixture._confirm(pg, args)
    assert again.status in {"exists", "existing"} and again.cut.id == cut.cut.id
    assert await pg.conn.fetchval("SELECT count(*) FROM accounting_polizas") == 1
    assert await pg.conn.fetchval("SELECT count(*) FROM amex_direct_cut_items") == 3
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM amex_recognition_consumptions")
        == 3
    )
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM budget_movement_assignments") == 2
    )
    assert await pg.conn.fetchval("SELECT count(*) FROM documentos") == 0
    amounts = defaultdict(lambda: [Decimal(0), Decimal(0)])
    for line in await pg.conn.fetch(
        "SELECT cuenta_codigo,debe,haber FROM accounting_poliza_lines"
    ):
        amounts[line["cuenta_codigo"]][0] += Decimal(str(line["debe"] or 0))
        amounts[line["cuenta_codigo"]][1] += Decimal(str(line["haber"] or 0))
    expected_amounts = {
        pg.gasto.codigo: [Decimal("448.00"), Decimal(0)],
        pg.iva.codigo: [Decimal("16.00"), Decimal(0)],
        debtor.codigo: [Decimal("232.00"), Decimal(0)],
        pg.liability.codigo: [Decimal(0), Decimal("696.00")],
    }
    assert dict(amounts) == expected_amounts
    async with pg.factory() as session:
        actuals = await build_budget_actuals_snapshot(
            session,
            edition_year=2026,
            version_id=str(pg.version_id),
            tournament_id=str(pg.tournament.id),
            tournament_name=pg.tournament.name,
        )
    expenses = {str(p["expense_id"]) for p in cut.cut.snapshot_json["partidas"]}
    own = [m for m in actuals["movements"] if str(m.get("expense_id")) in expenses]
    assert len(own) == 2
    assert sum(Decimal(str(m["amount"])) for m in own) == Decimal("448.00")
    assert await stable_evidence(pg, view.statement.id) == before
    assert dict(
        await pg.conn.fetchrow(
            "SELECT cfdi_id,status_receipt_id,match_proof FROM amex_direct_items WHERE id=$1",
            benefit_id,
        )
    ) == dict(fiscal_link)

    exported = {}
    for fmt in ("xlsx", "csv"):
        async with pg.factory() as session:
            blob, filename, _ = await accounting.export_direct_cut(
                session, actor_id=pg.actor.id, import_id=view.statement.id, format=fmt
            )
        assert blob and filename.endswith("." + fmt)
        exported[fmt] = blob
    book = load_workbook(io.BytesIO(exported["xlsx"]))
    rows = list(book["Poliza COI"].values)
    fiscal_rows = [r for r in rows if r[8] == cfdi.cfdi_uuid]
    assert len(fiscal_rows) == 1 and fiscal_rows[0][5:7] == (
        "AAA010101AAA",
        "BBB010101BBB",
    )
    assert all(c.data_type != "f" for sheet in book for row in sheet for c in row)
    csv_rows = list(csv.reader(io.StringIO(exported["csv"].decode("latin-1"))))
    for exported_rows in (rows, csv_rows):
        assert sum(r[0] == "Eg" for r in exported_rows) == 1
        assert sum(r[1] == "FIN_PARTIDAS" for r in exported_rows) == 1
        totals = defaultdict(lambda: [Decimal(0), Decimal(0)])
        for r in exported_rows:
            if r[1] in expected_amounts:
                totals[r[1]][0] += Decimal(str(r[5] or 0))
                totals[r[1]][1] += Decimal(str(r[6] or 0))
        assert dict(totals) == expected_amounts
    csv_cfdi = [r for r in csv_rows if r[8] == cfdi.cfdi_uuid]
    assert len(csv_cfdi) == 1
    assert tuple(v.strip() for v in csv_cfdi[0][5:7]) == (
        "AAA010101AAA",
        "BBB010101BBB",
    )
    await pg.conn.execute(
        "UPDATE cuentas_contables SET nombre='Mutable later catalog label' WHERE id=$1",
        pg.gasto.id,
    )
    async with pg.factory() as session:
        frozen, _, _ = await accounting.export_direct_cut(
            session, actor_id=pg.actor.id, import_id=view.statement.id, format="xlsx"
        )
    assert list(load_workbook(io.BytesIO(frozen))["Poliza COI"].values) == rows
    async with pg.factory() as session:
        csv_frozen, _, _ = await accounting.export_direct_cut(
            session, actor_id=pg.actor.id, import_id=view.statement.id, format="csv"
        )
    assert csv_frozen == exported["csv"]
    recognized = await current(pg, view.statement.id)
    with pytest.raises(DirectConflict):
        await plan(
            pg,
            decision_request(
                pg,
                recognized,
                [partner_id],
                "partner_receivable",
                debtor,
                action="revoke",
            ),
        )
    assert await pg.conn.fetchval("SELECT count(*) FROM accounting_polizas") == 1


@pytest.mark.asyncio
async def test_acceptance_benefit_send_needs_no_project_but_only_CP_can_assign(
    direct_pg,
):
    pg = direct_pg
    await setup_policy(pg)
    view, _ = await import_rows(pg, "2026-10-02,BENEFIT,-116.00\n")
    item_id = view.items[0].id
    view = await choose(pg, view, [item_id], "benefit_expense")
    view = await no_invoice(pg, view, [item_id])
    row = view.items[0]
    async with pg.factory.begin() as session:
        sent = await budget.send_direct_items_to_budget(
            session,
            actor_id=pg.actor.id,
            import_id=view.statement.id,
            item_ids=[item_id],
            expected_versions={str(item_id): row.source_version},
            confirmed=True,
        )
    assert sent.status == "sent", sent.reason
    view = await current(pg, view.statement.id)
    async with pg.factory() as session:
        queued = await budget.list_direct_budget_queue(session, actor_id=pg.cp_actor.id)
    assert any(i.id == item_id for i in queued)
    row = view.items[0]
    async with pg.factory.begin() as session:
        with pytest.raises(DirectAccessError):
            await budget.assign_direct_budget_line(
                session,
                actor_id=pg.actor.id,
                item_id=item_id,
                budget_line_id=pg.line_id,
                expected_source_version=row.source_version,
                expected_source_hash=row.source_hash,
                idempotency_key="not-CP",
            )
    blocked = await preview_cut(pg, view)
    assert (
        blocked.status == "pending"
        and "budget_assignment_required" in blocked.progress["exceptions"]
    )
    assert await pg.conn.fetchval("SELECT count(*) FROM accounting_polizas") == 0
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM budget_movement_assignments") == 0
    )
    # The same source becomes eligible only after its real CP owner assigns it.
    async with pg.factory.begin() as session:
        assigned = await budget.assign_direct_budget_line(
            session,
            actor_id=pg.cp_actor.id,
            item_id=item_id,
            budget_line_id=pg.line_id,
            expected_source_version=row.source_version,
            expected_source_hash=row.source_hash,
            idempotency_key="actual-CP-assignment",
        )
    assert assigned.status == "assigned", assigned.reason
    ready = await current(pg, view.statement.id)
    recognized = await fixture._confirm(pg, await fixture._cut_args(pg, ready))
    assert recognized.status == "recognized", recognized.reason
    absence = await pg.conn.fetchrow(
        "SELECT i.tournament_id,i.project_code,e.proyecto FROM amex_direct_items i "
        "JOIN amex_direct_effective_sources s ON s.id=i.id "
        "JOIN expense_reports e ON e.id=s.imported_expense_id WHERE i.id=$1",
        item_id,
    )
    assert absence["tournament_id"] is None and absence["project_code"] is None
    assert not str(absence["proyecto"] or "").strip()
    lines = await pg.conn.fetch(
        "SELECT cuenta_codigo,debe,haber FROM accounting_poliza_lines"
    )
    assert len(lines) == 2
    assert {(r["cuenta_codigo"], r["debe"], r["haber"]) for r in lines} == {
        (pg.gasto.codigo, Decimal("116.00"), Decimal(0)),
        (pg.liability.codigo, Decimal(0), Decimal("116.00")),
    }
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM budget_movement_assignments") == 1
    )


@pytest.mark.asyncio
@pytest.mark.parametrize("bad_kind", ["ungranted", "inactive", "wrong_type"])
async def test_acceptance_invalid_partner_account_never_mutates_source(
    direct_pg, bad_kind
):
    pg = direct_pg
    debtor, _ = await setup_policy(pg)
    view, _ = await import_rows(pg, "2026-10-02,PERSONAL,-232.00\n")
    if bad_kind == "ungranted":
        other = models.CuentaContable(
            id=uuid4(),
            codigo="1170-002-788",
            nombre="Outside explicit grant",
            tipo="deudor",
            activo=True,
        )
        async with pg.factory.begin() as session:
            session.add(other)
        debtor = other
        error = DirectAccessError
    elif bad_kind == "inactive":
        await pg.conn.execute(
            "UPDATE cuentas_contables SET activo=false WHERE id=$1", debtor.id
        )
        error = DirectConflict
    else:
        await pg.conn.execute(
            "UPDATE cuentas_contables SET codigo='5300-999-999' WHERE id=$1", debtor.id
        )
        error = DirectConflict
    with pytest.raises(error):
        await plan(
            pg,
            decision_request(
                pg, view, [view.items[0].id], "partner_receivable", debtor
            ),
        )
    assert (
        await pg.conn.fetchval(
            "SELECT count(*) FROM amex_direct_account_treatment_decisions"
        )
        == 0
    )
    assert await pg.conn.fetchval("SELECT count(*) FROM accounting_polizas") == 0
    assert (
        await pg.conn.fetchval(
            "SELECT count(*) FROM expense_reports WHERE cuenta_contable_id IS NOT NULL"
        )
        == 0
    )


@pytest.mark.asyncio
async def test_acceptance_late_stale_batch_is_atomic(direct_pg):
    pg = direct_pg
    debtor, _ = await setup_policy(pg)
    view, _ = await import_rows(
        pg, "2026-10-02,FIRST,-116.00\n2026-10-03,SECOND,-232.00\n"
    )
    ids = [i.id for i in view.items]
    request = decision_request(pg, view, ids, "benefit_expense")
    before = await plan(pg, request)
    # Another legitimate decision changes only the last selected row after preview.
    view = await choose(pg, view, [ids[-1]], "partner_receivable", debtor)
    with pytest.raises(DirectConflict):
        await apply(pg, request, before, "old-entire-batch")
    assert (
        await pg.conn.fetchval(
            "SELECT count(*) FROM amex_direct_account_treatment_decisions"
        )
        == 1
    )
    assert (
        await pg.conn.fetchval(
            "SELECT account_treatment_decision_id FROM amex_direct_items WHERE id=$1",
            ids[0],
        )
        is None
    )
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM budget_movement_assignments") == 0
    )
    assert await pg.conn.fetchval("SELECT count(*) FROM accounting_polizas") == 0


@pytest.mark.asyncio
@pytest.mark.parametrize("same_key", [True, False])
async def test_acceptance_concurrent_confirmations_have_one_financial_destination(
    direct_pg, same_key
):
    pg = direct_pg
    debtor, _ = await setup_policy(pg)
    view, _ = await import_rows(pg, "2026-10-02,PERSONAL,-232.00\n")
    request = decision_request(
        pg, view, [view.items[0].id], "partner_receivable", debtor
    )
    planned = await plan(pg, request)
    results = await asyncio.gather(
        apply(pg, request, planned, "race-first"),
        apply(pg, request, planned, "race-first" if same_key else "race-second"),
        return_exceptions=True,
    )
    successful = [r for r in results if not isinstance(r, Exception)]
    assert any(r.status == "ready" for r in successful)
    if same_key:
        assert sorted(r.status for r in successful) == ["exists", "ready"]
    else:
        assert len(successful) == 1 and any(
            isinstance(r, DirectConflict) for r in results
        )
    assert (
        await pg.conn.fetchval(
            "SELECT count(*) FROM amex_direct_account_treatment_decisions"
        )
        == 1
    )
    assert await pg.conn.fetchval("SELECT count(*) FROM accounting_polizas") == 0
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM budget_movement_assignments") == 0
    )


@pytest.mark.asyncio
async def test_acceptance_grant_upgrade_preserves_links_and_old_identity_changes_fail(
    direct_pg,
):
    pg = direct_pg
    matched, _, _ = await fixture._matched_statement(pg)
    view = await current(pg, matched.statement.id)
    context_before = await pg.conn.fetchval(
        "SELECT to_jsonb(c)::text FROM amex_statement_direct_contexts c"
    )
    fiscal_before = await pg.conn.fetchrow(
        "SELECT cfdi_id,status_receipt_id,match_proof FROM amex_direct_items"
    )
    old_events = {
        str(r["id"]): r["data"]
        for r in await pg.conn.fetch(
            "SELECT id,to_jsonb(e)::text data FROM amex_direct_events e"
        )
    }
    await setup_policy(pg)
    view = await current(pg, matched.statement.id)
    view = await choose(pg, view, [view.items[0].id], "benefit_expense")
    assert (
        await pg.conn.fetchval(
            "SELECT to_jsonb(c)::text FROM amex_statement_direct_contexts c"
        )
        == context_before
    )
    assert dict(
        await pg.conn.fetchrow(
            "SELECT cfdi_id,status_receipt_id,match_proof FROM amex_direct_items"
        )
    ) == dict(fiscal_before)
    config = json.loads(os.environ["AMEX_DIRECT_AUTHORIZED_CONTEXTS"])
    for mutation in ("receiver", "actor_scope", "card_scope"):
        altered = copy.deepcopy(config)
        if mutation == "receiver":
            altered[0]["receiver_rfc"] = "CCC010101CCC"
        elif mutation == "actor_scope":
            altered[0]["actor_scopes"][0]["tournament_ids"] = []
        else:
            altered[0]["card_account_ids"] = [str(uuid4())]
        pg.monkeypatch.setenv("AMEX_DIRECT_AUTHORIZED_CONTEXTS", json.dumps(altered))
        with pytest.raises((DirectConflict, DirectAccessError)):
            await current(pg, matched.statement.id)
    pg.monkeypatch.setenv("AMEX_DIRECT_AUTHORIZED_CONTEXTS", json.dumps(config))
    assert (
        await pg.conn.fetchval(
            "SELECT to_jsonb(c)::text FROM amex_statement_direct_contexts c"
        )
        == context_before
    )
    assert await pg.conn.fetchval("SELECT count(*) FROM accounting_polizas") == 0

    after_events = {
        str(r["id"]): r["data"]
        for r in await pg.conn.fetch(
            "SELECT id,to_jsonb(e)::text data FROM amex_direct_events e"
        )
    }
    assert all(after_events.get(key) == data for key, data in old_events.items())


@pytest.mark.asyncio
async def test_acceptance_removed_grant_blocks_cut_but_allows_explicit_revocation(
    direct_pg,
):
    pg = direct_pg
    debtor, config = await setup_policy(pg)
    view, _ = await import_rows(pg, "2026-10-02,PERSONAL,-232.00\n")
    view = await choose(pg, view, [view.items[0].id], "partner_receivable", debtor)
    view = await no_invoice(pg, view, [view.items[0].id])
    for row in config:
        row["partner_debtor_account_ids"] = []
        for actor in row["actor_scopes"]:
            actor["partner_debtor_account_ids"] = []
    pg.monkeypatch.setenv("AMEX_DIRECT_AUTHORIZED_CONTEXTS", json.dumps(config))
    await fixture.install_account_treatment_policy(pg)
    view = await current(pg, view.statement.id)
    blocked = await preview_cut(pg, view)
    assert blocked.status == "pending"
    assert set(blocked.progress["exceptions"]) & {
        "account_treatment_policy_changed",
        "account_treatment_source_changed",
    }
    # Fiscal declaration must be explicitly revoked before changing destination.
    row = view.items[0]
    async with pg.factory.begin() as session:
        result = await fiscal.revoke_no_cfdi_items(
            session,
            actor_id=pg.actor.id,
            import_id=view.statement.id,
            item_ids=[row.id],
            expected_versions={str(row.id): row.source_version},
            expected_source_version=view.source_version,
            reason="Withdrawn grant explicit fiscal release",
            request_key="release-invoice",
            confirmed=True,
        )
    assert result.status == "ready", result.reason
    view = await current(pg, view.statement.id)
    request = decision_request(
        pg, view, [row.id], "partner_receivable", debtor, action="revoke"
    )
    withdrawn = await apply(
        pg, request, await plan(pg, request), "withdraw-destination"
    )
    assert withdrawn.status == "ready", withdrawn.reason
    assert (
        await pg.conn.fetchval(
            "SELECT count(*) FROM amex_direct_account_treatment_decisions"
        )
        == 2
    )
    assert await pg.conn.fetchval("SELECT count(*) FROM accounting_polizas") == 0


@pytest.mark.asyncio
async def test_acceptance_corrected_sign_source_keeps_google_charge_and_amazon_credit(
    direct_pg,
):
    from devnous.gastos.services import amex_direct_sign_correction_service as signs
    from devnous.gastos.services.amex_direct_matching_service import (
        exclude_direct_credit,
    )

    pg = direct_pg
    debtor, _ = await setup_policy(pg)
    view, _ = await import_rows(
        pg,
        "2026-10-02,GOOGLE,939.00\n2026-10-03,BENEFIT,116.00\n2026-10-04,AMAZON REFUND,-46.53\n",
    )
    original = await stable_evidence(pg, view.statement.id)
    assert [i.movement_kind for i in view.items].count("debit") == 1
    request = dict(
        actor_id=pg.actor.id,
        import_id=view.statement.id,
        profile="debit_positive",
        reason="Explicit correction of source sign declaration",
        expected_source_version=view.source_version,
    )
    async with pg.factory() as session:
        planned = await signs.preview_direct_sign_correction(session, **request)
    assert planned.status == "preview", planned.reason
    async with pg.factory.begin() as session:
        fixed = await signs.confirm_direct_sign_correction(
            session,
            **request,
            preview_hash=planned.progress["preview_hash"],
            idempotency_key="sign-before-destinations",
            confirmed=True,
        )
    assert fixed.status == "corrected", fixed.reason
    view = await current(pg, view.statement.id)
    by_name = {i.descriptor: i for i in view.items}
    assert by_name["GOOGLE"].movement_kind == "debit" and by_name[
        "GOOGLE"
    ].amount == Decimal("939")
    assert (
        by_name["AMAZON REFUND"].movement_kind == "credit"
        and by_name["AMAZON REFUND"].imported_expense_id is None
    )
    google, benefit, refund = [
        by_name[k].id for k in ("GOOGLE", "BENEFIT", "AMAZON REFUND")
    ]
    view = await choose(pg, view, [google], "partner_receivable", debtor)
    view = await no_invoice(pg, view, [google])
    view = await choose(pg, view, [benefit], "benefit_expense")
    view = await no_invoice(pg, view, [benefit])
    async with pg.factory.begin() as session:
        result = await exclude_direct_credit(
            session,
            actor_id=pg.actor.id,
            import_id=view.statement.id,
            item_id=refund,
            expected_source_version=view.source_version,
            reason="Actual refund is not a new expense",
            confirmed=True,
        )
    assert result.status == "ready", result.reason
    view = await cp_assign(pg, await current(pg, view.statement.id), [benefit])
    cut = await fixture._confirm(pg, await fixture._cut_args(pg, view))
    assert cut.status == "recognized", cut.reason
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM amex_direct_sign_revisions") == 1
    )
    assert await pg.conn.fetchval("SELECT count(*) FROM amex_direct_sign_rows") == 3
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM budget_movement_assignments") == 1
    )
    assert await stable_evidence(pg, view.statement.id) == original
    lines = await pg.conn.fetch(
        "SELECT cuenta_codigo,debe,haber FROM accounting_poliza_lines"
    )
    assert (
        sum(r["debe"] for r in lines)
        == sum(r["haber"] for r in lines)
        == Decimal("1055.00")
    )
    assert sum(
        r["debe"] for r in lines if r["cuenta_codigo"] == debtor.codigo
    ) == Decimal("939.00")
    assert sum(
        r["debe"] for r in lines if r["cuenta_codigo"] == pg.gasto.codigo
    ) == Decimal("116.00")
    assert len(cut.cut.snapshot_json["partidas"]) == 3


@pytest.mark.asyncio
async def test_acceptance_partner_WITH_CFDi_posts_only_full_asset_and_card(direct_pg):
    pg = direct_pg
    debtor, _ = await setup_policy(pg)
    matched, _, cfdi = await fixture._matched_statement(pg)
    view = await current(pg, matched.statement.id)
    original = await stable_evidence(pg, view.statement.id)
    link = await pg.conn.fetchrow(
        "SELECT cfdi_id,status_receipt_id,match_proof FROM amex_direct_items"
    )
    item_id = view.items[0].id
    view = await choose(pg, view, [item_id], "partner_receivable", debtor)
    cut = await fixture._confirm(pg, await fixture._cut_args(pg, view))
    assert cut.status == "recognized", cut.reason
    lines = await pg.conn.fetch(
        "SELECT cuenta_codigo,debe,haber FROM accounting_poliza_lines"
    )
    assert len(lines) == 2
    assert {(r["cuenta_codigo"], r["debe"], r["haber"]) for r in lines} == {
        (debtor.codigo, Decimal("116.00"), Decimal(0)),
        (pg.liability.codigo, Decimal(0), Decimal("116.00")),
    }
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM budget_movement_assignments") == 0
    )
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM amex_direct_no_cfdi_decisions")
        == 0
    )
    assert await stable_evidence(pg, view.statement.id) == original
    assert dict(
        await pg.conn.fetchrow(
            "SELECT cfdi_id,status_receipt_id,match_proof FROM amex_direct_items"
        )
    ) == dict(link)
    part = cut.cut.snapshot_json["partidas"][0]
    assert part["cfdi"]["uuid"] == cfdi.cfdi_uuid
    assert part["no_cfdi_decision_id"] is None
    assert part["budget_assignment_id"] is None
    async with pg.factory() as session:
        blob, _, _ = await accounting.export_direct_cut(
            session, actor_id=pg.actor.id, import_id=view.statement.id, format="xlsx"
        )
    rows = list(load_workbook(io.BytesIO(blob))["Poliza COI"].values)
    assert any(r[8] == cfdi.cfdi_uuid for r in rows)


@pytest.mark.asyncio
async def test_acceptance_withdrawn_benefit_line_can_be_returned_but_not_reassigned(
    direct_pg,
):
    pg = direct_pg
    _, config = await setup_policy(pg)
    view, _ = await import_rows(pg, "2026-10-02,EXISTING BENEFIT,-116.00\n")
    item_id = view.items[0].id
    view = await choose(pg, view, [item_id], "benefit_expense")
    view = await no_invoice(pg, view, [item_id])
    view = await cp_assign(pg, view, [item_id])
    old_assignment = view.items[0].budget_assignment_id
    old_decision = view.items[0].account_treatment_decision_id
    for row in config:
        row["benefit_budget_line_ids"] = []
        row["benefit_budget_pairs"] = []
        for scope in row["actor_scopes"]:
            scope["benefit_budget_line_ids"] = []
            scope["benefit_budget_pairs"] = []
    pg.monkeypatch.setenv("AMEX_DIRECT_AUTHORIZED_CONTEXTS", json.dumps(config))
    await fixture.install_account_treatment_policy(pg)
    view = await current(pg, view.statement.id)
    blocked = await preview_cut(pg, view)
    assert blocked.status == "pending"
    assert await pg.conn.fetchval("SELECT count(*) FROM accounting_polizas") == 0
    async with pg.factory() as session:
        queue = await budget.list_direct_budget_queue(session, actor_id=pg.cp_actor.id)
    assert any(
        row.id == item_id and row.budget_assignment_id == old_assignment
        for row in queue
    )
    row = view.items[0]
    with pytest.raises(DirectAccessError):
        async with pg.factory.begin() as session:
            await budget.reject_direct_budget_item(
                session,
                actor_id=pg.actor.id,
                item_id=item_id,
                expected_source_version=row.source_version,
                reason="Finance is not the configured CP return authority",
            )
    async with pg.factory.begin() as session:
        returned = await budget.reject_direct_budget_item(
            session,
            actor_id=pg.cp_actor.id,
            item_id=item_id,
            expected_source_version=row.source_version,
            reason="Return the exact assignment after its grant was withdrawn",
        )
    assert returned.status == "rejected", returned.reason
    view = await current(pg, view.statement.id)
    row = view.items[0]
    assert row.budget_assignment_id is None and row.budget_line_id is None
    assert row.budget_status == "classified"
    try:
        async with pg.factory.begin() as session:
            refused = await budget.assign_direct_budget_line(
                session,
                actor_id=pg.cp_actor.id,
                item_id=item_id,
                budget_line_id=pg.line_id,
                expected_source_version=row.source_version,
                expected_source_hash=row.source_hash,
                idempotency_key="must-not-regrant",
            )
    except (DirectAccessError, DirectConflict):
        pass
    else:
        assert refused.status == "pending"
    try:
        async with pg.factory.begin() as session:
            refused = await budget.send_direct_items_to_budget(
                session,
                actor_id=pg.actor.id,
                import_id=view.statement.id,
                item_ids=[item_id],
                expected_versions={str(item_id): row.source_version},
                confirmed=True,
            )
    except (DirectAccessError, DirectConflict):
        pass
    else:
        assert refused.status == "pending"
    async with pg.factory.begin() as session:
        revoked_fiscal = await fiscal.revoke_no_cfdi_items(
            session,
            actor_id=pg.actor.id,
            import_id=view.statement.id,
            item_ids=[item_id],
            expected_versions={str(item_id): row.source_version},
            expected_source_version=view.source_version,
            reason="Release old fiscal treatment after budget return",
            request_key="release-old-benefit-fiscal",
            confirmed=True,
        )
    assert revoked_fiscal.status == "ready", revoked_fiscal.reason
    view = await current(pg, view.statement.id)
    request = decision_request(pg, view, [item_id], "benefit_expense", action="revoke")
    revoked = await apply(
        pg, request, await plan(pg, request), "withdraw-benefit-destination"
    )
    assert revoked.status == "ready", revoked.reason
    decisions = await pg.conn.fetch(
        "SELECT id,action FROM amex_direct_account_treatment_decisions ORDER BY revision"
    )
    assert len(decisions) == 2
    assert decisions[0]["id"] == old_decision and decisions[0]["action"] == "set"
    assert decisions[1]["action"] == "revoke"
    events = await pg.conn.fetch(
        "SELECT payload_json::text FROM amex_direct_events WHERE item_id=$1 AND action='budget_return'",
        item_id,
    )
    assert len(events) == 1
    returned_proof = json.loads(events[0]["payload_json"])
    assert returned_proof["prior"]["assignment_id"] == str(old_assignment)
    assert await pg.conn.fetchval("SELECT count(*) FROM accounting_polizas") == 0
    assert (
        await pg.conn.fetchval(
            "SELECT count(*) FROM amex_direct_account_treatment_policies"
        )
        == 2
    )


@pytest.mark.asyncio
async def test_acceptance_two_benefits_are_one_confirmed_batch_and_budgeted_once(
    direct_pg,
):
    pg = direct_pg
    await setup_policy(pg)
    view, _ = await import_rows(
        pg, "2026-10-02,FIRST BENEFIT,-116.00\n2026-10-03,SECOND BENEFIT,-232.00\n"
    )
    ids = [row.id for row in view.items]
    request = decision_request(pg, view, ids, "benefit_expense")
    planned = await plan(pg, request)
    assert len(planned.progress["manifest"]["items"]) == 2
    accepted = await apply(pg, request, planned, "one-two-benefit-batch")
    assert accepted.status == "ready", accepted.reason
    replay = await apply(pg, request, planned, "one-two-benefit-batch")
    assert replay.status == "exists", replay.reason
    assert (
        await pg.conn.fetchval(
            "SELECT count(*) FROM amex_direct_account_treatment_decisions"
        )
        == 2
    )
    assert (
        await pg.conn.fetchval(
            "SELECT count(*) FROM amex_direct_items WHERE account_treatment_decision_id IS NOT NULL"
        )
        == 2
    )
    view = await current(pg, view.statement.id)
    view = await no_invoice(pg, view, ids)
    view = await cp_assign(pg, view, ids)
    cut = await fixture._confirm(pg, await fixture._cut_args(pg, view))
    assert cut.status == "recognized", cut.reason
    assert await pg.conn.fetchval("SELECT count(*) FROM accounting_polizas") == 1
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM budget_movement_assignments") == 2
    )
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM amex_recognition_consumptions")
        == 2
    )
    assert (
        await pg.conn.fetchval(
            "SELECT count(*) FROM amex_direct_account_treatment_decisions"
        )
        == 2
    )
    actual = await pg.conn.fetchrow(
        "SELECT sum(debe)::numeric debit,sum(haber)::numeric credit FROM accounting_poliza_lines"
    )
    assert actual["debit"] == actual["credit"] == Decimal("348.00")
    projects = await pg.conn.fetch(
        "SELECT i.tournament_id,i.project_code,e.proyecto FROM amex_direct_items i JOIN amex_direct_effective_sources s ON s.id=i.id JOIN expense_reports e ON e.id=s.imported_expense_id"
    )
    assert all(
        r["tournament_id"] is None
        and r["project_code"] is None
        and not str(r["proyecto"] or "").strip()
        for r in projects
    )
