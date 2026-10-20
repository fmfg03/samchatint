"""Private original148-row acceptance with explicit synthetic fiscal decisions.

Set AMEX_QA_ACTIVITY_CSV locally; CI must not embed the private bank statement.
Only external SAT transport is substituted. Original owner receipts stay intact.
"""

import hashlib
import importlib.util
import io
import json
import os
from datetime import date
from decimal import Decimal
from pathlib import Path

import asyncpg
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
_spec = importlib.util.spec_from_file_location(
    "amex_full_statement_owner",
    ROOT / "tests/acceptance/test_amex_sat_budget_direct_postgres.py",
)
owner = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(owner)
direct_pg = owner.direct_pg
REPORT = Path("/tmp/amex-remediation-full-statement-result.json")


def record(key, value):
    data = json.loads(REPORT.read_text()) if REPORT.exists() else {}
    data[key] = value
    REPORT.write_text(json.dumps(data, indent=2, default=str) + "\n")


@pytest.mark.asyncio
async def test_original_statement_recovers_five_and_downloads_all_formats(direct_pg):
    """All original145debits; synthetic fiscal/budget facts only, no serving repair."""
    import time
    from uuid import uuid4

    from sqlalchemy import text
    from test_amex_direct_period_recovery_postgres import seed_previous_period_omission

    from devnous.gastos.services import amex_direct_accounting_service as accounting
    from devnous.gastos.services import amex_direct_intake_service as intake
    from devnous.gastos.services import amex_direct_matching_service as matching
    from devnous.gastos.services import amex_direct_recovery_service as recovery
    from devnous.gastos.services import amex_direct_status_service as status
    from devnous.gastos.services.cfdi_ingestion_service import ingest_cfdi_xml

    pg = direct_pg
    started = time.monotonic()
    path = os.environ.get("AMEX_QA_ACTIVITY_CSV")
    if not path:
        pytest.skip(
            "Set AMEX_QA_ACTIVITY_CSV to the private original; "
            "never commit the statement"
        )
    content = Path(path).read_bytes()
    assert (
        hashlib.sha256(content).hexdigest()
        == "ab641fa6ccafcfc046bf213764f41a61a969d05ce022899d217aa8ab3ea3cd4c"
    )
    parsed = intake.parse_direct_csv(content, "debit_positive", "MXN")
    denied = await owner._direct_import(
        pg,
        content,
        "debit_positive",
        period_start=date(2026, 9, 1),
        period_end=date(2026, 9, 30),
    )
    assert denied.status == "pending" and denied.reason == "statement_period_mismatch"
    assert await owner._count(pg, owner.models.ExpenseReport) == 0
    import_id, version = await seed_previous_period_omission(
        pg,
        content,
        start=date(2026, 9, 1),
        end=date(2026, 9, 30),
        profile="debit_positive",
    )
    originals = await pg.conn.fetch(
        "SELECT id,imported_expense_id,source_json FROM amex_direct_items ORDER BY id"
    )
    original_header = await pg.conn.fetchval(
        "SELECT row_to_json(h)::text FROM amex_statement_imports h WHERE id=$1",
        import_id,
    )
    initial_ids = {
        r["imported_expense_id"] for r in originals if r["imported_expense_id"]
    }
    assert len(originals) == 148 and len(initial_ids) == 140
    recovery_args = dict(
        actor_id=pg.actor.id,
        import_id=import_id,
        period_start=date(2026, 8, 31),
        period_end=date(2026, 9, 30),
        reason="Synthetic QA explicit coverage correction",
        expected_source_version=version,
    )
    async with pg.factory() as session:
        preview = await recovery.preview_direct_period_recovery(
            session, **recovery_args
        )
    assert preview.status == "preview", preview.reason
    assert preview.progress["expected_count"] == 5
    recovery_args.update(
        preview_hash=preview.progress["preview_hash"],
        idempotency_key="full-source-period-recovery",
        confirmed=True,
    )
    async with pg.factory.begin() as session:
        recovered = await recovery.confirm_direct_period_recovery(
            session, **recovery_args
        )
    assert recovered.status == "recovered", recovered.reason
    async with pg.factory.begin() as session:
        replay = await recovery.confirm_direct_period_recovery(session, **recovery_args)
    assert replay.status == "exists"
    assert originals == await pg.conn.fetch(
        "SELECT id,imported_expense_id,source_json FROM amex_direct_items ORDER BY id"
    )
    assert original_header == await pg.conn.fetchval(
        "SELECT row_to_json(h)::text FROM amex_statement_imports h WHERE id=$1",
        import_id,
    )
    assert await pg.conn.fetchval("SELECT count(*) FROM expense_reports") == 145
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM amex_direct_recovered_rows") == 5
    )
    assert await pg.conn.fetchval("SELECT count(*) FROM amex_statement_imports") == 1
    current = await owner._direct_current(pg, import_id)
    assert sum(i.movement_kind == "debit" for i in current.items) == 145
    assert len(current.items) == 148
    pairs = {}
    debit_rows = [row for row in parsed if row["movement_kind"] == "debit"]
    descriptors = set()
    async with pg.factory.begin() as session:
        session.add(
            owner.models.AccountingClosePeriod(
                fiscal_year=2026, fiscal_month=9, status="open"
            )
        )
        await session.execute(
            text("UPDATE budget_lines SET budget_amount=1000000000 WHERE id=:id"),
            {"id": pg.line_id},
        )
        rows_by_number = {r["row_number"]: r for r in debit_rows}
        for item in current.items:
            if item.movement_kind != "debit":
                continue
            source = item.source_json
            row = rows_by_number[source["row_number"]]
            fiscal_uuid = str(uuid4()).upper()
            amount = str(row["amount"])
            stamp = row["date"].strftime("%Y-%m-%d")
            xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<cfdi:Comprobante xmlns:cfdi="http://www.sat.gob.mx/cfd/4"
 xmlns:tfd="http://www.sat.gob.mx/TimbreFiscalDigital" Version="4.0"
 Fecha="{stamp}T10:30:00" Moneda="MXN" SubTotal="{amount}" Total="{amount}"
 TipoDeComprobante="I">
<cfdi:Emisor Rfc="AAA010101AAA" Nombre="SYNTHETIC QA" RegimenFiscal="601"/>
<cfdi:Receptor Rfc="BBB010101BBB" Nombre="SYNTHETIC QA" UsoCFDI="G03"
 DomicilioFiscalReceptor="06600" RegimenFiscalReceptor="601"/>
<cfdi:Conceptos><cfdi:Concepto ClaveProdServ="01010101" Cantidad="1" ClaveUnidad="H87"
 Descripcion="SYNTHETIC no tax assumption" ValorUnitario="{amount}" Importe="{amount}"
 ObjetoImp="01"/></cfdi:Conceptos>
<cfdi:Complemento><tfd:TimbreFiscalDigital UUID="{fiscal_uuid}"
 FechaTimbrado="{stamp}T10:31:00"/></cfdi:Complemento></cfdi:Comprobante>"""
            ingested = await ingest_cfdi_xml(session, xml, source="sat")
            assert ingested.cfdi_report, ingested.reason
            pairs[item.id] = ingested.cfdi_report.id
            descriptors.add(row["description"])
    for descriptor in sorted(descriptors):
        async with pg.factory.begin() as session:
            alias = await matching.confirm_merchant_alias(
                session,
                actor_id=pg.actor.id,
                import_id=import_id,
                descriptor=descriptor,
                issuer_rfc="AAA010101AAA",
                reason=(
                    "SYNTHETIC QA only: invented merchant identity "
                    "to exercise145rows"
                ),
                confirmed=True,
            )
            assert alias.status not in {"pending", "error"}, alias.reason
    observed = owner.ObservedSAT()
    chunks = 0
    for step in range(10):
        current = await owner._direct_current(pg, import_id)
        await status.consult_direct_statement_status(
            pg.factory,
            actor_id=pg.actor.id,
            import_id=import_id,
            expected_source_version=current.source_version,
            request_key="qa-full-status",
            consented=True,
            status_service=observed,
        )
        chunks += 1
        if (
            await pg.conn.fetchval(
                "SELECT count(DISTINCT cfdi_id) FROM amex_cfdi_status_receipts"
            )
            == 145
        ):
            break
    assert (
        await pg.conn.fetchval(
            "SELECT count(DISTINCT cfdi_id) FROM amex_cfdi_status_receipts"
        )
        == 145
    )
    current = await owner._direct_current(pg, import_id)
    automatic = sum(
        i.match_status == "matched" for i in current.items if i.movement_kind == "debit"
    )
    manual = 0
    for identity, cfdi_id in pairs.items():
        current = await owner._direct_current(pg, import_id)
        item = next(i for i in current.items if i.id == identity)
        if item.cfdi_id == cfdi_id and item.match_status == "matched":
            continue
        async with pg.factory.begin() as session:
            linked = await matching.link_direct_cfdi_manual(
                session,
                actor_id=pg.actor.id,
                import_id=import_id,
                item_id=identity,
                cfdi_id=cfdi_id,
                expected_source_version=current.source_version,
                reason=(
                    "SYNTHETIC QA declared1:1row-CFDI identity, "
                    "no real fiscal assertion"
                ),
                confirmed=True,
            )
            assert linked.status == "matched", linked.reason
        manual += 1
    current = await owner._direct_current(pg, import_id)
    for item in current.items:
        if item.movement_kind != "credit":
            continue
        snapshot = await owner._direct_current(pg, import_id)
        async with pg.factory.begin() as session:
            excluded = await matching.exclude_direct_credit(
                session,
                actor_id=pg.actor.id,
                import_id=import_id,
                item_id=item.id,
                expected_source_version=snapshot.source_version,
                reason="SYNTHETIC QA only full-state credit exclusion",
                confirmed=True,
            )
            assert excluded.status != "pending", excluded.reason
    current = await owner._assign_all_items(
        pg, await owner._direct_current(pg, import_id)
    )
    async with pg.factory() as session:
        preview = await accounting.preview_direct_accounting(
            session, actor_id=pg.actor.id, import_id=import_id
        )
    assert preview.status == "ready", (preview.reason, preview.progress)
    args = dict(
        actor_id=pg.actor.id,
        import_id=import_id,
        expected_source_key=preview.source_key,
        expected_source_version=preview.source_version,
        accounting_date=date(2026, 9, 30),
        confirmed=True,
        reason="SYNTHETIC QA full145rowpath",
        idempotency_key="qa-full-cut",
    )
    async with pg.factory.begin() as session:
        cut = await accounting.confirm_direct_accounting(session, **args)
        assert cut.status == "recognized", cut.reason
    async with pg.factory.begin() as session:
        replay = await accounting.confirm_direct_accounting(session, **args)
        assert replay.cut.id == cut.cut.id
    assert await owner._count(pg, owner.models.AccountingPoliza) == 1
    debit = await pg.conn.fetchval("SELECT sum(debe) FROM accounting_poliza_lines")
    credit = await pg.conn.fetchval("SELECT sum(haber) FROM accounting_poliza_lines")
    assert debit == credit
    assert Decimal(str(debit)).quantize(Decimal(".01")) == sum(
        r["amount"] for r in debit_rows
    )
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM amex_recognition_consumptions")
        == 145
    )
    assert (
        await pg.conn.fetchval(
            "SELECT count(*) FROM amex_direct_items "
            "WHERE accounting_status='recognized'"
        )
        == 145
    )
    async with pg.factory() as session:
        xlsx, name, mime = await accounting.export_direct_cut(
            session, actor_id=pg.actor.id, import_id=import_id, format="xlsx"
        )
    artifact = Path("/tmp/amex-remediation-FULL-CSV-SYNTHETIC-poliza-coi.xlsx")
    artifact.write_bytes(xlsx)
    sheet = load_workbook(io.BytesIO(xlsx))["Poliza COI"]
    from devnous.gastos.services import amex_direct_budget_service as budget_owner

    async with pg.factory() as session:
        cp_queue = await budget_owner.list_direct_budget_queue(
            session, actor_id=pg.cp_actor.id
        )
    assert cp_queue == []
    from test_amex_qa_browser_accounting_journey import assert_exported_policy

    uuids = [
        r["cfdi_uuid"]
        for r in await pg.conn.fetch("SELECT cfdi_uuid FROM cfdi_reports")
    ]
    for fmt in ("xlsx", "csv", "zip"):
        async with pg.factory() as session:
            content_bytes, name, mime = await accounting.export_direct_cut(
                session, actor_id=pg.actor.id, import_id=import_id, format=fmt
            )
        assert_exported_policy(
            content_bytes, fmt, uuids, cut.cut.snapshot_json["lines"]
        )
    assert await pg.conn.fetchval("SELECT count(*) FROM accounting_polizas") == 1
    assert await pg.conn.fetchval("SELECT count(*) FROM documentos") == 0
    recovered_expense = await pg.conn.fetchval(
        "SELECT imported_expense_id FROM amex_direct_recovered_rows LIMIT 1"
    )
    assert recovered_expense is not None
    with pytest.raises(asyncpg.CheckViolationError):
        await pg.conn.execute(
            "UPDATE expense_reports SET gasto_cantidad=gasto_cantidad+1 WHERE id=$1",
            recovered_expense,
        )
    assert await pg.conn.fetchval("SELECT count(*) FROM accounting_polizas") == 1
    assert initial_ids <= {
        r["id"] for r in await pg.conn.fetch("SELECT id FROM expense_reports")
    }
    record(
        "recovery_integrity",
        dict(
            original140_expense_ids_preserved=True,
            original148_physical_items_unchanged=True,
            original_header_unchanged=True,
            new_expenses=5,
            effective_membership=148,
            effective_debits=145,
            original_bytes_preserved=True,
            one_import=True,
            recovery_idempotent=True,
            cp_queue_after_cut=0,
            exported_formats=["xlsx", "csv", "zip"],
            recovered_expense_frozen_after_cut=True,
        ),
    )
    record(
        "full_actual_csv_synthetic_dependencies",
        dict(
            label=(
                "Original CSV with synthetic fiscal/budget dependencies; "
                "no production posting"
            ),
            source_sha256=hashlib.sha256(content).hexdigest(),
            explicit_recovery_period=["2026-08-31", "2026-09-30"],
            total_rows=148,
            debits=145,
            credits=3,
            skipped=0,
            synthetic_cfdis=145,
            synthetic_tax_treatment="Zero tax purely fixture, no real IVA assertion",
            synthetic_merchant_aliases=len(descriptors),
            status_transport="SYNTHETIC vigente stub only external transport",
            status_requests=len(observed.requests),
            status_chunks=chunks,
            initial_automatic_matched=automatic,
            manual_link_operations=manual,
            finance_cp_distinct=True,
            budget_assignments=145,
            polizas=1,
            recognized_debits=145,
            consumptions=145,
            journal_balanced=True,
            total_matches_original_debits=True,
            xlsx=str(artifact),
            xlsx_bytes=len(xlsx),
            xlsx_rows=sheet.max_row,
            replay_same_cut=True,
            seconds=round(time.monotonic() - started, 2),
            production_writes=0,
        ),
    )
