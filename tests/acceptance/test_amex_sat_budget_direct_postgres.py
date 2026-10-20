"""Independent direct-AMEX acceptance against owner SQL in disposable schemas.

No runtime credentials, production fixtures, real SAT calls, or startup DDL.
Business APIs stay real; only the external SAT transport may be substituted.
"""

import hashlib
import json
import os
from datetime import date
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4

import asyncpg
import pytest
import pytest_asyncio
from sqlalchemy import ForeignKeyConstraint, MetaData, func, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from devnous.gastos import models
from devnous.gastos.amex_schema_policy import OWNER_MIGRATION_TABLES

ROOT = Path(__file__).resolve().parents[2]
DIRECT_TABLES = frozenset(
    {
        "amex_statement_source_files",
        "amex_statement_direct_contexts",
        "amex_statement_direct_state",
        "amex_merchant_identity_aliases",
        "amex_cfdi_status_receipts",
        "amex_direct_items",
        "amex_direct_events",
        "amex_direct_accounting_cuts",
        "amex_direct_cut_items",
    }
)


@pytest_asyncio.fixture
async def direct_pg(monkeypatch):
    raw = os.getenv("TEST_AMEX_DATABASE_URL")
    if not raw:
        pytest.skip("Set isolated TEST_AMEX_DATABASE_URL")
    url = make_url(raw)
    if (
        url.host not in {"localhost", "127.0.0.1"}
        or url.port != 55471
        or url.username != "test_user"
        or url.database != "devnous_test"
    ):
        pytest.fail(
            "Direct acceptance requires isolated localhost:55471/test_user/devnous_test"
        )
    schema = "direct_acceptance_" + uuid4().hex
    admin = create_async_engine(url)
    async with admin.begin() as conn:
        await conn.execute(text(f'CREATE SCHEMA "{schema}"'))
    engine = create_async_engine(
        url, connect_args={"server_settings": {"search_path": schema}}
    )
    connection = None
    try:
        legacy = MetaData()
        for table in models.Base.metadata.tables.values():
            if table.name in OWNER_MIGRATION_TABLES | DIRECT_TABLES:
                continue
            clone = table.to_metadata(legacy)
            # Unrelated legacy domains do not need operational seeds. AMEX-owned
            # FKs/guards are created by the actual owner SQL below.
            for constraint in list(clone.constraints):
                if isinstance(constraint, ForeignKeyConstraint):
                    clone.constraints.remove(constraint)
                    for element in constraint.elements:
                        clone.foreign_keys.discard(element)
                        element.parent.foreign_keys.discard(element)
        async with engine.begin() as conn:
            await conn.run_sync(legacy.create_all)
        connection = await asyncpg.connect(
            user=url.username,
            password=url.password,
            host=url.host,
            port=url.port,
            database=url.database,
            server_settings={"search_path": schema},
        )
        for filename in (
            "20261001_amex_shared_recognition.sql",
            "20261013_amex_approved_informe_recognition.sql",
            "20261013_amex_statement_workflow.sql",
            "20261014_amex_sat_budget_direct.sql",
            "20261015_amex_direct_period_recovery.sql",
            "20261020_amex_expense_account_settlements.sql",
        ):
            path = ROOT / "database/migrations" / filename
            assert path.is_file(), f"Owner migration absent: {filename}"
            if filename == "20261014_amex_sat_budget_direct.sql":
                from samchat.budgets.service import ensure_budget_schema

                setup_factory = async_sessionmaker(engine, expire_on_commit=False)
                async with setup_factory.begin() as fixture_session:
                    await ensure_budget_schema(fixture_session)
            await connection.execute(path.read_text())
        factory = async_sessionmaker(engine, expire_on_commit=False)
        actor = models.Empleado(
            id=uuid4(), nombre="Contabilidad", rol="finanzas", activo=True
        )
        owner = models.Empleado(
            id=uuid4(), nombre="Synthetic Owner", rol="empleado", activo=True
        )
        liability = models.CuentaContable(
            id=uuid4(), codigo="2120-002-062", nombre="AMEX", tipo="pasivo", activo=True
        )
        card = models.AmexCardAccount(
            id=uuid4(),
            card_label="Synthetic direct",
            cardholder_key="empresa",
            last4="5007",
            liability_cuenta_contable_id=liability.id,
            active=True,
        )
        tournament = models.Tournament(id=uuid4(), name="Acceptance direct")
        cp_actor = models.Empleado(
            id=uuid4(), nombre="Juan Pablo synthetic", rol="empleado", activo=True
        )
        gasto = models.CuentaContable(
            id=uuid4(),
            codigo="5300-001-001",
            nombre="Alimentos",
            tipo="gasto",
            activo=True,
        )
        iva = models.CuentaContable(
            id=uuid4(), codigo="1180-001-001", nombre="IVA", tipo="iva", activo=True
        )
        concept = models.BudgetConcept(
            id=uuid4(),
            tournament_id=tournament.id,
            tournament_name=tournament.name,
            concept_name="Alimentos",
            concept_key="alimentos",
            budget_direction="expense",
            active=True,
            cuenta_contable_id=gasto.id,
        )
        version_id, line_id = uuid4(), uuid4()
        async with factory.begin() as session:
            session.add_all(
                [
                    actor,
                    owner,
                    liability,
                    card,
                    tournament,
                    cp_actor,
                    gasto,
                    iva,
                    concept,
                    models.AccountingClosePeriod(
                        fiscal_year=2026, fiscal_month=10, status="open"
                    ),
                ]
            )
            await session.flush()
            await session.execute(
                text(
                    "INSERT INTO budget_versions(id,edition_year,version_name,status) "
                    "VALUES(:id,2026,'Approved acceptance','approved')"
                ),
                {"id": version_id},
            )
            await session.execute(
                text(
                    "INSERT INTO budget_lines(id,budget_version_id,budget_concept_id,"
                    "tournament_id,tournament_name,concept_name,line_direction,"
                    "budget_amount) VALUES(:id,:version,:concept,:tournament,"
                    "'Acceptance direct','Alimentos','expense',1000)"
                ),
                {
                    "id": line_id,
                    "version": version_id,
                    "concept": concept.id,
                    "tournament": tournament.id,
                },
            )
        scopes = [
            {
                "actor_id": str(person.id),
                "tournament_ids": [str(tournament.id)],
                "edition_years": [2026],
            }
            for person in (actor, cp_actor)
        ]
        monkeypatch.setenv("ASSISTANT_ACCOUNTING_IVA_CUENTA", iva.codigo)
        monkeypatch.setenv(
            "AMEX_DIRECT_AUTHORIZED_CONTEXTS",
            json.dumps(
                [
                    {
                        "entity_code": "TEST",
                        "company_code": "TEST",
                        "receiver_rfc": "BBB010101BBB",
                        "card_account_ids": [str(card.id)],
                        "tournament_ids": [str(tournament.id)],
                        "edition_years": [2026],
                        "actor_scopes": scopes,
                        "budget_actor_ids": [str(cp_actor.id)],
                    }
                ]
            ),
        )
        yield SimpleNamespace(
            factory=factory,
            conn=connection,
            actor=actor,
            owner=owner,
            card=card,
            liability=liability,
            schema=schema,
            monkeypatch=monkeypatch,
            tournament=tournament,
            cp_actor=cp_actor,
            gasto=gasto,
            iva=iva,
            concept=concept,
            version_id=version_id,
            line_id=line_id,
        )
    finally:
        if connection is not None:
            if connection.is_in_transaction():
                await connection.execute("ROLLBACK")
            await connection.close()
        await engine.dispose()
        async with admin.begin() as conn:
            await conn.execute(text(f'DROP SCHEMA "{schema}" CASCADE'))
        await admin.dispose()


async def _direct_import(
    pg,
    content,
    profile="debit_negative",
    key="acceptance-intake",
    period_start=date(2026, 10, 1),
    period_end=date(2026, 10, 31),
):
    from devnous.gastos.services.amex_direct_intake_service import (
        import_direct_statement,
    )

    async with pg.factory.begin() as session:
        return await import_direct_statement(
            session,
            actor_id=pg.actor.id,
            employee_id=pg.owner.id,
            card_account_id=pg.card.id,
            period_start=period_start,
            period_end=period_end,
            content=content,
            filename="original.csv",
            mime_type="text/csv",
            profile=profile,
            company_code="TEST",
            receiver_rfc="BBB010101BBB",
            currency="MXN",
            idempotency_key=key,
            confirmed=True,
        )


async def _count(pg, model):
    async with pg.factory() as session:
        return await session.scalar(select(func.count()).select_from(model))


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "profile,amount,kind",
    [
        ("debit_negative", "-116.00", "debit"),
        ("debit_positive", "116.00", "debit"),
        ("debit_negative", "116.00", "credit"),
        ("debit_positive", "-116.00", "credit"),
    ],
)
async def test_signed_intake_preserves_exact_file_and_never_fabricates_informe(
    direct_pg, profile, amount, kind
):
    pg = direct_pg
    original = (
        "\ufeffDate,Description,Amount\r\n2026-10-02,COMERCIO ACME CDMX,"
        + amount
        + "\r\n"
    ).encode()
    result = await _direct_import(pg, original, profile)
    assert result.status in {"created", "imported"}, result.reason
    statement = result.statement
    source = await pg.conn.fetchrow(
        "SELECT * FROM amex_statement_source_files WHERE import_id=$1", statement.id
    )
    assert source["content"] == original
    assert source["file_sha256"] == hashlib.sha256(original).hexdigest()
    assert source["byte_size"] == len(original)
    assert await _count(pg, models.Documento) == 0
    assert await _count(pg, models.AccountingPoliza) == 0
    assert await _count(pg, models.ExpenseReport) == int(kind == "debit")
    rows = await pg.conn.fetch(
        "SELECT * FROM amex_direct_items WHERE import_id=$1", statement.id
    )
    assert len(rows) == 1 and rows[0]["movement_kind"] == kind
    assert rows[0]["raw_signed_amount"] == Decimal(amount)
    assert rows[0]["profile"] == profile
    assert statement.imported_count == int(kind == "debit")
    assert statement.credit_count == int(kind == "credit")
    replay = await _direct_import(pg, original, profile)
    assert replay.status in {"exists", "existing"}, replay.reason
    assert replay.statement.id == statement.id
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM amex_statement_source_files") == 1
    )


def test_direct_owner_tables_never_participate_in_runtime_create_all():
    from devnous.gastos.amex_schema_policy import runtime_managed_tables

    declared = set(models.Base.metadata.tables)
    assert (
        DIRECT_TABLES <= declared
    ), "Direct tables need explicit owner models before acceptance"
    managed = {table.name for table in runtime_managed_tables(models.Base.metadata)}
    assert not (DIRECT_TABLES & managed)


@pytest.mark.asyncio
async def test_all_unknown_keeps_explicit_missing_sign_without_zero_gasto(direct_pg):
    original = b"Date,Description,Amount\n2026-10-02,COMERCIO ACME CDMX,116.00\n"
    result = await _direct_import(direct_pg, original, profile=None)
    assert result.status in {"created", "imported"}, result.reason
    assert result.statement.imported_count == 0
    assert result.statement.unknown_count == 1
    item = await direct_pg.conn.fetchrow("SELECT * FROM amex_direct_items")
    assert item["movement_kind"] == "unknown"
    assert item["imported_expense_id"] is None
    assert item["profile"] is None
    assert await _count(direct_pg, models.ExpenseReport) == 0
    assert await _count(direct_pg, models.AccountingPoliza) == 0
    assert await _count(direct_pg, models.Documento) == 0


def _synthetic_sat_xml(uuid="10000000-0000-4000-8000-000000000001"):
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<cfdi:Comprobante xmlns:cfdi="http://www.sat.gob.mx/cfd/4"
 xmlns:tfd="http://www.sat.gob.mx/TimbreFiscalDigital"
 Version="4.0" Fecha="2026-10-03T10:30:00" Moneda="MXN"
 SubTotal="100.00" Total="116.00" TipoDeComprobante="I">
 <cfdi:Emisor Rfc="AAA010101AAA" Nombre="ACME" RegimenFiscal="601"/>
 <cfdi:Receptor Rfc="BBB010101BBB" Nombre="RECEPTOR" UsoCFDI="G03"
  DomicilioFiscalReceptor="06600" RegimenFiscalReceptor="601"/>
 <cfdi:Conceptos><cfdi:Concepto ClaveProdServ="01010101" Cantidad="1"
  ClaveUnidad="H87" Descripcion="Alimentos" ValorUnitario="100.00"
  Importe="100.00" ObjetoImp="02"><cfdi:Impuestos><cfdi:Traslados>
  <cfdi:Traslado Base="100.00" Impuesto="002" TipoFactor="Tasa"
   TasaOCuota="0.160000" Importe="16.00"/>
 </cfdi:Traslados></cfdi:Impuestos></cfdi:Concepto></cfdi:Conceptos>
 <cfdi:Impuestos TotalImpuestosTrasladados="16.00"><cfdi:Traslados>
  <cfdi:Traslado Base="100.00" Impuesto="002" TipoFactor="Tasa"
   TasaOCuota="0.160000" Importe="16.00"/>
 </cfdi:Traslados></cfdi:Impuestos>
 <cfdi:Complemento><tfd:TimbreFiscalDigital UUID="{uuid}"
  FechaTimbrado="2026-10-03T10:31:00"/></cfdi:Complemento>
</cfdi:Comprobante>"""


class ObservedSAT:
    """Explicit external service substitute; never asserts live SAT validity."""

    def __init__(self, status="vigente"):
        self.status = status
        self.requests = []

    async def consult_status(self, **request):
        self.requests.append(request)
        return {
            "status": self.status,
            "request": request,
            "raw": {"Estado": self.status, "fixture": "synthetic"},
        }


async def _matched_statement(
    pg, *, alias=True, status="vigente", observed_override=None, content=None
):
    from devnous.gastos.services import amex_direct_context_service as context
    from devnous.gastos.services import amex_direct_matching_service as matching
    from devnous.gastos.services import amex_direct_status_service as status_owner
    from devnous.gastos.services.cfdi_ingestion_service import ingest_cfdi_xml

    direct = SimpleNamespace(
        list_direct_statement=context.list_direct_statement,
        confirm_merchant_alias=matching.confirm_merchant_alias,
        reconcile_direct_statement=matching.reconcile_direct_statement,
        consult_direct_statement_status=status_owner.consult_direct_statement_status,
    )

    async with pg.factory.begin() as session:
        ingested = await ingest_cfdi_xml(session, _synthetic_sat_xml(), source="sat")
    imported = await _direct_import(
        pg,
        content or b"Date,Description,Amount\n2026-10-02,COMERCIO ACME CDMX,-116.00\n",
    )
    assert imported.status in {"created", "imported"}, imported.reason
    if alias:
        async with pg.factory.begin() as session:
            confirmed = await direct.confirm_merchant_alias(
                session,
                actor_id=pg.actor.id,
                import_id=imported.statement.id,
                descriptor="COMERCIO ACME CDMX",
                issuer_rfc="AAA010101AAA",
                reason="Exact synthetic merchant declaration",
                confirmed=True,
            )
            assert confirmed.status not in {"pending", "error"}, confirmed.reason
    async with pg.factory() as session:
        before_status = await direct.list_direct_statement(
            session, actor_id=pg.actor.id, import_id=imported.statement.id
        )
    observed = observed_override or ObservedSAT(status)
    await direct.consult_direct_statement_status(
        pg.factory,
        actor_id=pg.actor.id,
        import_id=imported.statement.id,
        expected_source_version=before_status.source_version,
        request_key="acceptance-observed-status",
        consented=True,
        status_service=observed,
    )
    async with pg.factory.begin() as session:
        current = await direct.list_direct_statement(
            session, actor_id=pg.actor.id, import_id=imported.statement.id
        )
        matched = await direct.reconcile_direct_statement(
            session,
            actor_id=pg.actor.id,
            import_id=imported.statement.id,
            expected_source_version=current.source_version,
        )
    return matched, observed, ingested.cfdi_report


@pytest.mark.asyncio
async def test_ordinary_csv_no_uuid_uses_confirmed_alias_and_observed_status(direct_pg):
    result, observed, cfdi = await _matched_statement(direct_pg)
    assert len(observed.requests) == 1
    assert observed.requests[0]["uuid"] == cfdi.cfdi_uuid
    assert await _count(direct_pg, models.Documento) == 0
    assert await _count(direct_pg, models.AccountingPoliza) == 0
    item = await direct_pg.conn.fetchrow("SELECT * FROM amex_direct_items")
    assert item["declared_uuid"] is None
    assert item["match_status"] == "matched", item["match_reason"]
    assert item["cfdi_id"] == cfdi.id
    assert item["merchant_alias_id"] is not None
    assert item["status_receipt_id"] is not None
    assert item["budget_status"] == "unclassified"
    assert item["accounting_status"] == "pending"


@pytest.mark.asyncio
@pytest.mark.parametrize("alias,status", [(False, "vigente"), (True, "cancelado")])
async def test_matching_never_promotes_economic_similarity_or_cancelled_cfdi(
    direct_pg, alias, status
):
    await _matched_statement(direct_pg, alias=alias, status=status)
    item = await direct_pg.conn.fetchrow("SELECT * FROM amex_direct_items")
    assert item["match_status"] != "matched"
    assert item["match_reason"]
    assert item["cfdi_id"] is None
    assert await _count(direct_pg, models.AccountingPoliza) == 0


async def _direct_current(pg, import_id):
    from devnous.gastos.services import amex_direct_context_service as direct

    async with pg.factory() as session:
        return await direct.list_direct_statement(
            session, actor_id=pg.actor.id, import_id=import_id
        )


async def _assign_all_items(pg, current):
    from devnous.gastos.services import amex_direct_budget_service as direct

    import_id = current.statement.id
    selected = [item for item in current.items if item.movement_kind == "debit"]
    async with pg.factory.begin() as session:
        classified = await direct.classify_direct_items(
            session,
            actor_id=pg.actor.id,
            import_id=import_id,
            item_ids=[item.id for item in selected],
            tournament_id=pg.tournament.id,
            project_code=None,
            budget_concept_id=pg.concept.id,
            edition_year=2026,
            expected_versions={str(item.id): item.source_version for item in selected},
            reason="Explicit accounting project/category",
        )
        assert classified.status not in {"pending", "error"}, classified.reason
    current = await _direct_current(pg, import_id)
    selected = [item for item in current.items if item.movement_kind == "debit"]
    async with pg.factory.begin() as session:
        submitted = await direct.send_direct_items_to_budget(
            session,
            actor_id=pg.actor.id,
            import_id=import_id,
            item_ids=[item.id for item in selected],
            expected_versions={str(item.id): item.source_version for item in selected},
            confirmed=True,
        )
        assert submitted.status not in {"pending", "error"}, submitted.reason
    current = await _direct_current(pg, import_id)
    for item in current.items:
        if item.movement_kind != "debit":
            continue
        async with pg.factory.begin() as session:
            assigned = await direct.assign_direct_budget_line(
                session,
                actor_id=pg.cp_actor.id,
                item_id=item.id,
                budget_line_id=pg.line_id,
                expected_source_version=item.source_version,
                expected_source_hash=item.source_hash,
                idempotency_key="acceptance-CP-line-" + str(item.id),
            )
            assert assigned.status not in {"pending", "error"}, assigned.reason
    return await _direct_current(pg, import_id)


async def _assigned_statement(pg):
    matched, observed, cfdi = await _matched_statement(pg)
    return await _assign_all_items(pg, await _direct_current(pg, matched.statement.id))


@pytest.mark.asyncio
async def test_actual_classification_cp_assignment_and_atomic_direct_cut(direct_pg):
    import io
    from openpyxl import load_workbook
    from devnous.gastos.services import amex_direct_accounting_service as direct

    pg = direct_pg
    current = await _assigned_statement(pg)
    item = current.items[0]
    assert item.tournament_id == pg.tournament.id
    assert item.budget_concept_id == pg.concept.id
    assert item.budget_line_id == pg.line_id
    source = await pg.conn.fetchrow(
        "SELECT * FROM expense_reports WHERE id=$1", item.imported_expense_id
    )
    assert source["cuenta_contable_id"] == pg.gasto.id
    assert await _count(pg, models.Documento) == 0
    assert await _count(pg, models.AccountingPoliza) == 0
    async with pg.factory() as session:
        preview = await direct.preview_direct_accounting(
            session, actor_id=pg.actor.id, import_id=current.statement.id
        )
    assert preview.status == "ready", preview.reason
    args = dict(
        actor_id=pg.actor.id,
        import_id=current.statement.id,
        expected_source_key=preview.source_key,
        expected_source_version=preview.source_version,
        accounting_date=date(2026, 10, 31),
        confirmed=True,
        reason="Complete synthetic direct cut",
        idempotency_key="acceptance-direct-cut",
    )
    async with pg.factory.begin() as session:
        result = await direct.confirm_direct_accounting(session, **args)
        assert result.status in {"created", "recognized"}, result.reason
    async with pg.factory.begin() as session:
        retry = await direct.confirm_direct_accounting(session, **args)
        assert retry.status in {"exists", "existing"}, retry.reason
        assert retry.cut.id == result.cut.id
    assert await _count(pg, models.AccountingPoliza) == 1
    assert await _count(pg, models.Documento) == 0
    assert await _count(pg, models.ExpenseReport) == 1
    lines = await pg.conn.fetch("SELECT * FROM accounting_poliza_lines")
    assert sum(Decimal(str(line["debe"] or 0)) for line in lines) == Decimal("116")
    assert sum(Decimal(str(line["haber"] or 0)) for line in lines) == Decimal("116")
    assert sum(
        Decimal(str(line["debe"] or 0))
        for line in lines
        if line["cuenta_codigo"] == pg.gasto.codigo
    ) == Decimal("100")
    assert {line["cuenta_codigo"] for line in lines if line["haber"]} == {
        pg.liability.codigo
    }
    consumption = await pg.conn.fetchrow("SELECT * FROM amex_recognition_consumptions")
    assert consumption["binding_mode"] == "direct"
    assert consumption["source_kind"] == "statement"
    assert consumption["imported_expense_id"] == item.imported_expense_id
    assert (
        await pg.conn.fetchval(
            "SELECT count(*) FROM amex_recognition_representations WHERE role='report'"
        )
        == 0
    )
    async with pg.factory() as session:
        exported = await direct.export_direct_cut(
            session, actor_id=pg.actor.id, import_id=current.statement.id, format="xlsx"
        )
    workbook = load_workbook(io.BytesIO(exported[0]))
    rows = list(workbook["Poliza COI"].values)
    fiscal_rows = [
        row for row in rows if row[8] == "10000000-0000-4000-8000-000000000001"
    ]
    assert len(fiscal_rows) == 1
    assert fiscal_rows[0][5:7] == ("AAA010101AAA", "BBB010101BBB")
    await pg.conn.execute(
        "UPDATE cuentas_contables SET nombre='Later catalogue name' WHERE id=$1",
        pg.gasto.id,
    )
    async with pg.factory() as session:
        repeated = await direct.export_direct_cut(
            session, actor_id=pg.actor.id, import_id=current.statement.id, format="xlsx"
        )
    assert list(load_workbook(io.BytesIO(repeated[0]))["Poliza COI"].values) == rows

    from devnous.gastos.services.amex_cut_export_service import cut_expense_cfdis

    async with pg.factory() as session:
        actual_cut = await session.get(models.AmexDirectAccountingCut, result.cut.id)
        frozen_rows = cut_expense_cfdis(actual_cut)
    assert len(frozen_rows) == 1
    assert frozen_rows[0].poliza_group_key == f"amex-direct-cut:{result.cut.id}"
    assert frozen_rows[0].poliza_description == (
        f"Corte AMEX directo · Estado {current.statement.id}"
    )
    assert not hasattr(actual_cut, "informe_id")


@pytest.mark.asyncio
async def test_mixed_credit_preserves_signed_bank_value_without_invented_expense(
    direct_pg,
):
    from devnous.gastos.services.amex_direct_matching_service import (
        exclude_direct_credit,
    )

    pg = direct_pg
    matched, observed, cfdi = await _matched_statement(
        pg,
        content=(
            b"Date,Description,Amount\n"
            b"2026-10-02,COMERCIO ACME CDMX,-116.00\n"
            b"2026-10-03,KNOWN CREDIT,20.00\n"
        ),
    )
    current = await _direct_current(pg, matched.statement.id)
    credit = next(item for item in current.items if item.movement_kind == "credit")
    async with pg.factory.begin() as session:
        result = await exclude_direct_credit(
            session,
            actor_id=pg.actor.id,
            import_id=current.statement.id,
            item_id=credit.id,
            expected_source_version=current.source_version,
            reason="Explicit known bank credit, no new expense",
            confirmed=True,
        )
        assert result.status == "ready", result.reason
    current = await _assign_all_items(
        pg, await _direct_current(pg, current.statement.id)
    )
    result = await _confirm(pg, await _cut_args(pg, current))
    assert result.status == "recognized", result.reason
    frozen_credit = next(
        item
        for item in result.cut.snapshot_json["partidas"]
        if item["coverage_kind"] == "credit_exclusion"
    )
    assert frozen_credit["expense_id"] is None
    assert frozen_credit["amount"] is None
    assert Decimal(frozen_credit["bank_signed_amount"]) == Decimal("20")
    assert frozen_credit["journal_lines"] == []
    assert await _count(pg, models.ExpenseReport) == 1
    assert await _count(pg, models.Documento) == 0
    assert await _count(pg, models.AccountingPoliza) == 1


@pytest.mark.asyncio
async def test_sql_original_and_context_receipts_are_append_only(direct_pg):
    pg = direct_pg
    imported = await _direct_import(
        pg, b"Date,Description,Amount\n2026-10-02,ACME,-116.00\n"
    )
    assert imported.status in {"created", "imported"}, imported.reason
    with pytest.raises(asyncpg.PostgresError):
        await pg.conn.execute(
            "UPDATE amex_statement_source_files SET content=$1 WHERE import_id=$2",
            b"replacement",
            imported.statement.id,
        )
    with pytest.raises(asyncpg.PostgresError):
        await pg.conn.execute(
            "UPDATE amex_statement_direct_contexts "
            "SET receiver_rfc='CCC010101CCC' WHERE import_id=$1",
            imported.statement.id,
        )
    with pytest.raises(asyncpg.PostgresError):
        await pg.conn.execute(
            "DELETE FROM amex_statement_direct_contexts WHERE import_id=$1",
            imported.statement.id,
        )
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM amex_statement_direct_contexts")
        == 1
    )
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM amex_statement_source_files") == 1
    )


@pytest.mark.asyncio
@pytest.mark.parametrize("denial", ["inactive", "same_name", "explicit_rule"])
async def test_actual_access_rules_and_actor_identity_deny_intake_without_effects(
    direct_pg, denial
):
    from sqlalchemy import update
    from devnous.gastos.services.access_control_service import (
        empleado_area,
        empleado_role,
    )

    pg = direct_pg
    async with pg.factory.begin() as session:
        if denial == "explicit_rule":
            session.add(
                models.AccessControlRule(
                    tool_key="amex_direct",
                    action_key="cargar",
                    role_key=empleado_role(pg.actor),
                    area_key=empleado_area(pg.actor),
                    allowed=False,
                    active=True,
                )
            )
        else:
            values = (
                {"activo": False}
                if denial == "inactive"
                else {"rol": "empleado", "nombre": "Juan Pablo"}
            )
            await session.execute(
                update(models.Empleado)
                .where(models.Empleado.id == pg.actor.id)
                .values(**values)
            )
    with pytest.raises(PermissionError):
        await _direct_import(pg, b"Date,Description,Amount\n2026-10-02,ACME,-116.00\n")
    assert await _count(pg, models.ExpenseReport) == 0
    assert await _count(pg, models.Documento) == 0
    assert await _count(pg, models.AccountingPoliza) == 0
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM amex_statement_source_files") == 0
    )


@pytest.mark.asyncio
async def test_original_download_is_exact_and_denial_is_reloaded(direct_pg):
    from devnous.gastos.services.amex_direct_intake_service import (
        download_direct_source,
    )
    from devnous.gastos.services.access_control_service import (
        empleado_area,
        empleado_role,
    )

    pg = direct_pg
    original = b"Date,Description,Amount\r\n2026-10-02,ACME,116.00\r\n"
    imported = await _direct_import(pg, original)
    async with pg.factory() as session:
        content, filename, mime = await download_direct_source(
            session, actor_id=pg.actor.id, import_id=imported.statement.id
        )
        assert content == original and filename == "original.csv" and mime == "text/csv"
    async with pg.factory.begin() as session:
        session.add(
            models.AccessControlRule(
                tool_key="amex_direct",
                action_key="fuente",
                role_key=empleado_role(pg.actor),
                area_key=empleado_area(pg.actor),
                allowed=False,
                active=True,
            )
        )
    async with pg.factory() as session:
        with pytest.raises(PermissionError):
            await download_direct_source(
                session, actor_id=pg.actor.id, import_id=imported.statement.id
            )


@pytest.mark.asyncio
async def test_sat_runs_outside_financial_locks_and_revalidates_revoked_actor(
    direct_pg,
):
    from sqlalchemy import update

    pg = direct_pg

    class RevokingSAT(ObservedSAT):
        async def consult_status(self, **request):
            # NOWAIT makes the absence of the import lock observable during the
            # external call; this does not rely on a guessed sleep duration.
            async with pg.factory.begin() as session:
                await session.execute(
                    text("SELECT id FROM amex_statement_imports FOR UPDATE NOWAIT")
                )
                await session.execute(
                    update(models.Empleado)
                    .where(models.Empleado.id == pg.actor.id)
                    .values(activo=False)
                )
            return await super().consult_status(**request)

    observed = RevokingSAT()
    with pytest.raises(PermissionError):
        await _matched_statement(pg, observed_override=observed)
    assert len(observed.requests) == 1
    item = await pg.conn.fetchrow("SELECT * FROM amex_direct_items")
    assert item["match_status"] != "matched" and item["cfdi_id"] is None
    assert await _count(pg, models.AccountingPoliza) == 0
    assert await _count(pg, models.Documento) == 0


@pytest.mark.asyncio
async def test_one_cfdi_two_ordinary_charges_remain_globally_ambiguous(direct_pg):
    original = (
        b"Date,Description,Amount\n"
        b"2026-10-02,COMERCIO ACME CDMX,-116.00\n"
        b"2026-10-02,COMERCIO ACME CDMX,-116.00\n"
    )
    await _matched_statement(direct_pg, content=original)
    rows = await direct_pg.conn.fetch("SELECT * FROM amex_direct_items")
    assert len(rows) == 2
    assert all(
        row["match_status"] != "matched" and row["cfdi_id"] is None for row in rows
    )
    assert await _count(direct_pg, models.ExpenseReport) == 2
    assert await _count(direct_pg, models.AccountingPoliza) == 0
    assert await _count(direct_pg, models.Documento) == 0


async def _cut_args(pg, current):
    from devnous.gastos.services.amex_direct_accounting_service import (
        preview_direct_accounting,
    )

    async with pg.factory() as session:
        preview = await preview_direct_accounting(
            session, actor_id=pg.actor.id, import_id=current.statement.id
        )
    assert preview.status == "ready", preview.reason
    return dict(
        actor_id=pg.actor.id,
        import_id=current.statement.id,
        expected_source_key=preview.source_key,
        expected_source_version=preview.source_version,
        accounting_date=date(2026, 10, 31),
        confirmed=True,
        reason="Acceptance complete cut",
        idempotency_key="acceptance-canonical-cut",
    )


async def _confirm(pg, args):
    from devnous.gastos.services.amex_direct_accounting_service import (
        confirm_direct_accounting,
    )

    async with pg.factory.begin() as session:
        return await confirm_direct_accounting(session, **args)


@pytest.mark.asyncio
async def test_concurrent_direct_confirmations_share_one_cut_and_journal(direct_pg):
    import asyncio

    pg = direct_pg
    current = await _assigned_statement(pg)
    args = await _cut_args(pg, current)
    results = await asyncio.wait_for(
        asyncio.gather(_confirm(pg, args), _confirm(pg, args)), timeout=10
    )
    assert sorted(result.status for result in results) == ["existing", "recognized"]
    assert results[0].cut.id == results[1].cut.id
    assert await _count(pg, models.AccountingPoliza) == 1
    assert await _count(pg, models.AmexRecognitionConsumption) == 1
    assert await _count(pg, models.Documento) == 0


@pytest.mark.asyncio
async def test_cut_failure_after_journal_flush_rolls_back_every_new_accounting_effect(
    direct_pg,
):
    from sqlalchemy.exc import DBAPIError

    pg = direct_pg
    current = await _assigned_statement(pg)
    args = await _cut_args(pg, current)
    source_before = await pg.conn.fetchrow("SELECT * FROM expense_reports")
    await pg.conn.execute(
        """CREATE FUNCTION acceptance_fail_cut() RETURNS trigger
      LANGUAGE plpgsql AS $$ BEGIN
        RAISE EXCEPTION 'acceptance fail after journal'; END $$;
      CREATE TRIGGER acceptance_fail_cut BEFORE INSERT ON amex_direct_accounting_cuts
      FOR EACH ROW EXECUTE FUNCTION acceptance_fail_cut();"""
    )
    with pytest.raises(DBAPIError, match="acceptance fail after journal"):
        await _confirm(pg, args)
    assert await _count(pg, models.AccountingPoliza) == 0
    assert await _count(pg, models.AccountingPolizaLine) == 0
    assert await _count(pg, models.AmexRecognitionConsumption) == 0
    assert await _count(pg, models.AmexDirectAccountingCut) == 0
    assert await _count(pg, models.AmexDirectCutItem) == 0
    assert dict(await pg.conn.fetchrow("SELECT * FROM expense_reports")) == dict(
        source_before
    )
    assert await _count(pg, models.Documento) == 0
    await pg.conn.execute(
        "DROP TRIGGER acceptance_fail_cut ON amex_direct_accounting_cuts"
    )
    assert (await _confirm(pg, args)).status == "recognized"


@pytest.mark.asyncio
async def test_intake_failure_after_source_persistence_leaves_no_file_or_debit(
    direct_pg,
):
    from sqlalchemy.exc import DBAPIError

    pg = direct_pg
    await pg.conn.execute(
        """CREATE FUNCTION acceptance_fail_context() RETURNS trigger
      LANGUAGE plpgsql AS $$ BEGIN RAISE EXCEPTION 'acceptance context failure'; END $$;
      CREATE TRIGGER acceptance_fail_context
      BEFORE INSERT ON amex_statement_direct_contexts
      FOR EACH ROW EXECUTE FUNCTION acceptance_fail_context();"""
    )
    with pytest.raises(DBAPIError, match="acceptance context failure"):
        await _direct_import(pg, b"Date,Description,Amount\n2026-10-02,ACME,-116.00\n")
    assert await _count(pg, models.ExpenseReport) == 0
    assert await _count(pg, models.Documento) == 0
    assert await _count(pg, models.AccountingPoliza) == 0
    for table in (
        "amex_statement_imports",
        "amex_statement_import_items",
        "amex_statement_source_files",
        "amex_direct_items",
        "amex_direct_events",
    ):
        assert await pg.conn.fetchval(f"SELECT count(*) FROM {table}") == 0


@pytest.mark.asyncio
async def test_all_credits_excluded_still_never_create_zero_journal(direct_pg):
    from devnous.gastos.services.amex_direct_matching_service import (
        exclude_direct_credit,
    )
    from devnous.gastos.services.amex_direct_accounting_service import (
        preview_direct_accounting,
    )

    pg = direct_pg
    imported = await _direct_import(
        pg, b"Date,Description,Amount\n2026-10-02,ACME,116.00\n"
    )
    current = await _direct_current(pg, imported.statement.id)
    async with pg.factory.begin() as session:
        excluded = await exclude_direct_credit(
            session,
            actor_id=pg.actor.id,
            import_id=imported.statement.id,
            item_id=current.items[0].id,
            expected_source_version=current.source_version,
            reason="Known credit is not new expense",
            confirmed=True,
        )
        assert excluded.status == "ready", excluded.reason
    async with pg.factory() as session:
        preview = await preview_direct_accounting(
            session, actor_id=pg.actor.id, import_id=imported.statement.id
        )
    assert preview.status == "pending" and preview.reason
    assert await _count(pg, models.AccountingPoliza) == 0
    assert await _count(pg, models.ExpenseReport) == 0
    assert await _count(pg, models.Documento) == 0


@pytest.mark.asyncio
async def test_real_budget_reader_replaces_pending_projection_with_one_fiscal_journal(
    direct_pg,
):
    from samchat.budgets.service import build_budget_actuals_snapshot

    pg = direct_pg
    current = await _assigned_statement(pg)

    async def snapshot():
        async with pg.factory() as session:
            return await build_budget_actuals_snapshot(
                session,
                edition_year=2026,
                version_id=str(pg.version_id),
                tournament_id=str(pg.tournament.id),
                tournament_name=pg.tournament.name,
            )

    before = await snapshot()
    pending = [
        row
        for row in before["movements"]
        if row.get("expense_id") == str(current.items[0].imported_expense_id)
    ]
    assert len(pending) == 1
    args = await _cut_args(pg, current)
    assert (await _confirm(pg, args)).status == "recognized"
    after = await snapshot()
    actual = [
        row
        for row in after["movements"]
        if row.get("expense_id") == str(current.items[0].imported_expense_id)
    ]
    assert len(actual) == 1, after
    assert actual[0]["kind"] == "ledger_expense"
    assert actual[0]["accounting_line_id"] is not None
    assert Decimal(str(actual[0]["amount"])) == Decimal("100"), actual[0]


async def _cross_calendar_state(pg, *, two_debits=True):
    from devnous.gastos.services.cfdi_ingestion_service import ingest_cfdi_xml
    from devnous.gastos.services.amex_direct_matching_service import (
        confirm_merchant_alias,
        reconcile_direct_statement,
    )
    from devnous.gastos.services.amex_direct_status_service import (
        consult_direct_statement_status,
    )

    first = _synthetic_sat_xml().replace("2026-10-03", "2026-09-30")
    second = (
        _synthetic_sat_xml("10000000-0000-4000-8000-000000000002")
        .replace("2026-10-03", "2026-10-02")
        .replace("100.00", "200.00")
        .replace("116.00", "232.00")
        .replace("16.00", "32.00")
    )
    async with pg.factory.begin() as session:
        await ingest_cfdi_xml(session, first, source="sat")
        if two_debits:
            await ingest_cfdi_xml(session, second, source="sat")
    original = b"Date,Description,Amount\n2026-09-29,COMERCIO ACME CDMX,-116.00\n"
    if two_debits:
        original += b"2026-10-01,COMERCIO ACME CDMX,-232.00\n"
    imported = await _direct_import(
        pg, original, period_start=date(2026, 9, 29), period_end=date(2026, 10, 29)
    )
    assert imported.status in {"created", "imported"}, imported.reason
    async with pg.factory.begin() as session:
        await confirm_merchant_alias(
            session,
            actor_id=pg.actor.id,
            import_id=imported.statement.id,
            descriptor="COMERCIO ACME CDMX",
            issuer_rfc="AAA010101AAA",
            reason="Declared exact merchant",
            confirmed=True,
        )
    current = await _direct_current(pg, imported.statement.id)
    await consult_direct_statement_status(
        pg.factory,
        actor_id=pg.actor.id,
        import_id=imported.statement.id,
        expected_source_version=current.source_version,
        request_key="acceptance-cross-calendar",
        consented=True,
        status_service=ObservedSAT(),
    )
    current = await _direct_current(pg, imported.statement.id)
    async with pg.factory.begin() as session:
        matched = await reconcile_direct_statement(
            session,
            actor_id=pg.actor.id,
            import_id=imported.statement.id,
            expected_source_version=current.source_version,
        )
    assert all(item.match_status == "matched" for item in matched.items)
    return await _assign_all_items(pg, await _direct_current(pg, imported.statement.id))


@pytest.mark.asyncio
async def test_two_economic_months_one_declared_state_single_october_policy(direct_pg):
    pg = direct_pg
    current = await _cross_calendar_state(pg)
    args = await _cut_args(pg, current)
    args["accounting_date"] = date(2026, 10, 15)
    result = await _confirm(pg, args)
    assert result.status == "recognized", result.reason
    assert result.cut.accounting_date == date(2026, 10, 15)
    assert await _count(pg, models.AccountingPoliza) == 1
    assert await _count(pg, models.Documento) == 0
    dates = await pg.conn.fetch("SELECT fecha FROM expense_reports ORDER BY fecha")
    assert [row["fecha"].date() for row in dates] == [
        date(2026, 9, 29),
        date(2026, 10, 1),
    ]
    assert (
        await pg.conn.fetchval("SELECT sum(debe) FROM accounting_poliza_lines") == 348
    )
    assert (
        await pg.conn.fetchval("SELECT sum(haber) FROM accounting_poliza_lines") == 348
    )


@pytest.mark.asyncio
async def test_declared_accounting_date_closed_rejects_even_when_economic_month_open(
    direct_pg,
):
    pg = direct_pg
    current = await _cross_calendar_state(pg, two_debits=False)
    args = await _cut_args(pg, current)
    args["accounting_date"] = date(2026, 10, 15)
    async with pg.factory.begin() as session:
        session.add(
            models.AccountingClosePeriod(
                fiscal_year=2026, fiscal_month=10, status="closed"
            )
        )
    result = await _confirm(pg, args)
    assert result.status == "pending" and result.reason
    assert await _count(pg, models.AccountingPoliza) == 0
    assert await _count(pg, models.AmexRecognitionConsumption) == 0
    assert (await pg.conn.fetchval("SELECT fecha FROM expense_reports")).date() == date(
        2026, 9, 29
    )
    assert await _count(pg, models.Documento) == 0


@pytest.mark.asyncio
async def test_context_revision_releases_own_application_and_assignment(
    direct_pg,
):
    from devnous.gastos.services.amex_direct_context_service import set_direct_context

    pg = direct_pg
    current = await _assigned_statement(pg)
    contexts = json.loads(os.environ["AMEX_DIRECT_AUTHORIZED_CONTEXTS"])
    revised = dict(contexts[0], receiver_rfc="CCC010101CCC")
    pg.monkeypatch.setenv(
        "AMEX_DIRECT_AUTHORIZED_CONTEXTS", json.dumps(contexts + [revised])
    )
    before = await pg.conn.fetchrow("SELECT * FROM amex_statement_direct_contexts")
    async with pg.factory.begin() as session:
        result = await set_direct_context(
            session,
            actor_id=pg.actor.id,
            import_id=current.statement.id,
            company_code="TEST",
            receiver_rfc="CCC010101CCC",
            currency="MXN",
            profile="debit_negative",
            expected_source_version=current.source_version,
            reason="Explicit receiver correction before recognition",
            confirmed=True,
        )
    assert result.status == "ready", result.reason
    history = await pg.conn.fetch(
        "SELECT * FROM amex_statement_direct_contexts ORDER BY version"
    )
    assert len(history) == 2 and dict(history[0]) == dict(before)
    assert history[1]["version"] == 2 and history[1]["receiver_rfc"] == "CCC010101CCC"
    item = await pg.conn.fetchrow("SELECT * FROM amex_direct_items")
    assert item["context_id"] == history[1]["id"]
    assert item["match_status"] != "matched" and item["cfdi_id"] is None
    assert (
        item["budget_status"] == "unclassified" and item["budget_assignment_id"] is None
    )
    source = await pg.conn.fetchrow("SELECT * FROM expense_reports")
    assert source["cfdi_report_id"] is None and source["cfdi_uuid_manual"] is None
    assert (
        await pg.conn.fetchval("SELECT count(*) FROM budget_movement_assignments") == 1
    )
    assert (
        await pg.conn.fetchval(
            "SELECT count(*) FROM amex_direct_events WHERE action='budget_assignment'"
        )
        == 1
    )
    assert await _count(pg, models.AccountingPoliza) == 0
    # Historical SQL assignment remains evidence; it is not a current decision.
    from devnous.gastos.services.amex_direct_budget_service import classify_direct_items
    from samchat.budgets.service import build_budget_actuals_snapshot

    refreshed = await _direct_current(pg, current.statement.id)
    current_item = refreshed.items[0]
    async with pg.factory.begin() as session:
        classified = await classify_direct_items(
            session,
            actor_id=pg.actor.id,
            import_id=current.statement.id,
            item_ids=[current_item.id],
            tournament_id=pg.tournament.id,
            project_code=None,
            budget_concept_id=pg.concept.id,
            edition_year=2026,
            expected_versions={str(current_item.id): current_item.source_version},
            reason="Explicit reclassification after receiver correction",
        )
        assert classified.status == "classified", classified.reason
    async with pg.factory() as session:
        snapshot = await build_budget_actuals_snapshot(
            session,
            edition_year=2026,
            version_id=str(pg.version_id),
            tournament_id=str(pg.tournament.id),
            tournament_name=pg.tournament.name,
        )
    movements = [
        row
        for row in snapshot["movements"]
        if row.get("expense_id") == str(current_item.imported_expense_id)
    ]
    assert len(movements) == 1
    assert movements[0].get("budget_line_id") is None, movements[0]
    assert movements[0]["kind"] == "pending_fiscal_evidence"
    assert movements[0]["amount"] is None


@pytest.mark.asyncio
@pytest.mark.parametrize("with_cut", [False, True])
async def test_payment_owner_schedules_and_posts_only_liability(direct_pg, with_cut):
    from devnous.gastos.services.amex_card_payment_service import (
        AmexCardPaymentRequest,
        create_amex_card_payment_request,
    )
    from devnous.gastos.services.amex_accounting_posting_service import (
        ensure_amex_payment_posting,
    )

    pg = direct_pg
    if with_cut:
        current = await _assigned_statement(pg)
        assert (await _confirm(pg, await _cut_args(pg, current))).status == "recognized"
    else:
        await _direct_import(pg, b"Date,Description,Amount\n2026-10-02,ACME,-116.00\n")
    journal_before = await _count(pg, models.AccountingPoliza)
    async with pg.factory.begin() as session:
        session.add(
            models.CuentaContable(
                id=uuid4(),
                codigo="1120-001-001",
                nombre="Santander",
                tipo="banco",
                activo=True,
            )
        )
    async with pg.factory() as session:
        actor = await session.get(models.Empleado, pg.actor.id)
        scheduled = await create_amex_card_payment_request(
            session,
            actor=actor,
            request=AmexCardPaymentRequest(
                card_account_id=pg.card.id,
                amount=Decimal("80"),
                fecha_pago=date(2026, 10, 20),
            ),
        )
    assert scheduled.documento.tipo == "SOLICITUD"
    assert scheduled.documento.estado == "aprobado"
    assert scheduled.documento.pagado_en is None
    assert scheduled.documento.monto_solicitado == 80
    assert await _count(pg, models.AccountingPoliza) == journal_before
    assert await _count(pg, models.ExpenseReport) == 1
    async with pg.factory.begin() as session:
        documento = await session.get(models.Documento, scheduled.documento.id)
        paid_posting = await ensure_amex_payment_posting(
            session, documento=documento, payment_date=date(2026, 10, 20)
        )
        assert paid_posting.status == "created", paid_posting.reason
    lines = await pg.conn.fetch(
        "SELECT cuenta_codigo,debe,haber FROM accounting_poliza_lines "
        "WHERE poliza_id=$1 ORDER BY line_no",
        paid_posting.poliza.id,
    )
    assert [
        (row["cuenta_codigo"], row["debe"] or 0, row["haber"] or 0) for row in lines
    ] == [(pg.liability.codigo, 80, 0), ("1120-001-001", 0, 80)]
    assert await _count(pg, models.ExpenseReport) == 1
    assert await _count(pg, models.AccountingPoliza) == journal_before + 1


@pytest.mark.asyncio
async def test_sql_cut_projection_and_membership_cannot_be_rewritten(direct_pg):
    pg = direct_pg
    current = await _assigned_statement(pg)
    result = await _confirm(pg, await _cut_args(pg, current))
    assert result.status == "recognized"
    original = await pg.conn.fetchrow("SELECT * FROM amex_direct_accounting_cuts")
    for command in (
        "UPDATE amex_direct_items SET accounting_status='pending'",
        "UPDATE amex_statement_direct_state SET source_version=source_version+1",
        "DELETE FROM amex_direct_cut_items",
        "DELETE FROM amex_direct_accounting_cuts",
    ):
        with pytest.raises(asyncpg.PostgresError):
            await pg.conn.execute(command)
    assert dict(
        await pg.conn.fetchrow("SELECT * FROM amex_direct_accounting_cuts")
    ) == dict(original)
    assert await _count(pg, models.AccountingPoliza) == 1


@pytest.mark.asyncio
async def test_unresolved_original_row_blocks_whole_cut_even_when_debit_ready(
    direct_pg,
):
    from devnous.gastos.services.amex_direct_accounting_service import (
        preview_direct_accounting,
    )

    pg = direct_pg
    original = (
        b"Date,Description,Amount\n"
        b"2026-10-02,COMERCIO ACME CDMX,-116.00\n"
        b"2026-10-02,Unresolved original amount,not-money\n"
    )
    matched, observed, cfdi = await _matched_statement(pg, content=original)
    current = await _assign_all_items(
        pg, await _direct_current(pg, matched.statement.id)
    )
    assert len(current.items) == 2
    assert sum(item.budget_status == "assigned" for item in current.items) == 1
    async with pg.factory() as session:
        preview = await preview_direct_accounting(
            session, actor_id=pg.actor.id, import_id=current.statement.id
        )
    assert preview.status == "pending", preview.progress
    assert "statement_row_unresolved" in preview.progress["exceptions"]
    result = await _confirm(
        pg,
        dict(
            actor_id=pg.actor.id,
            import_id=current.statement.id,
            expected_source_key=preview.source_key,
            expected_source_version=preview.source_version,
            accounting_date=date(2026, 10, 31),
            confirmed=True,
            reason="Full original scope",
            idempotency_key="acceptance-pending-full-state",
        ),
    )
    assert result.status == "pending"
    assert await _count(pg, models.AccountingPoliza) == 0
    assert await _count(pg, models.AmexDirectAccountingCut) == 0
    assert await _count(pg, models.AmexRecognitionConsumption) == 0
    assert await _count(pg, models.ExpenseReport) == 1
    assert await _count(pg, models.Documento) == 0


@pytest.mark.asyncio
async def test_observed_receipt_ttl_boundary_and_changed_xml_invalidate_status(
    direct_pg,
):
    from datetime import timedelta
    from sqlalchemy import update
    from devnous.gastos.services.amex_direct_matching_service import (
        current_status_receipt,
    )

    pg = direct_pg
    result, observed, cfdi = await _matched_statement(pg)
    async with pg.factory() as session:
        stored = await session.get(models.CFDIReport, cfdi.id)
        receipt = await current_status_receipt(session, stored)
        assert receipt is not None
        assert (
            await current_status_receipt(
                session, stored, now=receipt.observed_at + timedelta(hours=24)
            )
            is not None
        )
        assert (
            await current_status_receipt(
                session,
                stored,
                now=receipt.observed_at + timedelta(hours=24, microseconds=1),
            )
            is None
        )
    async with pg.factory.begin() as session:
        await session.execute(
            update(models.CFDIReport)
            .where(models.CFDIReport.id == cfdi.id)
            .values(xml_raw=cfdi.xml_raw + "\n")
        )
    async with pg.factory() as session:
        changed = await session.get(models.CFDIReport, cfdi.id)
        assert await current_status_receipt(session, changed) is None
    assert await _count(pg, models.AmexCFDIStatusReceipt) == 1
    assert await _count(pg, models.AccountingPoliza) == 0


@pytest.mark.asyncio
async def test_classify_budget_confirm_export_execute_no_schema_ddl(direct_pg):
    from sqlalchemy import event
    from devnous.gastos.services.amex_direct_accounting_service import export_direct_cut

    pg = direct_pg
    matched, observed, cfdi = await _matched_statement(pg)
    current = await _direct_current(pg, matched.statement.id)
    engine = pg.factory.kw["bind"]
    ddl = []

    def observe_sql(_connection, _cursor, statement, _parameters, _context, _many):
        if statement.lstrip().split(None, 1)[0].upper() in {
            "CREATE",
            "ALTER",
            "DROP",
            "TRUNCATE",
        }:
            ddl.append(statement)

    event.listen(engine.sync_engine, "before_cursor_execute", observe_sql)
    try:
        assigned = await _assign_all_items(pg, current)
        assert (
            await _confirm(pg, await _cut_args(pg, assigned))
        ).status == "recognized"
        async with pg.factory() as session:
            content, name, mime = await export_direct_cut(
                session,
                actor_id=pg.actor.id,
                import_id=current.statement.id,
                format="xlsx",
            )
        assert content and name.endswith(".xlsx")
    finally:
        event.remove(engine.sync_engine, "before_cursor_execute", observe_sql)
    assert ddl == [], "Business actions must consume owner-installed schema"
