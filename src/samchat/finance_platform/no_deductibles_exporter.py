"""XLSX export for the No Deducibles control."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def generate_no_deductibles_xlsx(report: dict[str, Any]) -> bytes:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Resumen"
    summary_sheet.append(["Control de No Deducibles"])
    summary_sheet["A1"].font = Font(size=16, bold=True)
    period = report.get("period") or {}
    summary = report.get("summary") or {}
    summary_sheet.append(["Periodo de fecha de gasto", f"{period.get('year')}-{int(period.get('month') or 0):02d}"])
    summary_sheet.append(["Gasto total", summary.get("total_amount") or 0])
    summary_sheet.append(["Deducible (CFDI vinculado)", summary.get("deductible_amount") or 0])
    summary_sheet.append(["No deducible (sin CFDI)", summary.get("non_deductible_amount") or 0])
    summary_sheet.append(["% no deducible", (summary.get("non_deductible_percent") or 0) / 100])
    summary_sheet["B6"].number_format = "0.00%"
    for cell in (summary_sheet["B3"], summary_sheet["B4"], summary_sheet["B5"]):
        cell.number_format = '$#,##0.00'
    summary_sheet.column_dimensions["A"].width = 34
    summary_sheet.column_dimensions["B"].width = 22

    detail_sheet = workbook.create_sheet("Detalle no deducible")
    headers = [
        "Fecha gasto", "Torneo", "Fase", "Documento", "Referencia documento",
        "Gasto", "Concepto", "Responsable", "Monto", "Estatus", "Motivo", "UUID CFDI",
    ]
    header_fill = PatternFill("solid", fgColor="0F766E")
    detail_sheet.append(headers)
    for cell in detail_sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    for row in report.get("non_deductible_rows") or []:
        detail_sheet.append([
            str(row.get("expense_date") or "")[:10], row.get("tournament_name") or "",
            row.get("phase") or "", row.get("source_type") or "",
            row.get("source_reference") or "", row.get("reference") or "",
            row.get("concept") or "", row.get("employee_name") or "", row.get("amount") or 0,
            "No deducible", row.get("fiscal_reason") or "", row.get("cfdi_uuid") or "",
        ])
    for row in detail_sheet.iter_rows(min_row=2, min_col=9, max_col=9):
        row[0].number_format = '$#,##0.00'
    detail_sheet.freeze_panes = "A2"
    for column, width in {"A": 14, "B": 28, "C": 18, "D": 14, "E": 22, "F": 22, "G": 36, "H": 28, "I": 15, "J": 16, "K": 38, "L": 38}.items():
        detail_sheet.column_dimensions[column].width = width

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
