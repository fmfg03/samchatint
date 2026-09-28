"""Review-only DR workpaper assembled from an informe's active expenses.

This is deliberately separate from the importable COI export: incomplete
accounting classification must not silently discard expenses or create postings.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


@dataclass(frozen=True)
class InformeWorkpaperExpense:
    source_id: str
    reference: str
    date: str
    description: str
    amount: float
    vat: float | None
    expense_account: str
    vat_account: str
    counterpart_account: str
    company_amex: bool
    cfdi_uuid: str


def _money(value: float) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _safe_text(value: str) -> str:
    """Keep untrusted source descriptions as literal text in Excel."""
    text = " ".join(str(value or "").split())
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def generate_informe_poliza_workpaper(
    expenses: Sequence[InformeWorkpaperExpense],
    *,
    reference: str,
    title: str,
    currency: str,
) -> bytes:
    """Build one balanced, editable DR-style paper with live review formulas."""
    if not expenses:
        raise ValueError("El informe no contiene gastos activos.")
    if currency.upper() != "MXN":
        raise ValueError("El papel DR requiere importes MXN; revise la moneda del informe.")

    wb = Workbook()
    wb.calculation.calcMode = "auto"
    sheet = wb.active
    sheet.title = "Papel DR"
    sheet["A3"] = "Dr"
    sheet["B3"] = 3
    sheet["C3"] = _safe_text(f"{reference} / {title}")
    sheet["D3"] = sum(1 + (1 if _money(item.vat or 0) > 0 else 0) for item in expenses) + len(
        {item.company_amex for item in expenses}
    )

    origin = wb.create_sheet("Origen y revisión")
    origin.append(
        [
            "ID gasto", "Referencia", "Fecha", "Concepto", "Importe",
            "IVA capturado", "CFDI UUID", "Medio", "Cuenta gasto guardada",
            "Cuenta IVA guardada", "Contrapartida guardada",
        ]
    )

    totals: dict[bool, Decimal] = {False: Decimal("0.00"), True: Decimal("0.00")}
    group_accounts: dict[bool, set[str]] = {False: set(), True: set()}
    row = 4
    for expense in expenses:
        gross = _money(expense.amount)
        vat = _money(expense.vat or 0)
        if gross <= 0 or vat < 0 or vat > gross:
            raise ValueError(f"Importe o IVA inválido en gasto {expense.reference}.")
        group = bool(expense.company_amex)
        totals[group] += gross
        group_accounts[group].add(expense.counterpart_account.strip())
        description = _safe_text(
            " / ".join(part for part in (expense.reference, expense.description) if part)
        )
        for account, concept, debit in (
            (expense.expense_account, description, gross - vat),
            (expense.vat_account, _safe_text(f"IVA / {description}"), vat),
        ):
            if debit == 0:
                continue
            sheet.append([])  # Rows 1-3 are the DR header, then one movement per row.
            sheet.cell(row, 2, account.strip() or None)
            sheet.cell(row, 3, "0")
            sheet.cell(row, 4, concept)
            sheet.cell(row, 5, "1")
            sheet.cell(row, 6, float(debit))
            row += 1
        origin.append(
            [
                expense.source_id,
                _safe_text(expense.reference),
                expense.date,
                _safe_text(expense.description),
                float(gross),
                float(vat) if expense.vat is not None else "Sin desglose",
                _safe_text(expense.cfdi_uuid),
                "AMEX empresa" if group else "Empleado / anticipo / reembolso",
                expense.expense_account,
                expense.vat_account if vat else "",
                expense.counterpart_account,
            ]
        )

    for group in (False, True):
        if not totals[group]:
            continue
        candidates = group_accounts[group]
        account = next(iter(candidates)) if len(candidates) == 1 else ""
        sheet.cell(row, 2, account or None)
        sheet.cell(row, 3, "0")
        sheet.cell(
            row,
            4,
            "Contrapartida AMEX empresa" if group else "Contrapartida empleado / anticipo / reembolso",
        )
        sheet.cell(row, 5, "1")
        sheet.cell(row, 7, float(totals[group]))
        row += 1
    last = row - 1
    sheet.cell(row, 2, "FIN_PARTIDAS")

    review = wb.create_sheet("Cuadre")
    review.append(["Papel de trabajo / revisión contable", "Valor"])
    review.append(["Informe", _safe_text(reference)])
    review.append(["Moneda", currency.upper()])
    review.append(["Gastos activos incluidos", len(expenses)])
    review.append(["Debe", f"=SUM('Papel DR'!F4:F{last})"])
    review.append(["Haber", f"=SUM('Papel DR'!G4:G{last})"])
    review.append(["Diferencia", "=ROUND(B5-B6,2)"])
    review.append(["Cuentas pendientes", f'=COUNTBLANK(\'Papel DR\'!B4:B{last})'])
    review.append(
        ["Estado", '=IF(AND(ABS(B7)<0.01,B8=0),"CUADRE Y CUENTAS COMPLETAS","PENDIENTE DE REVISIÓN")']
    )
    review.append(["Origen", "Sólo gastos activos del informe; no crea una póliza contable."])
    review.append(
        ["IVA", "Se desglosa únicamente si existe un importe de IVA capturado en el gasto."]
    )
    review.append(
        ["Contrapartida", "Separada por AMEX y otros medios; revisar anticipo/reembolso antes de COI."]
    )
    review.append(["Carga COI", "Revisar cuentas, impuestos y estado del informe antes de importar."])

    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 8
    sheet.column_dimensions["D"].width = 82
    sheet.column_dimensions["E"].width = 8
    sheet.column_dimensions["F"].width = 18
    sheet.column_dimensions["G"].width = 18
    sheet.freeze_panes = "D4"
    for movement_row in sheet.iter_rows(min_row=4, max_row=last):
        for amount_cell in (movement_row[5], movement_row[6]):
            amount_cell.number_format = '#,##0.00'
    for tab in (origin, review):
        for cell in tab[1]:
            cell.fill = PatternFill("solid", fgColor="1F2937")
            cell.font = Font(color="FFFFFF", bold=True)
        tab.freeze_panes = "A2"
    for column in "ABCDEFGHIJK":
        origin.column_dimensions[column].width = 22
    origin.column_dimensions["D"].width = 55
    review.column_dimensions["A"].width = 32
    review.column_dimensions["B"].width = 82
    for i in (5, 6, 7):
        review.cell(i, 2).number_format = '#,##0.00'
    wb.active = 0
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
