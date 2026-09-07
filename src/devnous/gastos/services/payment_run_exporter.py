"""Workbook export for a closed Payment Run."""

from __future__ import annotations

import io
from collections import defaultdict
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill("solid", fgColor="0F766E")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(size=16, bold=True)


def _text(value: Any) -> str:
    return "" if value is None else str(value)


def _safe_cell_text(value: Any) -> str:
    """Prevent spreadsheet programs from evaluating untrusted text as formulas."""
    text = _text(value)
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def _money(value: Any) -> float:
    try:
        return round(float(value or 0), 2)
    except (TypeError, ValueError):
        return 0.0


def _totals_by_currency(items: list[dict[str, Any]]) -> str:
    totals: dict[str, float] = defaultdict(float)
    for item in items:
        currency = _text(item.get("currency") or "MXN").upper()
        totals[currency] += _money(item.get("monto"))
    return " | ".join(
        f"{currency} {amount:,.2f}" for currency, amount in sorted(totals.items())
    ) or "Sin partidas"


def generate_payment_run_order_xlsx(*, closure: dict[str, Any]) -> bytes:
    """Build the controlled payment-order workbook for one closed cutoff."""
    items = list(closure.get("items") or [])
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Orden de pago"

    sheet["A1"] = "Orden de pago - SamChat"
    sheet["A1"].font = _TITLE_FONT
    sheet["A2"] = f"Corte: {_text(closure.get('id'))}"
    sheet["A3"] = f"Fecha de corte: {_text(closure.get('run_date') or '-')}"
    sheet["A4"] = f"Totales por moneda: {_totals_by_currency(items)}"

    headers = [
        "Solicitud",
        "Referencia operaciones",
        "Solicitante",
        "Beneficiario",
        "Banco",
        "Cuenta",
        "CLABE",
        "Concepto",
        "Fecha programada",
        "Moneda",
        "Monto pagable",
        "Estatus datos de pago",
    ]
    header_row = 6
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    for row_number, item in enumerate(
        items, start=header_row + 1
    ):
        row = [
            _safe_cell_text(item.get("numero_referencia")),
            _safe_cell_text(item.get("referencia_operaciones")),
            _safe_cell_text(item.get("solicitante")),
            _safe_cell_text(item.get("beneficiario")),
            _safe_cell_text(item.get("banco")),
            _safe_cell_text(item.get("cuenta_bancaria")),
            _safe_cell_text(item.get("cuenta_clabe")),
            _safe_cell_text(item.get("concepto_pago")),
            _safe_cell_text(item.get("fecha_pago")),
            _safe_cell_text(item.get("currency") or "MXN"),
            _money(item.get("monto")),
            _safe_cell_text(item.get("payment_data_status")),
        ]
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_number, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.cell(row=row_number, column=11).number_format = '#,##0.00'

    for column in range(1, len(headers) + 1):
        max_length = max(
            len(_text(sheet.cell(row=row, column=column).value))
            for row in range(1, sheet.max_row + 1)
        )
        sheet.column_dimensions[get_column_letter(column)].width = min(
            max(max_length + 2, 14),
            42,
        )
    sheet.freeze_panes = "A7"
    sheet.auto_filter.ref = f"A{header_row}:L{sheet.max_row}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
